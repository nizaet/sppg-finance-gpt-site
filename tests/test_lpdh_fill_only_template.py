"""Synthetic fill-only contract tests; no private source or live financial writes."""
import sys
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import unittest
import base64
from contextlib import ExitStack
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch
from copy import deepcopy
import test_generated_document_api as fixture
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from backend.lpdh_logic import compute_preview, fallback_lpdh_workbook
from backend.lpdh_template import (prepare_template, fill_template, input_cells,
                                   HEADERS, YELLOW, NS, source_hash)
from backend.generated_document_logic import merge_final_documents
from backend import lpdh_api
from fastapi import HTTPException


def synthetic_template():
    wb = fallback_lpdh_workbook()
    for (s,c),label in HEADERS.items(): wb[s][c] = label
    # Native source financial formulas and ordinary yellow inputs.
    for s,cc in input_cells().items():
        if s in ('I_RegisterBukti','Lampiran_Dokumen'): continue
        for c in cc:
            wb[s][c] = None
            wb[s][c].fill = PatternFill('solid',fgColor=YELLOW)
    wb['Identitas']['B15'].number_format='dd-mm-yyyy'
    wb['B_BahanBaku']['I6'] = '=IF(OR(F6="",H6=""),0,F6*H6)'
    wb['B_BahanBaku']['I7'] = ArrayFormula(ref='I7',text='=F7*H7')
    wb['I_RegisterBukti']['C5'] = '=B_BahanBaku!K6'
    wb['Petunjuk']['A10'] = 'UNRELATED CONTENT MUST STAY'
    out = BytesIO(); wb.save(out); return out.getvalue()


class FillOnlyTests(unittest.TestCase):
    def setUp(self):
        self.raw = synthetic_template()
        self.raw_digest = source_hash(self.raw)
        self.template = prepare_template(self.raw)
        doc = {'id':12,'site':'MAJA','serviceDate':'2026-10-05','documentType':'UPAH_RELAWAN',
               'documentNumber':'047/TEST/X/2026','status':'FINAL','header':{'paymentSnapshotVersion':2},
               'total':200,'items':[{'itemName':n,'quantity':1,'unit':'hari','unitPrice':100,'lineTotal':100,'metadata':{}} for n in ('Synthetic A','Synthetic B')]}
        self.daily = merge_final_documents({},[doc])
        self.daily['lpdhNumber'] = '001/LPDH/TEST/X/2026'
        for x in self.daily['volunteerPayments']:
            x.update(evidenceLink='https://example.test/signed.pdf',paymentReference='SHARED-REF')
        self.preview = compute_preview({},self.daily,'2026-10-05',False)

    def fill(self): return fill_template(self.template,{},self.daily,self.preview,'2026-10-05')

    def test_exports_only_change_yellow_inputs_and_keep_all_formulas_styles_parts(self):
        output = self.fill()
        before = load_workbook(BytesIO(self.template)); after = load_workbook(BytesIO(output))
        self.assertEqual(before.sheetnames,after.sheetnames)
        allow = input_cells()
        for s in before.sheetnames:
            self.assertEqual(list(before[s].merged_cells.ranges),list(after[s].merged_cells.ranges))
            for row in before[s]:
                for cell in row:
                    target = after[s][cell.coordinate]
                    self.assertEqual(cell._style,target._style,(s,cell.coordinate))
                    if cell.data_type == 'f':
                        signature=lambda v:v if isinstance(v,str) else dict(v)
                        self.assertEqual(signature(cell.value),signature(target.value),(s,cell.coordinate))
                    elif cell.value != target.value:
                        self.assertIn(cell.coordinate,allow.get(s,set()),(s,cell.coordinate))
                        self.assertEqual(cell.fill.fgColor.rgb,YELLOW)
                        self.assertNotEqual(cell.data_type,'f')
        with ZipFile(BytesIO(self.template)) as a,ZipFile(BytesIO(output)) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if not name.startswith('xl/worksheets/sheet'):
                    self.assertEqual(a.read(name),b.read(name),name)
                else:
                    aa=ET.fromstring(a.read(name)); bb=ET.fromstring(b.read(name))
                    for element in ('mergeCells','dataValidations','conditionalFormatting','sheetViews','cols','pageSetup'):
                        self.assertEqual([ET.tostring(x) for x in aa.findall('s:'+element,NS)],
                                         [ET.tostring(x) for x in bb.findall('s:'+element,NS)],(name,element))
        self.assertEqual(source_hash(self.raw),self.raw_digest)

    def test_aggregate_receipt_one_register_entry_shared_number_and_reference(self):
        w = load_workbook(BytesIO(self.fill()))
        for r in (6,7):
            self.assertEqual(w['C1_Relawan'][f'J{r}'].value,'047/TEST/X/2026')
            self.assertEqual(w['C1_Relawan'][f'K{r}'].value,'https://example.test/signed.pdf')
        self.assertEqual([w['I_RegisterBukti'][f'C{r}'].value for r in range(5,131)].count('047/TEST/X/2026'),1)
        self.assertEqual(w['I_RegisterBukti']['E5'].value,200)
        for r in (2,3): self.assertEqual(w['Lampiran_Dokumen'][f'I{r}'].value,'SHARED-REF')

    def test_no_preparation_or_formula_creation_during_fill(self):
        from unittest.mock import patch
        with patch('backend.lpdh_template.prepare_template',side_effect=AssertionError('must not prepare on download')):
            self.fill()
        with self.assertRaisesRegex(ValueError,'belum disiapkan'):
            fill_template(self.raw,{},self.daily,self.preview,'2026-10-05')

    def test_changed_mapping_formula_input_and_capacity_rejected(self):
        for c,v in [('D5','SHIFTED HEADER'),('D6','=1+1')]:
            w=load_workbook(BytesIO(self.raw)); w['B_BahanBaku'][c]=v
            out=BytesIO(); w.save(out)
            with self.assertRaises(ValueError): prepare_template(out.getvalue())
        self.preview['rawMaterials']=[{'qty':1,'price':1}]*41
        with self.assertRaisesRegex(ValueError,'kapasitas'):
            self.fill()
        w=load_workbook(BytesIO(self.raw)); w['A_PM']['S6'].fill=PatternFill('solid',fgColor=YELLOW)
        out=BytesIO(); w.save(out)
        with self.assertRaisesRegex(ValueError,'belum dipetakan'): prepare_template(out.getvalue())

    def test_literal_strings_dates_zero_and_fresh_day_no_old_values(self):
        self.daily['lpdhNumber']='=SUM(1,2)'
        w=load_workbook(BytesIO(self.fill()))
        self.assertEqual(w['Identitas']['B5'].data_type,'s')
        self.assertEqual(w['Identitas']['B5'].value,'=SUM(1,2)')
        self.assertEqual(w['Identitas']['B15'].value.date().isoformat(),'2026-10-05')
        p=compute_preview({}, {}, '2026-10-06',False)
        next_day=load_workbook(BytesIO(fill_template(self.template,{}, {},p,'2026-10-06')))
        self.assertIsNone(next_day['C1_Relawan']['C6'].value)
        self.assertIsNone(next_day['Lampiran_Dokumen']['I2'].value)
        self.assertEqual(next_day['D_Insentif']['C21'].value,0)
        self.daily['lpdhNumber']='invalid\x00text'
        with self.assertRaisesRegex(ValueError,'tidak kompatibel'): self.fill()


class GenerationTemplateTests(unittest.TestCase):
    def setUp(self):
        raw=synthetic_template()
        self.masters={'_officialTemplateBase64':base64.b64encode(raw).decode(),
                      '_preparedTemplateBase64':base64.b64encode(prepare_template(raw)).decode(),
                      '_preparedTemplateSourceHash':source_hash(raw),'_preparedTemplateVersion':1}
        self.cur=fixture.FakeConnection()
        self.request=SimpleNamespace(state=SimpleNamespace(sppg_role='MAJA'))
        self.payload=lpdh_api.GenerateIn(site='MAJA',service_date=date(2026,10,5))
        self.stack=ExitStack(); self.addCleanup(self.stack.close)
        p=compute_preview({}, {}, '2026-10-05',True); p.update(ready=True)
        for name,value in {'_load_master':lambda site:{'data':self.masters},
                          '_load_daily':lambda *args:{'data':{},'status':'DRAFT'},
                          '_load_final_plan':lambda *args:{'payload':{}},
                          '_daily_with_hpe':lambda site,day,daily:(daily,{'effective':True}),
                          'daily_number':lambda *args:'001/LPDH/TEST/X/2026',
                          'claim_number':lambda *args:None,'compute_preview':lambda *args:p}.items():
            self.stack.enter_context(patch.object(lpdh_api,name,value))
        self.stack.enter_context(patch.object(fixture.api,'load_documents',return_value=[]))

    def generate(self): return lpdh_api._generate_locked(self.payload,self.request,'MAJA',self.cur)

    def test_draft_with_failed_checks_does_not_issue_or_lock_report(self):
        self.payload.draft_only=True
        p=compute_preview({}, {}, '2026-10-05',False)
        p.update(ready=False,errorCount=4)
        with patch.object(lpdh_api,'compute_preview',return_value=p), patch.object(lpdh_api,'_load_final_plan',return_value=None), patch.object(lpdh_api,'_daily_with_hpe',side_effect=lambda site,day,daily:(daily,{'effective':False})), patch.object(lpdh_api,'daily_number',side_effect=AssertionError('draft must not reserve a number')), patch.object(lpdh_api,'claim_number',side_effect=AssertionError('draft must not claim a number')):
            result=self.generate()
        self.assertTrue(result['draft'])
        self.assertIn('_DRAFT_',result['filename'])
        self.assertEqual(result['validation']['errorCount'],4)
        self.assertFalse(any('insert into lpdh_daily_state' in sql or 'insert into lpdh_generation_log' in sql for sql,_ in self.cur.calls))
        w=load_workbook(BytesIO(base64.b64decode(result['contentBase64'])))
        self.assertTrue(any(cell.data_type=='f' for sheet in w for row in sheet for cell in row))

    def test_installed_template_download_does_not_prepare_or_build_formulas(self):
        with patch.object(lpdh_api,'prepare_template',side_effect=AssertionError('no preparation')):
            result=self.generate()
        w=load_workbook(BytesIO(base64.b64decode(result['contentBase64'])))
        self.assertEqual(w['Identitas']['B5'].value,'001/LPDH/TEST/X/2026')
        self.assertFalse(any(sql.startswith('update lpdh_site_state') for sql,_ in self.cur.calls))

    def test_existing_upload_prepared_once_source_kept(self):
        original=self.masters['_officialTemplateBase64']
        del self.masters['_preparedTemplateBase64']
        self.generate()
        update=next(params for sql,params in self.cur.calls if sql.startswith('update lpdh_site_state'))
        self.assertEqual(update[2],original)
        self.assertEqual(self.masters['_officialTemplateBase64'],original)

    def test_missing_or_corrupt_template_blocks_no_financial_state_written(self):
        for data in ({},{'_officialTemplateBase64':'not-base64'}):
            self.masters=data; self.cur.calls=[]
            with self.assertRaises(HTTPException) as err: self.generate()
            self.assertEqual(err.exception.status_code,409)
            self.assertFalse(any('insert into lpdh_daily_state' in sql or 'insert into lpdh_generation_log' in sql for sql,_ in self.cur.calls))


if __name__=='__main__': unittest.main()
