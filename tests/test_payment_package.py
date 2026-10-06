"""Synthetic package fixtures: no identities or live transactions."""
import sys
import unittest
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from payment_package import school_daily_rate, incentive_default, payment_sections
from generated_document_logic import merge_final_documents
from test_generated_document_api import api
GeneratedDocumentIn = api.GeneratedDocumentIn
from generated_document_pdf import render_document_pdf
from generated_document_excel import render_document_excel
from generated_document_selftest import fixtures
from openpyxl import load_workbook
from pypdf import PdfReader
from lpdh_selftest import fixture
from lpdh_logic import compute_preview
from backend.legacy_payment_reconciliation import legacy_payment_plan, replace_legacy_payments


class PaymentPackageTests(unittest.TestCase):
    def package(self):
        docs = fixtures()
        package = docs[3]
        package['header'].update(combinedPayments=True, paymentSnapshotVersion=2)
        package['header'].update(issuerName='Penerbit Uji',recipientName='Dapur Uji',recipientAddress='Alamat Uji',senderSignatory='Petugas Uji')
        for item in package['items']:
            item['metadata']['recipientType'] = 'Relawan'
        package['items'] += docs[4]['items']
        package['total'] = sum(x['lineTotal'] for x in package['items'])
        return package

    def test_rates(self):
        self.assertEqual([school_daily_rate(x) for x in (0, 1, 99, 100, 101, 500, 501)], [20000,20000,20000,20000,30000,30000,40000])
        self.assertEqual(incentive_default({'smallPortions':90,'largePortions':5,'staffLarge':6},'Guru')['dailyAmount'],30000)
        self.assertEqual(incentive_default({'balitaSmall':64,'pregnantLarge':8,'breastfeedingLarge':24},'Kader')['dailyAmount'],96000)

    def test_one_number_separate_accounting_and_exports(self):
        doc = self.package()
        daily = merge_final_documents({}, [doc])
        self.assertEqual(len(daily['volunteerPayments']),2)
        self.assertEqual(len(daily['incentiveRecipients']),2)
        rows = daily['volunteerPayments'] + daily['incentiveRecipients']
        self.assertEqual([row['sourceLine'] for row in rows],[1,2,3,4])
        self.assertTrue(all(row['receiptNo'] == doc['documentNumber'] for row in rows))
        self.assertEqual(merge_final_documents(daily,[doc]),daily)
        date, masters, initial, plan = fixture()
        preview = compute_preview(masters, daily, date, True, plan)
        package_entries = [row for row in preview['register'] if row.get('sourceDocumentId') == doc['id']]
        self.assertEqual(len(package_entries), 1)
        self.assertEqual(package_entries[0]['amount'],doc['total'])
        self.assertEqual(package_entries[0]['proofStatus'],'UNIK')
        pdf = PdfReader(BytesIO(render_document_pdf(doc)))
        self.assertEqual(len(pdf.pages),7)
        self.assertIn('INVOICE UPAH DAN INSENTIF HARIAN',' '.join(pdf.pages[0].extract_text().split()))
        wb = load_workbook(BytesIO(render_document_excel(doc)))
        self.assertEqual(len(wb.sheetnames),7)
        self.assertEqual(wb.sheetnames[0],'Invoice Utama')
        self.assertEqual(sum(section['total'] for title,section in payment_sections(doc)), doc['total'])

    def test_package_validation(self):
        doc = self.package()
        payload = dict(site=doc['site'],document_type=doc['documentType'],service_date=doc['serviceDate'],document_number=doc['documentNumber'],header_payload=doc['header'],items=[dict(item_name=i['itemName'],category_code=i['category'],quantity=i['quantity'],unit=i['unit'],unit_price=i['unitPrice'],metadata=i['metadata']) for i in doc['items']])
        GeneratedDocumentIn(**payload)
        payload['items'][0]['metadata']['recipientType']='Tidak valid'
        with self.assertRaises(ValueError):
            GeneratedDocumentIn(**payload)

    def test_legacy_replacement_requires_fresh_confirmation(self):
        doc = self.package()
        daily = {'volunteerPayments':[{'name':'Relawan Uji A','workDays':1,'dailyRate':90000}],
                 'incentiveRecipients':[{'name':'Guru Uji','type':'Guru','amount':10000}]}
        plan = legacy_payment_plan(daily,doc)
        self.assertEqual(plan['legacyTotal'],100000)
        replaced = replace_legacy_payments(daily,doc,plan['snapshotHash'],'MAJA')
        self.assertEqual(len(merge_final_documents(replaced,[doc])['incentiveRecipients']),2)
        daily['incentiveRecipients'][0]['amount']=20000
        with self.assertRaises(ValueError):
            replace_legacy_payments(daily,doc,plan['snapshotHash'],'MAJA')


if __name__ == '__main__':
    unittest.main()
