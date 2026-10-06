"""Isolated API contract tests; no production database or session is used."""
import base64
import importlib.util
import json
import sys
import types
import unittest
from contextlib import contextmanager
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
package = types.ModuleType("backend")
package.__path__ = [str(ROOT / "backend")]


def session_role(token):
    if token not in {"Bearer MAJA", "Bearer CEMPLANG", "Bearer OWNER"}:
        raise HTTPException(401, "unauthorized")
    return token.split()[1]


auth = types.ModuleType("backend.auth_api")
auth.session_role = session_role
db = types.ModuleType("backend.db")
db.connection = lambda: None
db.database_ready = lambda: True
drive_module = types.ModuleType("backend.accountant_drive")
drive_module.upload_accountant_artifact = lambda **kwargs: {"driveUri": "https://drive.google.com/file/d/test-archive-" + kwargs.get('artifact_key','') + "/view", 'folderId':'test-folder'}
# Run this test file in its own process: importing backend/__init__ would load
# unrelated services, so install only the package and authentication/DB fixtures.
sys.modules.update({"backend": package, "backend.auth_api": auth, "backend.db": db, "backend.accountant_drive": drive_module})
spec = importlib.util.spec_from_file_location("backend.accountant_generated_document_api", ROOT / "backend/accountant_generated_document_api.py")
api = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = api
spec.loader.exec_module(api)
pdf_spec = importlib.util.spec_from_file_location("backend.generated_document_pdf", ROOT / "backend/generated_document_pdf.py")
pdf_module = importlib.util.module_from_spec(pdf_spec)
sys.modules[pdf_spec.name] = pdf_module
pdf_spec.loader.exec_module(pdf_module)


class FakeConnection:
    def __init__(self):
        self.committed = False
        self.calls = []
        self.row = {"id": 1, "site": "MAJA", "document_type": "OPERASIONAL", "document_number": "INV-OPS-MAJA-20261005-001", "service_date": date(2026, 10, 5), "status": "DRAFT", "header_payload": {"issuerName": "Penerbit", "recipientName": "Dapur", "recipientAddress": "Alamat", "senderSignatory": "Pengirim"}, "total_amount": 1000}
        self.items = [{"item_name": "Gas", "category_code": "Gas", "quantity": 1, "unit": "tabung", "unit_price": 1000, "line_total": 1000, "item_payload": {}}]
        self.result = None
        self.numbers = {}
        self.profiles = {}
        self.assets = {}
        self.serials = []
        self.anchors = {}

    @contextmanager
    def cursor(self):
        yield self

    def execute(self, sql, params=()):
        self.calls.append((sql, params))
        if sql.startswith("select full_number from document_number_anchors"):
            self.result = self.anchors.get(tuple(params))
        elif sql.startswith("insert into document_number_anchors"):
            self.anchors[tuple(params[:2])] = {'full_number':params[2]}
        elif sql.startswith("select owner_key,full_number from document_number_serials"):
            self.result = next((x for x in self.serials if x["site"] == params[0] and x["namespace"] == params[1] and (x["normalized_number"] == params[2] or x["serial"] is not None and x["serial"] == params[3])), None)
        elif sql.startswith("select full_number from document_number_serials"):
            rows = [x for x in self.serials if x["site"] == params[0] and x["namespace"] == params[1] and x["serial"] is not None]
            self.result = max(rows, key=lambda x: x["serial"]) if rows else None
        elif sql.startswith("insert into document_number_serials"):
            self.serials.append(dict(zip(("site", "namespace", "serial", "full_number", "normalized_number", "owner_key"), params)))
        elif sql.startswith("update document_number_serials"):
            for x in self.serials:
                if x["site"] == params[1] and x["namespace"] == params[2] and x["owner_key"] == params[3] and x["serial"] == params[4]: x["full_number"] = params[0]
        elif sql.startswith("select * from generated_accountant_documents where id"):
            self.result = deepcopy(self.row) if params[0] == 1 else None
        elif sql.startswith("select * from generated_accountant_documents where request_key"):
            self.result = deepcopy(self.row) if self.row.get("request_key") == params[0] else None
        elif sql.startswith("select item_name"):
            self.result = deepcopy(self.items)
        elif sql.startswith("select data from lpdh_site_state"):
            self.result = {"data": {}}
        elif "from calculator_master_catalog" in sql:
            self.result = []
        elif sql.startswith("select count(*)"):
            self.result = {"n": 1}
        elif sql.startswith("select document_id from generated_accountant_document_numbers"):
            self.result = {"document_id": self.numbers[params[0]]} if params[0] in self.numbers else None
        elif sql.startswith("delete from generated_accountant_document_numbers"):
            self.numbers = {key: owner for key, owner in self.numbers.items() if owner != params[0]}
        elif sql.startswith("insert into generated_accountant_document_numbers"):
            self.numbers[params[0]] = params[1]
        elif sql.startswith("insert into generated_document_profiles"):
            self.profiles[(params[0], params[1])] = json.loads(params[2])
        elif sql.startswith("select profile_key,header_payload"):
            self.result = [{"profile_key": key[1], "header_payload": header} for key, header in self.profiles.items() if key[0] == params[0]]
        elif sql.startswith("insert into generated_document_assets"):
            asset_id = len(self.assets) + 1
            self.assets[asset_id] = dict(zip(("site","asset_kind","filename","mime_type","content","created_by"), params))
            self.result = {"id": asset_id}
        elif "from generated_document_assets where id" in sql:
            self.result = deepcopy(self.assets.get(params[0]))
            if self.result and len(params) > 1 and (self.result['site'] != params[1] or self.result['asset_kind'] != params[2]):
                self.result = None
        elif sql.startswith("update generated_accountant_documents set document_number"):
            self.row.update(document_number=params[0], header_payload=json.loads(params[1]), total_amount=params[2])
            self.result = deepcopy(self.row)
        elif sql.startswith("insert into generated_accountant_documents"):
            self.row.update(site=params[0], document_type=params[1], document_number=params[2], service_date=params[3], header_payload=json.loads(params[4]), total_amount=params[5], request_key=params[6], request_hash=params[7])
            self.items = []
            self.result = deepcopy(self.row)
        elif sql.startswith("delete from generated_accountant_document_items"):
            self.items = []
        elif "insert into generated_accountant_document_items" in sql:
            self.items.append(dict(zip(("document_id", "item_name", "category_code", "quantity", "unit", "unit_price", "line_total", "item_payload"), params)))
            self.items[-1]["item_payload"] = json.loads(params[-1])
        elif sql.startswith("update generated_accountant_documents set status"):
            self.row["status"] = "CANCELLED" if "'CANCELLED'" in sql else "FINAL"
        elif sql.startswith("update generated_accountant_documents set drive_uri"):
            self.row["drive_uri"] = params[0]
            self.row.update(drive_excel_uri=params[1],drive_folder_id=params[2],drive_upload_status=params[3],drive_upload_error=params[4])
        elif sql.startswith("update generated_accountant_documents set drive_upload_status"):
            self.row["drive_upload_status"] = "FAILED"
            self.row["drive_upload_error"] = params[0]

    def fetchone(self):
        return self.result

    def fetchall(self):
        return self.result

    def commit(self):
        self.committed = True


class DocumentApiTests(unittest.TestCase):
    def setUp(self):
        app = FastAPI()
        app.include_router(api.router, prefix="/v1")
        self.client = TestClient(app)
        self.conn = FakeConnection()
        @contextmanager
        def connection():
            original = deepcopy((self.conn.row, self.conn.items, self.conn.numbers, self.conn.profiles, self.conn.assets, self.conn.serials))
            self.conn.committed = False
            try:
                yield self.conn
            finally:
                if not self.conn.committed:
                    self.conn.row, self.conn.items, self.conn.numbers, self.conn.profiles, self.conn.assets, self.conn.serials = original
        self.patcher = patch.object(api, "connection", connection)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        self.headers = {"Authorization": "Bearer MAJA"}
        self.payload = {"site": "MAJA", "document_type": "OPERASIONAL", "document_number": "INV/MAJA/001", "service_date": "2026-10-05", "header_payload": self.conn.row["header_payload"], "request_key": "test-request-001", "items": [{"item_name": "Gas", "category_code": "Gas", "quantity": 1, "unit": "tabung", "unit_price": 1000}]}

    def test_authentication_and_site_scope(self):
        self.assertEqual(self.client.post("/v1/accountant-documents", json=self.payload).status_code, 401)
        for action in ("finalize", "pdf"):
            response = self.client.request("PATCH" if action == "finalize" else "GET", f"/v1/accountant-documents/1/{action}", headers={"Authorization": "Bearer CEMPLANG"})
            self.assertEqual(response.status_code, 403)
        self.assertFalse(self.conn.committed)

    def test_reject_weekly_wages_and_invalid_category(self):
        for change in ({"document_type": "UPAH_RELAWAN", "quantity": 5, "unit": "hari"}, {"category_code": "Upah relawan"}, {"unit_price": 1.001}):
            payload = deepcopy(self.payload)
            if "document_type" in change:
                payload["document_type"] = change.pop("document_type")
            payload["items"][0].update(change)
            self.assertEqual(self.client.post("/v1/accountant-documents", json=payload, headers=self.headers).status_code, 422)
        self.assertEqual(self.conn.calls, [])

    def test_create_is_idempotent_and_conflicting_retry_rejected(self):
        first = self.client.post("/v1/accountant-documents", json=self.payload, headers=self.headers)
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(first.json()["document"]["documentNumber"], "INV/MAJA/001")
        second = self.client.post("/v1/accountant-documents", json=self.payload, headers=self.headers)
        self.assertEqual(first.json()["document"], second.json()["document"])
        self.payload["items"][0]["unit_price"] = 2000
        self.assertEqual(self.client.post("/v1/accountant-documents", json=self.payload, headers=self.headers).status_code, 409)

    def test_manual_number_required_and_duplicate_reserved_history_rejected(self):
        missing = deepcopy(self.payload); missing.pop('document_number')
        self.assertEqual(self.client.post('/v1/accountant-documents', json=missing, headers=self.headers).status_code, 422)
        for number in ('=SUM(A1)', '../folder', '  ', '@formula'):
            payload = {**self.payload, 'document_number': number}
            self.assertEqual(self.client.post('/v1/accountant-documents', json=payload, headers=self.headers).status_code, 422)
        self.conn.numbers['inv/maja/001'] = 42
        payload = {**self.payload, 'document_number': ' inv/MAJA/001 '}
        response = self.client.post('/v1/accountant-documents', json=payload, headers=self.headers)
        self.assertEqual(response.status_code, 409, response.text)
        self.assertFalse(self.conn.committed)

    def test_draft_manual_number_can_change_but_final_number_is_locked(self):
        response = self.client.put('/v1/accountant-documents/1', json=self.payload, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['document']['documentNumber'], 'INV/MAJA/001')
        self.assertEqual(self.conn.numbers['inv/maja/001'], 1)
        self.conn.row['status'] = 'FINAL'
        self.payload['document_number'] = 'INV/MAJA/002'
        self.assertEqual(self.client.put('/v1/accountant-documents/1', json=self.payload, headers=self.headers).status_code, 409)
        self.assertEqual(self.conn.row['document_number'], 'INV/MAJA/001')

    def test_receipts_use_manual_individual_numbers_and_reject_duplicates(self):
        # Grandfather existing drafts, but no new legacy-shaped packages.
        self.payload['document_type'] = 'UPAH_RELAWAN'
        self.conn.row['document_type'] = 'UPAH_RELAWAN'
        self.payload['items'] = [{'item_name': 'Relawan Uji A', 'quantity': 1, 'unit': 'hari', 'unit_price': 1000, 'metadata': {'receiptNo':'KWT/101'}}, {'item_name': 'Relawan Uji B', 'quantity': 1, 'unit': 'hari', 'unit_price': 1000, 'metadata': {'receiptNo':'kwt/101'}}]
        self.assertEqual(self.client.post('/v1/accountant-documents', json=self.payload, headers=self.headers).status_code, 422)
        self.assertEqual(self.client.put('/v1/accountant-documents/1', json=self.payload, headers=self.headers).status_code, 409)
        self.assertFalse(self.conn.committed)
        self.payload['items'][1]['metadata']['receiptNo'] = 'KWT/102'
        response = self.client.put('/v1/accountant-documents/1', json=self.payload, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        doc = response.json()['document']
        self.assertEqual([api.receipt_number(doc, i) for i in range(2)], ['KWT/101', 'KWT/102'])

    def test_new_aggregate_common_number_and_server_subtype_validation(self):
        self.payload['document_type'] = 'INSENTIF_GURU_KADER'
        self.payload['header_payload'].update(paymentSnapshotVersion=2, recipientSubtype='Guru')
        self.payload['items'] = [{'item_name': name, 'quantity': 1, 'unit': 'hari', 'unit_price': 1000, 'metadata': {'recipientType':'Guru'}} for name in ('Guru A', 'Guru B')]
        self.payload['items'][1]['metadata']['recipientType'] = 'Kader'
        self.assertEqual(self.client.post('/v1/accountant-documents', json=self.payload, headers=self.headers).status_code, 422)
        self.payload['items'][1]['metadata']['recipientType'] = 'Guru'
        response = self.client.post('/v1/accountant-documents', json=self.payload, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        doc = response.json()['document']
        self.assertEqual([api.receipt_number(doc, i) for i in range(2)], [self.payload['document_number']] * 2)
        self.assertEqual(len(self.conn.numbers), 1)

    def test_daily_api_rejects_unsourced_financial_injection(self):
        from backend import lpdh_api as daily_api
        request = types.SimpleNamespace(state=types.SimpleNamespace(sppg_role="MAJA"))
        existing = {"data": {}, "status": "DRAFT"}
        with patch.object(daily_api, "connection", api.connection), patch.object(daily_api, "_load_daily", return_value=existing), patch.object(daily_api, "_load_master", return_value={"data": {}}), patch.object(api, "load_documents", return_value=[]), patch.object(daily_api, "_save_daily_locked", return_value={"saved": True}) as save:
            for key in ("rawMaterials", "operations", "volunteerPayments", "incentiveRecipients"):
                payload = daily_api.DailyStateIn(site="MAJA", service_date="2026-10-05", data={key: [{"name": "Injected", "qty": 1, "price": 10}]})
                with self.assertRaises(HTTPException) as rejected:
                    daily_api.save_daily(payload, request)
                self.assertEqual(rejected.exception.status_code, 409)
            save.assert_not_called()
            # A fabricated source id is discarded by the canonical FINAL document rebuild.
            payload = daily_api.DailyStateIn(site="MAJA", service_date="2026-10-05", data={"operations": [{"sourceDocumentId": 999, "price": 999}]})
            daily_api.save_daily(payload, request)
            self.assertEqual(payload.data["operations"], [])
            forged = daily_api.DailyStateIn(site="MAJA", service_date="2026-10-05", data={"_historicalGeneratedSnapshot": True, "pm": {"rows": [{"code": "KS-02", "targetPm": 999}]}})
            daily_api.save_daily(forged, request)
            self.assertNotIn("_historicalGeneratedSnapshot", forged.data)
            self.assertEqual(forged.data["_dailyWorkflowVersion"], 2)
            self.assertEqual(forged.data["pm"]["rows"][1]["targetPm"], 0)

    def test_private_artwork_upload_validation_and_site_isolation(self):
        stream = BytesIO(); Image.new('RGB', (40, 20), 'white').save(stream, format='PNG')
        payload = {'site': 'MAJA', 'asset_kind':'SIGNATURE', 'filename':'ttd-uji.png', 'content_base64':base64.b64encode(stream.getvalue()).decode()}
        self.assertEqual(self.client.post('/v1/accountant-documents/assets', json=payload).status_code, 401)
        self.assertEqual(self.client.post('/v1/accountant-documents/assets', json={**payload,'content_base64':'not base64'}, headers=self.headers).status_code, 422)
        response = self.client.post('/v1/accountant-documents/assets', json=payload, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        asset_id = response.json()['id']
        endpoint = f'/v1/accountant-documents/assets/{asset_id}'
        self.assertEqual(self.client.get(endpoint, headers={'Authorization':'Bearer CEMPLANG'}).status_code, 403)
        image = self.client.get(endpoint, headers=self.headers)
        self.assertEqual(image.headers['cache-control'], 'private, no-store')
        self.assertEqual(base64.b64decode(image.json()['contentBase64']), stream.getvalue(), 'original evidence bytes preserved')
        wrong_site = {'site':'CEMPLANG','header_payload':{'signatureAssetId':asset_id}}
        self.assertEqual(self.client.put('/v1/accountant-documents/profile', json=wrong_site, headers={'Authorization':'Bearer CEMPLANG'}).status_code, 422)
        wrong_kind = {'site':'MAJA','header_payload':{'stampAssetId':asset_id}}
        self.assertEqual(self.client.put('/v1/accountant-documents/profile', json=wrong_kind, headers=self.headers).status_code, 422)

    def test_saved_defaults_persist_and_never_copy_transaction_values(self):
        header = {**self.payload['header_payload'], 'documentProfileKey':'YAYASAN','evidenceLink':'https://example.test/evidence','paymentReference':'payment-123'}
        response = self.client.put('/v1/accountant-documents/profile', json={'site':'MAJA','header_payload':header}, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        result = self.client.get('/v1/accountant-documents/master?site=MAJA', headers=self.headers)
        saved = result.json()['profiles']['YAYASAN']
        self.assertEqual(saved['issuerName'], header['issuerName'])
        self.assertNotIn('evidenceLink', saved)
        self.assertNotIn('paymentReference', saved)
        self.assertFalse(any('lpdh_daily_state' in sql or 'update generated_accountant_documents' in sql for sql,_ in self.conn.calls))

    def test_status_read_is_scoped_and_cancellation_cannot_be_finalized(self):
        self.conn.row['status'] = 'CANCELLED'
        self.conn.row['cancellation_reason'] = 'Salah tanggal'
        response = self.client.get('/v1/accountant-documents/1', headers=self.headers)
        self.assertEqual(response.json()['document']['status'], 'CANCELLED')
        self.assertEqual(response.json()['document']['cancellationReason'], 'Salah tanggal')
        self.assertEqual(self.client.get('/v1/accountant-documents/1', headers={'Authorization':'Bearer CEMPLANG'}).status_code, 403)
        self.assertEqual(self.client.patch('/v1/accountant-documents/1/finalize', headers=self.headers).status_code, 409)
        self.assertFalse(self.conn.committed)

    def test_finalization_import_failure_rolls_back(self):
        with patch.object(api, "_sync_daily", side_effect=HTTPException(409, "duplicate daily payment")):
            response = self.client.patch("/v1/accountant-documents/1/finalize", headers=self.headers)
        self.assertEqual(response.status_code, 409)
        self.assertFalse(self.conn.committed)
        self.assertEqual(self.conn.row["status"], "DRAFT")

    def test_finalization_commits_document_and_daily_together(self):
        with patch.object(api, "_sync_daily", return_value={"data": {"operations": []}, "imported": 1}) as sync:
            response = self.client.patch("/v1/accountant-documents/1/finalize", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["syncedToDaily"])
        self.assertTrue(self.conn.committed)
        sync.assert_called_once_with(self.conn, "MAJA", date(2026, 10, 5), "MAJA")

    def test_final_is_immutable(self):
        self.conn.row["status"] = "FINAL"
        self.assertEqual(self.client.put("/v1/accountant-documents/1", json=self.payload, headers=self.headers).status_code, 409)
        self.assertFalse(self.conn.committed)

    def test_cancel_requires_reason_and_site_access(self):
        endpoint = "/v1/accountant-documents/1/cancel"
        self.assertEqual(self.client.patch(endpoint, headers=self.headers, json={"reason": "   "}).status_code, 422)
        self.assertEqual(self.client.patch(endpoint, headers={"Authorization": "Bearer CEMPLANG"}, json={"reason": "Koreksi"}).status_code, 403)
        self.assertFalse(self.conn.committed)

    def test_cancel_final_commits_rebuild_and_preserves_document(self):
        self.conn.row["status"] = "FINAL"
        with patch.object(api, "_sync_daily", return_value={}) as sync:
            response = self.client.patch("/v1/accountant-documents/1/cancel", headers=self.headers, json={"reason": "Harga keliru"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.conn.row["status"], "CANCELLED")
        sync.assert_called_once()
        self.assertFalse(any(sql.startswith("delete") for sql, _ in self.conn.calls))
        self.assertEqual(self.client.patch("/v1/accountant-documents/1/finalize", headers=self.headers).status_code, 409)

    def test_cancel_generated_daily_rolls_back(self):
        self.conn.row["status"] = "FINAL"
        with patch.object(api, "_sync_daily", side_effect=HTTPException(409, "LPDH sudah digenerate")):
            response = self.client.patch("/v1/accountant-documents/1/cancel", headers=self.headers, json={"reason": "Koreksi"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.conn.row["status"], "FINAL")
        self.assertFalse(self.conn.committed)

    def test_archive_only_final_and_retry_is_idempotent(self):
        endpoint = "/v1/accountant-documents/1/archive"
        self.assertEqual(self.client.post(endpoint, headers=self.headers).status_code, 409)
        self.conn.row["status"] = "FINAL"
        with patch.object(drive_module, "upload_accountant_artifact", return_value={"driveUri": "https://drive.google.com/file/d/test/view", 'folderId':'test-folder'}) as upload:
            first = self.client.post(endpoint, headers=self.headers)
            second = self.client.post(endpoint, headers=self.headers)
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(first.json()["driveUri"], second.json()["driveUri"])
        self.assertEqual(upload.call_count,2)
        self.assertEqual(upload.call_args.kwargs["kind"], "invoice")
        self.assertEqual(upload.call_args.kwargs["site"], "MAJA")
        self.assertTrue(upload.call_args_list[0].kwargs["data"].startswith(b"%PDF-"))
        self.assertTrue(upload.call_args_list[1].kwargs["data"].startswith(b"PK"))
        self.assertEqual(upload.call_args_list[1].kwargs['target_folder_id'],'test-folder')
        self.assertEqual(first.json()['driveExcelUri'],second.json()['driveExcelUri'])

    def test_excel_is_final_only_same_snapshot_and_site_scoped(self):
        endpoint='/v1/accountant-documents/1/excel'
        self.assertEqual(self.client.get(endpoint,headers=self.headers).status_code,409)
        self.assertEqual(self.client.get(endpoint,headers={'Authorization':'Bearer CEMPLANG'}).status_code,403)
        self.conn.row['status']='FINAL'
        response=self.client.get(endpoint,headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(base64.b64decode(response.json()['contentBase64'])))
        self.assertEqual(wb.sheetnames,['Invoice'])
        self.assertIn(1000,[c.value for row in wb['Invoice'] for c in row])
        self.assertFalse(self.conn.committed,'download is read-only')

    def test_partial_archive_retry_does_not_reupload_pdf(self):
        self.conn.row['status']='FINAL'
        endpoint='/v1/accountant-documents/1/archive'
        with patch.object(drive_module,'upload_accountant_artifact',side_effect=[{'driveUri':'https://drive.google.com/file/d/pdf/view','folderId':'date-folder'},RuntimeError('Excel failed')]) as upload:
            first=self.client.post(endpoint,headers=self.headers)
        self.assertEqual(first.json()['driveUploadStatus'],'PARTIAL')
        self.assertEqual(self.conn.row['drive_uri'],'https://drive.google.com/file/d/pdf/view')
        self.assertIsNone(self.conn.row['drive_excel_uri'])
        with patch.object(drive_module,'upload_accountant_artifact',return_value={'driveUri':'https://drive.google.com/file/d/excel/view','folderId':'date-folder'}) as upload:
            retry=self.client.post(endpoint,headers=self.headers)
        self.assertEqual(retry.json()['driveUploadStatus'],'UPLOADED')
        upload.assert_called_once()
        self.assertEqual(upload.call_args.kwargs['target_folder_id'],'date-folder')
        self.assertTrue(upload.call_args.kwargs['filename'].endswith('.xlsx'))

    def test_drive_failure_preserves_committed_final_and_allows_retry(self):
        with patch.object(api, "_sync_daily", return_value={"data": {}, "imported": 1}), patch.object(drive_module, "upload_accountant_artifact", side_effect=RuntimeError("Drive unavailable")):
            response = self.client.patch("/v1/accountant-documents/1/finalize", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["driveUploadStatus"], "FAILED")
        self.assertEqual(self.conn.row["status"], "FINAL")
        self.assertTrue(self.conn.committed)
        retry = self.client.post("/v1/accountant-documents/1/archive", headers=self.headers)
        self.assertEqual(retry.json()["driveUploadStatus"], "UPLOADED")

    def test_calendar_scope_month_and_cancelled_totals(self):
        endpoint = "/v1/accountant-documents/calendar?site=MAJA&month=2026-10"
        self.assertEqual(self.client.get(endpoint, headers={"Authorization": "Bearer CEMPLANG"}).status_code, 403)
        self.assertEqual(self.client.get(endpoint.replace("2026-10", "2026-13"), headers=self.headers).status_code, 422)
        self.conn.result = [{"service_date": date(2026, 10, 5), "status": status, "count": 1, "total": amount} for status, amount in (("FINAL", 1000), ("CANCELLED", 9000), ("DRAFT", 3000))]
        response = self.client.get(endpoint, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["items"], [{"serviceDate": "2026-10-05", "draft": 1, "final": 1, "cancelled": 1, "finalTotal": 1000}])

    def test_authenticated_download_returns_real_pdf(self):
        with patch.dict(sys.modules, {"backend.generated_document_pdf": pdf_module}):
            response = self.client.get("/v1/accountant-documents/1/pdf", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(base64.b64decode(response.json()["contentBase64"]).startswith(b"%PDF-"))

    def test_generated_daily_rejects_changed_final_documents(self):
        self.conn.result = {"data": {}, "status": "GENERATED"}
        doc = {"id": 1, "site": "MAJA", "documentType": "OPERASIONAL", "documentNumber": "INV-1", "serviceDate": "2026-10-05", "status": "FINAL", "header": {}, "items": [{"itemName": "Gas", "category": "Gas", "quantity": 1, "unit": "tabung", "unitPrice": 1000, "lineTotal": 1000, "metadata": {}}]}
        with patch.object(api, "load_documents", return_value=[doc]):
            with self.assertRaises(HTTPException) as exc:
                api._sync_daily(self.conn, "MAJA", date(2026, 10, 5), "MAJA")
        self.assertEqual(exc.exception.status_code, 409)
        self.assertTrue(any("pg_advisory_xact_lock" in sql for sql, _ in self.conn.calls))


if __name__ == "__main__":
    unittest.main()
