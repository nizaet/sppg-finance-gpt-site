"""Synthetic-only approval/default/Drive failure contracts."""
import base64
from contextlib import contextmanager, ExitStack
from copy import deepcopy
from datetime import date
from io import BytesIO
import unittest
from unittest.mock import patch
from types import SimpleNamespace

import test_generated_document_api as fixture
from backend import lpdh_api as api
from backend.lpdh_approval import incentive_defaults, attach_approval, cancel_approval, approval_hash, print_copy, render_approval
from openpyxl import Workbook, load_workbook
from PIL import Image
from fastapi import HTTPException


class ApprovalTests(unittest.TestCase):
    def test_locked_cancelled_approval_uses_frozen_financial_snapshot(self):
        state = {'status':'GENERATED', 'data':{'_historicalGeneratedSnapshot':True,
            '_approval':{'status':'CANCELLED','revision':1}, 'balance':{'openingRaw':123},
            'incentive':{'paidAmount':456}, 'rawMaterials':[{'amount':789}]}}
        original = deepcopy(state)
        with ExitStack() as stack:
            for name, value in {'_load_master':lambda *a:{'data':{'_officialTemplateBase64':'eA=='}},
                '_load_daily':lambda *a:state, '_daily_with_hpe':lambda *a:({}, {'effective':True}),
                '_load_final_plan':lambda *a:{}, 'compute_preview':lambda *a:{},
                'prepare_template':lambda *a:b'x', 'fill_template':lambda *a:b'x'}.items():
                stack.enter_context(patch.object(api,name,value))
            merge = stack.enter_context(patch('backend.generated_document_logic.merge_final_documents'))
            defaults = stack.enter_context(patch.object(api,'_with_incentive_defaults'))
            data = api._approval_inputs(None,'CEMPLANG',date(2026,10,7))[1]
            self.assertEqual(data, original['data']); self.assertEqual(state, original)
            merge.assert_not_called(); defaults.assert_not_called()
            state['data']['_approval'] = {}
            with self.assertRaises(HTTPException): api._approval_inputs(None,'CEMPLANG',date(2026,10,7))

    def test_cancelled_revision_gets_new_archive_hash_and_preserves_old(self):
        first = attach_approval({'incentive':{'paidAmount':100}}, {'driveUri':'https://drive.example/one'}, 'old', 'OWNER', 'now')
        cancelled = cancel_approval(first,'OWNER','later','Koreksi')
        self.assertNotEqual(approval_hash({},first,'2026-10-07',{}),approval_hash({},cancelled,'2026-10-07',{}))
        digest = approval_hash({},cancelled,'2026-10-07',{})
        revised = attach_approval(cancelled,{'driveUri':'https://drive.example/two'},digest,'OWNER','next')
        self.assertEqual(revised['_approval']['revision'],2)
        self.assertEqual(approval_hash({},revised,'2026-10-07',{}),digest)
        self.assertEqual(revised['_approvalHistory'][0]['pdfLink'],'https://drive.example/one')

    def test_print_conclusion_matches_review_without_changing_source(self):
        wb = Workbook(); wb.active.title = 'J_Pengesahan'
        wb.active['E23'] = 'PERLU PERBAIKAN'
        wb.create_sheet('G_CekPPK')['C33'] = '=IF(C32=0,"LENGKAP: DAPAT DIPROSES","PERLU PERBAIKAN")'
        source = BytesIO(); wb.save(source); content = source.getvalue()
        for ready, expected in [(True, 'LENGKAP: DAPAT DIPROSES'), (False, 'PERLU PERBAIKAN')]:
            result = load_workbook(BytesIO(print_copy(content, {}, {'ready': ready})))
            self.assertEqual(result['G_CekPPK']['C33'].value, expected)
            self.assertEqual(result['J_Pengesahan']['E23'].value, expected)
        original = load_workbook(BytesIO(content))
        self.assertEqual(original['J_Pengesahan']['E23'].value, 'PERLU PERBAIKAN')
        self.assertEqual(original['G_CekPPK']['C33'].data_type, 'f')

    def test_cancel_keeps_archive_and_manual_proof_and_can_refinalize(self):
        final = attach_approval({'incentive': {}}, {'driveUri': 'https://drive.example/a'}, 'hash', 'OWNER', 'now')
        cancelled = cancel_approval(final, 'OWNER', 'later', 'Koreksi TTD')
        self.assertEqual(final['_approval']['status'], 'FINAL')
        self.assertEqual(cancelled['_approval']['status'], 'CANCELLED')
        self.assertEqual(cancelled['_approvalHistory'][0]['pdfLink'], 'https://drive.example/a')
        self.assertEqual(cancelled['incentive']['evidenceLink'], '')
        self.assertEqual(cancelled['incentive']['approvalEvidenceLink'], '')
        final['incentive']['evidenceLink'] = 'https://manual.example/proof'
        self.assertEqual(cancel_approval(final, 'OWNER', 'later', 'Koreksi')['incentive']['evidenceLink'], 'https://manual.example/proof')
        again = attach_approval(cancelled, {'driveUri': 'https://drive.example/b'}, 'hash', 'OWNER', 'again')
        self.assertEqual(again['_approval']['status'], 'FINAL')
        with self.assertRaises(ValueError): cancel_approval(cancelled, 'OWNER', 'later', 'again')

    def test_cancel_endpoint_lock_stale_snapshot_and_retry(self):
        request = SimpleNamespace(state=SimpleNamespace(sppg_role='OWNER'))
        original = attach_approval({'incentive': {}}, {'driveUri': 'https://drive.example/a'}, 'a'*64, 'OWNER', 'now')
        state = {'status': 'DRAFT', 'data': original}
        writes = []; commits = []; locks = []
        @contextmanager
        def connection(): yield SimpleNamespace(cursor=lambda: cursor(), commit=lambda: commits.append(True))
        @contextmanager
        def cursor(): yield SimpleNamespace(execute=lambda *args: writes.append(args))
        payload = api.ApprovalCancelIn(site='MAJA', service_date=date(2026,10,5), expected_hash='a'*64, reason='Koreksi')
        with patch.object(api, 'connection', connection), patch.object(api, '_load_daily', return_value=state), patch.object(fixture.api, 'daily_lock', lambda *a: locks.append(a)):
            self.assertEqual(api.approval_cancel(payload, request)['approval']['status'], 'CANCELLED')
            self.assertEqual(len(commits), 1); self.assertEqual(len(locks), 1)
            payload.expected_hash = 'b'*64
            with self.assertRaises(HTTPException) as err: api.approval_cancel(payload, request)
            self.assertEqual(err.exception.status_code, 409)
            payload.expected_hash = 'a'*64; state['status'] = 'GENERATED'
            self.assertEqual(api.approval_cancel(payload, request)['approval']['status'], 'CANCELLED')
            state['status'] = 'DRAFT'; state['data'] = cancel_approval(original, 'OWNER', 'later', 'Koreksi')
            self.assertTrue(api.approval_cancel(payload, request)['alreadyCancelled'])
            self.assertEqual(len(commits), 2)
        with self.assertRaises(HTTPException) as err:
            api.approval_cancel(payload, SimpleNamespace(state=SimpleNamespace(sppg_role='CEMPLANG')))
        self.assertEqual(err.exception.status_code, 403)

    def test_defaults_update_automatic_values_and_preserve_manual_zero(self):
        result = incentive_defaults({}, 5572000, '2026-10-05', 'MAJA', '001/KW/2026')
        self.assertEqual(result['incentive']['paymentDate'], '2026-10-05')
        self.assertEqual(result['incentive']['receiptSigned'], 'Ya')
        self.assertEqual(result['incentive']['paidAmount'], 5572000)
        self.assertEqual(incentive_defaults(result, 6000000, '2026-10-05', 'MAJA')['incentive']['paidAmount'], 6000000)
        manual = {'incentive': {'paidAmount': 0, 'statementAmount': 100, 'paymentDate': '2026-10-06'}}
        changed = incentive_defaults(manual, 5572000, '2026-10-05', 'MAJA')
        self.assertEqual(changed['incentive']['paidAmount'], 0)
        self.assertEqual(changed['incentive']['statementAmount'], 100)
        self.assertEqual(changed['incentive']['paymentDate'], '2026-10-06')
        historic = {'_historicalGeneratedSnapshot': True, 'incentive': {}}
        self.assertEqual(incentive_defaults(historic, 100, '2026-10-05', 'MAJA'), historic)

    def test_links_manual_preservation_and_revision_history(self):
        archive = {'driveUri': 'https://drive.google.com/file/d/approval/view'}
        original = {'incentive': {}}
        final = attach_approval(original, archive, 'hash', 'OWNER', 'now')
        self.assertEqual(original, {'incentive': {}})
        self.assertEqual(final['incentive']['evidenceLink'], archive['driveUri'])
        self.assertEqual(approval_hash({}, original, '2026-10-05', {}), approval_hash({}, final, '2026-10-05', {}))
        final['incentive']['evidenceLink'] = 'https://manual.example/paid.pdf'
        next_version = attach_approval(final, {'driveUri': 'https://drive.google.com/file/d/new/view'}, 'hash2', 'OWNER', 'later')
        self.assertEqual(next_version['incentive']['evidenceLink'], 'https://manual.example/paid.pdf')
        self.assertEqual(len(next_version['_approvalHistory']), 1)
        with self.assertRaises(ValueError):
            attach_approval(original, {'driveUri': ''}, 'hash', 'OWNER', 'now')

    def test_print_copy_preserves_formulas_layout_and_only_j_visible(self):
        wb = Workbook(); wb.active.title = 'Identitas'; wb.active['B6'] = 'TEST'
        ws = wb.create_sheet('J_Pengesahan'); ws.merge_cells('A1:F1'); ws['A1'] = 'LEMBAR PENGESAHAN'
        ws['D5'] = '=Identitas!B6'; ws.column_dimensions['B'].width = 30
        source = BytesIO(); wb.save(source); raw = source.getvalue()
        picture = BytesIO(); Image.new('RGBA', (120, 40), (0, 0, 0, 255)).save(picture, 'PNG')
        asset = 'data:image/png;base64,' + base64.b64encode(picture.getvalue()).decode()
        printed = load_workbook(BytesIO(print_copy(raw, {'approvalSppgSignature': asset})))
        self.assertEqual(printed['J_Pengesahan']['D5'].value, '=Identitas!B6')
        self.assertEqual(printed['J_Pengesahan'].column_dimensions['B'].width, 30)
        self.assertEqual(printed['Identitas'].sheet_state, 'hidden')
        self.assertEqual(printed['J_Pengesahan'].sheet_state, 'visible')
        self.assertEqual(len(printed['J_Pengesahan']._images), 1)
        self.assertEqual(load_workbook(BytesIO(raw))['Identitas'].sheet_state, 'visible')
        wb['Identitas']['A1'] = '=_xlfn.TEXTJOIN(" | ",TRUE(),Identitas!B6:B7)'
        valid = BytesIO(); wb.save(valid)
        self.assertTrue(print_copy(valid.getvalue(), {}))
        wb['Identitas']['A1'] = "=cmd|'external'!A0"
        dde = BytesIO(); wb.save(dde)
        with self.assertRaises(ValueError): print_copy(dde.getvalue(), {})
        wb['Identitas']['A1'] = '=WEBSERVICE("https://example.com")'
        bad = BytesIO(); wb.save(bad)
        with self.assertRaises(ValueError): print_copy(bad.getvalue(), {})

    def test_missing_converter_fails_closed(self):
        with patch('backend.lpdh_approval.shutil.which', return_value=None):
            with self.assertRaises(ValueError): render_approval(b'', {})

    def test_all_artwork_centered_separate_and_padding_removed(self):
        from backend.lpdh_approval import ASSETS
        from PIL import ImageDraw
        wb = Workbook(); ws = wb.active; ws.title = 'J_Pengesahan'
        for col in ('B','D','F'): ws.column_dimensions[col].width = 30
        for row in range(29,33): ws.row_dimensions[row].height = 18
        ws['D31'] = '(cap SPPG)'; ws['F31'] = '(cap Yayasan)'; ws['D33'] = 'Nama Kepala'
        art = Image.new('RGB', (300,100), 'white'); ImageDraw.Draw(art).rectangle((100,30,199,69), fill='black')
        picture = BytesIO(); art.save(picture, 'PNG')
        asset = 'data:image/png;base64,' + base64.b64encode(picture.getvalue()).decode()
        source = BytesIO(); wb.save(source)
        result = load_workbook(BytesIO(print_copy(source.getvalue(), {key:asset for key in ASSETS})))['J_Pengesahan']
        self.assertEqual(len(result._images), 5)
        self.assertIsNone(result['D31'].value); self.assertIsNone(result['F31'].value)
        self.assertEqual(result['D33'].value, 'Nama Kepala')
        from openpyxl.utils.units import EMU_to_pixels
        for image in result._images:
            self.assertEqual((image.width,image.height), (100,40))
            anchor = image.anchor
            x = EMU_to_pixels(anchor._from.colOff); y = EMU_to_pixels(anchor._from.rowOff)
            w = EMU_to_pixels(anchor.ext.cx); h = EMU_to_pixels(anchor.ext.cy)
            self.assertGreaterEqual(x,0); self.assertLessEqual(x+w,215)
            self.assertGreaterEqual(y,0); self.assertLessEqual(y+h,96)
        # Stamp and signature share the same signing zone, drawn stamp first.
        self.assertEqual([image.anchor._from.row for image in result._images], [28]*5)
        self.assertGreater(EMU_to_pixels(result._images[0].anchor.ext.cx),65)
        self.assertEqual(load_workbook(BytesIO(source.getvalue())).active['D31'].value, '(cap SPPG)')

    def test_finalize_failure_stale_and_idempotence(self):
        request = SimpleNamespace(state=SimpleNamespace(sppg_role='OWNER'))
        self.commits = 0; self.writes = []
        @contextmanager
        def connection(): yield self
        self.cursor = lambda: connection()
        self.execute = lambda *args: self.writes.append(args)
        self.commit = lambda: setattr(self, 'commits', self.commits + 1)
        daily = {'incentive': {'receiptNo': '001/KW/2026', 'proofNo': 'BYR-TEST'}}
        inputs = ({'signers': [{'name':'A'},{'name':'B'},{'name':'C'}]}, daily, {}, b'xlsx', 'a'*64, {'effective': True})
        payload = api.ApprovalIn(site='MAJA', service_date=date(2026,10,5), expected_hash='a'*64)
        with ExitStack() as stack:
            for name, value in {'connection':connection, '_approval_inputs':lambda *a:deepcopy(inputs), 'render_approval':lambda *a:b'pdf', 'claim_number':lambda *a:None}.items():
                stack.enter_context(patch.object(api, name, value))
            stack.enter_context(patch.object(fixture.api, 'daily_lock', lambda *a:None))
            with patch('backend.accountant_drive.upload_accountant_artifact', side_effect=RuntimeError('failed')):
                with self.assertRaises(HTTPException) as error: api.approval_finalize(payload, request)
                self.assertEqual(error.exception.status_code, 502)
                self.assertEqual(self.commits, 0); self.assertEqual(self.writes, [])
            payload.expected_hash = 'b'*64
            with self.assertRaises(HTTPException) as error: api.approval_finalize(payload, request)
            self.assertEqual(error.exception.status_code, 409)
            payload.expected_hash = 'a'*64
            response = api.approval_finalize(payload, request)
            self.assertEqual(self.commits, 1)
            self.assertEqual(response['approval']['status'], 'FINAL')
            daily['_approval'] = response['approval']
            response = api.approval_finalize(payload, request)
            self.assertTrue(response['alreadyFinal']); self.assertEqual(self.commits, 1)
        with self.assertRaises(HTTPException) as error:
            api.approval_finalize(payload, SimpleNamespace(state=SimpleNamespace(sppg_role='CEMPLANG')))
        self.assertEqual(error.exception.status_code, 403)


if __name__ == '__main__': unittest.main()

