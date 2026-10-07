"""Synthetic regression fixtures; no user identities or production records."""
import sys
import unittest
from pathlib import Path
from datetime import date
from io import BytesIO
from copy import deepcopy
from openpyxl import load_workbook
from fastapi import HTTPException

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
import lpdh_logic as logic
from document_numbering import next_number, claim_number, daily_number, suggest_number, invoice_fallback


class Cursor:
    def __init__(self):
        self.rows = []
        self.anchors = {}
        self.result = None
        self.first_date = date(2026, 10, 5)
    def execute(self, sql, params):
        if sql.startswith("select full_number from document_number_anchors"):
            self.result = self.anchors.get(tuple(params))
        elif sql.startswith("insert into document_number_anchors"):
            self.anchors[tuple(params[:2])] = {'full_number': params[2], 'owner_key': params[3]}
        elif sql.startswith("select document_id"):
            self.result = None
        elif sql.startswith("select min"):
            self.result = {"first_date": self.first_date}
        elif sql.startswith("select full_number"):
            if "owner_key=" in sql:
                rows = [r for r in self.rows if r['site'] == params[0] and r['namespace'] == 'LPDH' and r['owner_key'] == params[1]]
                self.result = rows[-1] if rows else None
            else:
                rows = [r for r in self.rows if r['site'] == params[0] and r['namespace'] == params[1] and r['serial'] is not None]
                self.result = rows[-1] if rows else None
        elif sql.startswith("select owner_key"):
            self.result = next((r for r in self.rows if r['site'] == params[0] and r['namespace'] == params[1] and (r['normalized_number'] == params[2] or r['serial'] is not None and r['serial'] == params[3])), None)
        elif sql.startswith("insert"):
            keys = ('site','namespace','serial','full_number','normalized_number','owner_key')
            if not any(r['site'] == params[0] and r['namespace'] == params[1] and r['normalized_number'] == params[4] for r in self.rows): self.rows.append(dict(zip(keys,params)))
        elif sql.startswith("update"):
            for r in self.rows:
                if r['site'] == params[1] and r['namespace'] == params[2] and r['owner_key'] == params[3] and r['serial'] == params[4]: r['full_number'] = params[0]
    def fetchone(self): return self.result


class MasterAndNumberTests(unittest.TestCase):
    def test_staff_teacher_header_alias(self):
        from openpyxl import Workbook
        for header in ('TENAGA PENDIDIK/GURU', 'Tenaga Pendidik / Guru'):
            wb = Workbook()
            ws = wb.active
            ws.title = 'Master_Sekolah'
            ws.append(['Kode Unit', 'Nama Sekolah', 'Jenis Sekolah', header])
            ws.append(['TEST-1', 'Sekolah Uji', 'SD/MI', 12])
            ws.append(['TEST-2', 'Sekolah Nol', 'SD/MI', 0])
            content = BytesIO()
            wb.save(content)
            rows = logic.parse_master_workbook(content.getvalue())['schools']
            self.assertEqual([row['staffLarge'] for row in rows], [12, 0])

    def test_school_columns_roundtrip(self):
        wb = load_workbook(BytesIO(logic.make_master_template()))
        self.assertEqual(wb['Master_Sekolah']['F1'].value, 'Porsi Kecil')
        self.assertEqual(wb['Master_Sekolah']['G1'].value, 'Porsi Besar')
        wb['Master_Sekolah'].append(['S1','Sekolah Uji','SD/MI','','',100,70,'','Aktif',''])
        wb['Master_Sekolah'].append(['S2','Santri Uji','Santri','','',0,60,'','Aktif',''])
        wb['Master_Sekolah'].append(['S3','PTK Uji','PTK','','',0,10,'','Aktif',''])
        wb['Master_Posyandu'].append(['P1','Posyandu Uji','','','',20,5,7,'Aktif',''])
        out = BytesIO(); wb.save(out)
        data = logic.parse_master_workbook(out.getvalue())
        logic.validate_master_portions(data)
        targets = logic.master_target_by_group(data)
        self.assertEqual({k:targets[k] for k in ('KS-02','KS-03','KS-06','PTK','KS-09','KS-07','KS-08')}, {'KS-02':100,'KS-03':70,'KS-06':60,'PTK':10,'KS-09':20,'KS-07':5,'KS-08':7})
        self.assertEqual(len(data['volunteers']), 0)

    def test_portion_clause_and_zero_fallback(self):
        with self.assertRaises(ValueError): logic.validate_master_portions({'schools':[{'schoolType':'Santri','smallPortions':1}]})
        data = {'schools':[{'schoolType':'SD/MI','smallPortions':0,'largePortions':0}], 'groupTargets':{'KS-02':100,'KS-03':70,'PTK':10}}
        self.assertEqual(logic.master_target_by_group(data)['KS-02'],0)
        self.assertEqual(logic.master_target_by_group(data)['PTK'],0)

    def test_user_split_template_legacy_headers(self):
        from openpyxl import Workbook
        wb=Workbook(); school=wb.active; school.title='Master_Sekolah'
        school.append(['Kode Unit','Jenis Unit','Nama Sekolah / Posyandu','Kode Kelompok','Kelompok Sasaran','Porsi Kecil','Porsi Besar','TENAGA PENDIDIK','Jenis PIC','Target PM','Nama PIC','No. HP PIC','Alamat','Status Aktif','Catatan'])
        school.append(['S1','Sekolah','SD Uji','KS-02','SD/MI',100,70,12,'Sekolah',182,'Guru Uji','','','Aktif',''])
        pos=wb.create_sheet('Master Posyandu')
        pos.append(['Kode Unit','Jenis Unit','Nama Sekolah / Posyandu','Kode Kelompok','Ibu Hamil','Ibu Menyusui','Anak Balita (6–59 bulan)','Jenis PIC','Target PM','Nama PIC','No. HP PIC','Alamat','Status Aktif','Catatan'])
        pos.append(['P1','Posyandu','Pos Uji','KS-07',5,7,20,'3B',32,'Kader Uji','','','Aktif',''])
        out=BytesIO(); wb.save(out)
        data=logic.parse_master_workbook(out.getvalue()); logic.validate_master_portions(data)
        counts=logic.master_target_by_group(data)
        self.assertEqual([counts[k] for k in ('KS-02','KS-03','PTK','KS-07','KS-08','KS-09')],[100,70,12,5,7,20])

    def test_merge_preserves_edits_and_is_repeatable(self):
        parsed = {'identity':{'sppgId':'SOURCE'}, 'volunteers':[{'code':'R1','name':'Person A','dailyRate':100,'sourceRow':'C1_Relawan!6'}], 'groupTargets':{'KS-02':100}}
        original = {'identity':{'sppgId':'EXISTING'}, 'parameters':{'operationalPerPm':123}}
        merged, added, warnings = logic.merge_master_import(original,parsed,'source.xlsx')
        self.assertEqual(original, {'identity':{'sppgId':'EXISTING'}, 'parameters':{'operationalPerPm':123}})
        self.assertEqual(merged['identity']['sppgId'],'EXISTING')
        self.assertEqual(added['volunteers'],1)
        self.assertTrue(warnings)
        merged['volunteers'][0].update(name='Edited name',dailyRate=250)
        again, added, _ = logic.merge_master_import(merged,parsed,'source.xlsx')
        self.assertEqual(again['volunteers'][0]['dailyRate'],250)
        self.assertEqual(added['volunteers'],0)
        data = {'schools':[{'code':'S1','name':'Edited school','largePortions':99}]}
        parsed_school = {'schools':[{'code':'S1','name':'Source school','largePortions':50,'staffLarge':12}]}
        merged,_,_=logic.merge_master_import(data,parsed_school,'source.xlsx')
        self.assertEqual(merged['schools'][0]['largePortions'],99)
        self.assertEqual(merged['schools'][0]['staffLarge'],12)
        merged['schools'][0]['staffLarge']=0
        again,_,_=logic.merge_master_import(merged,parsed_school,'source.xlsx')
        self.assertEqual(again['schools'][0]['staffLarge'],0)

    def test_official_import_excludes_transactions_and_summary(self):
        wb = logic.fallback_lpdh_workbook()
        wb['Identitas']['B7'] = 'SPPG MAJA UJI'
        wb['C1_Relawan']['C6'] = 'Person A'
        wb['C1_Relawan']['G6'] = 100
        wb['C1_Relawan']['C7'] = 'Person B'
        wb['C1_Relawan']['G7'] = 0
        wb['C1_Relawan']['C66'] = 1
        wb['C1_Relawan']['J6'] = 'DO-NOT-IMPORT'
        wb['C_Operasional']['D13'] = 'Gas 50 kg'
        wb['C_Operasional']['E13'] = 1
        wb['C_Operasional']['F13'] = 'tabung'
        wb['C_Operasional']['G13'] = 1000
        out=BytesIO(); wb.save(out)
        data=logic.parse_master_workbook(out.getvalue())
        self.assertEqual(len(data['volunteers']),2)
        self.assertEqual(data['volunteers'][1]['dailyRate'],0)
        self.assertEqual(data['operations'][0]['category'],'Gas')
        self.assertNotIn('receiptNo',data['volunteers'][0])
        self.assertNotIn('balance',data)
        self.assertNotIn('rawMaterials',data)
        self.assertTrue(all(s['signed']=='Tidak' for s in data['signers']))

    def test_serials_manual_override_history_and_scope(self):
        cur=Cursor(); fallback='001/OP/TEST/X/2026'
        self.assertEqual(suggest_number(cur,'MAJA','OPERASIONAL',fallback),fallback)
        self.assertFalse(cur.rows, 'reads do not reserve')
        claim_number(cur,'MAJA','OPERASIONAL',fallback,'DOC:1')
        self.assertEqual(suggest_number(cur,'MAJA','OPERASIONAL',fallback),'002/OP/TEST/X/2026')
        with self.assertRaises(HTTPException): claim_number(cur,'MAJA','OPERASIONAL','001/CHANGED/X/2026','DOC:2')
        claim_number(cur,'MAJA','OPERASIONAL','010/OP/TEST/X/2026','DOC:1')
        with self.assertRaises(HTTPException): claim_number(cur,'MAJA','OPERASIONAL',fallback,'DOC:2')
        self.assertEqual(suggest_number(cur,'MAJA','OPERASIONAL',fallback),'011/OP/TEST/X/2026')
        self.assertEqual(suggest_number(cur,'CEMPLANG','OPERASIONAL',fallback),fallback)
        self.assertEqual(suggest_number(cur,'MAJA','BAHAN_BAKU',fallback),fallback)
        claim_number(cur,'MAJA','LPDH','MANUAL','DAY:2026-10-05')
        with self.assertRaises(HTTPException): claim_number(cur,'MAJA','LPDH','manual','DAY:2026-10-06')

    def test_daily_number_next_date_preserves_number_and_workbook(self):
        cur=Cursor(); masters={'identity':{'lpdhNumber':'001/LPDH/TEST/X/2026'}}
        self.assertEqual(daily_number(cur,'MAJA',date(2026,10,5),masters),'001/LPDH/TEST/X/2026')
        self.assertEqual(daily_number(cur,'MAJA',date(2026,10,6),masters),'002/LPDH/TEST/X/2026')
        claim_number(cur,'MAJA','LPDH','008/LPDH/TEST/X/2026','DAY:2026-10-06')
        self.assertEqual(daily_number(cur,'MAJA',date(2026,10,7),masters),'009/LPDH/TEST/X/2026')
        self.assertEqual(daily_number(cur,'MAJA',date(2026,10,6),masters),'008/LPDH/TEST/X/2026')
        self.assertEqual(next_number(['0099/LPDH/TEST/X/2026'],'001/X'),'0100/LPDH/TEST/X/2026')
        self.assertEqual(invoice_fallback({'vendor':{'rawInvoicePrefix':'/BB/TEST/X/2026'}},'MAJA','BAHAN_BAKU',date(2026,10,5)),'001/BB/TEST/X/2026')
        daily={'lpdhNumber':'002/LPDH/TEST/X/2026'}
        preview=logic.compute_preview(masters,daily,'2026-10-06',False,None)
        output=logic.populate_workbook(masters,daily,preview,'2026-10-06')
        self.assertEqual(load_workbook(BytesIO(output))['Identitas']['B5'].value,daily['lpdhNumber'])

    def test_latest_manual_anchor_not_highest_and_skip_claimed_serials(self):
        cur=Cursor(); fallback='001/BB/DEFAULT/X/2026'
        claim_number(cur,'MAJA','BAHAN_BAKU','999/BB/OLD/X/2026','DOC:1')
        claim_number(cur,'MAJA','BAHAN_BAKU','220/BB/MMD/IX/2026','DOC:2')
        self.assertEqual(suggest_number(cur,'MAJA','BAHAN_BAKU',fallback),'221/BB/MMD/X/2026')
        claim_number(cur,'MAJA','BAHAN_BAKU','221/BB/MMD/IX/2026','DOC:3')
        claim_number(cur,'MAJA','BAHAN_BAKU','219/BB/MANUAL/IX/2026','DOC:4')
        self.assertEqual(suggest_number(cur,'MAJA','BAHAN_BAKU',fallback),'222/BB/MANUAL/X/2026')
        claim_number(cur,'MAJA','BAHAN_BAKU','999/BB/OLD/X/2026','DOC:1')
        self.assertEqual(suggest_number(cur,'MAJA','BAHAN_BAKU',fallback),'222/BB/MANUAL/X/2026','unchanged historical saves do not reset anchor')
        claim_number(cur,'MAJA','BAHAN_BAKU','219/BB/EDITED/X/2026','DOC:4')
        self.assertEqual(suggest_number(cur,'MAJA','BAHAN_BAKU',fallback),'222/BB/EDITED/X/2026')


if __name__=='__main__': unittest.main()
