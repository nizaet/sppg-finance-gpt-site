"""Synthetic fixtures only; no Drive account or production ledger is touched."""
import ast
import base64
import importlib.util
import io
import re
import sys
import types
import unittest
from contextlib import contextmanager
from copy import deepcopy
from datetime import date
from pathlib import Path
from unittest.mock import Mock, patch

sys.path.insert(0,str(Path(__file__).parent))
import test_generated_document_api as fixture
from backend.document_routine_copy import routine_documents, routine_daily
from backend.generated_document_excel import render_document_excel
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]


def doc(kind='OPERASIONAL', subtype=None):
    return {'id':1,'site':'MAJA','documentType':kind,'documentNumber':'230/OP/TEST/X/2026','serviceDate':'2026-10-05','status':'FINAL',
      'header':{'issuerName':'Penerbit Uji','issuerAddress':'Alamat penerbit uji','recipientName':'Dapur Uji','recipientAddress':'Alamat uji',
                'senderSignatory':'Pengirim Uji','accountNumber':'0012345','recipientSubtype':subtype,'paymentSnapshotVersion':2,
                'evidenceLink':'https://example.test/paid','paymentReference':'OLD-TRANSFER','recipientSignatureAssetId':99},
      'total':1250.5,'items':[{'itemName':'Item uji','category':'Gas','quantity':1,'unit':'tabung','unitPrice':1250.5,'lineTotal':1250.5,
          'metadata':{'role':'Pengolah','recipientType':subtype,'unitName':'Unit Uji','receiptNo':'OLD','evidenceLink':'https://example.test/proof','paymentReference':'OLD','sourceDocumentId':23,'paymentId':'OLD'}}]}


class RoutineTests(unittest.TestCase):
    def test_upgrade_anchors_use_number_claim_not_archive_timestamp_and_keep_operator_edits(self):
        migration=(ROOT/'schema/document_number_anchor_seed_v044.sql').read_text(encoding='utf-8')
        self.assertIn('c.created_at desc',migration)
        self.assertNotIn('d.updated_at',migration)
        self.assertIn('document_number_anchors.selected_at <=',migration)
        self.assertIn("migration_name='schema/document_routine_archives_v043.sql'",migration)
        self.assertIn('btrim(d.document_number)',migration,'only current document numbers, not abandoned old claims')

    def test_multi_operational_documents_retained_not_collapsed(self):
        first=doc(); second=doc(); second['documentNumber']='231/OP/TEST/X/2026'
        raw=doc('BAHAN_BAKU'); wages=doc('UPAH_RELAWAN'); canceled=doc(); canceled['status']='CANCELLED'
        source=[first,second,raw,wages,canceled]; before=deepcopy(source)
        result=routine_documents(source)
        self.assertEqual(len(result),2)
        self.assertEqual([x['sourceNumber'] for x in result],[first['documentNumber'],second['documentNumber']])
        self.assertNotIn('documentNumber',result[0])
        self.assertNotIn('evidenceLink',result[0]['header'])
        self.assertNotIn('recipientSignatureAssetId',result[0]['header'])
        self.assertEqual(set(result[0]['items'][0]['metadata']),{'role','recipientType','unitName'})
        self.assertEqual(source,before)

    def test_old_mixed_package_becomes_separate_daily_guru_kader(self):
        source=doc('INSENTIF_GURU_KADER','Guru')
        other=deepcopy(source['items'][0]); other['metadata']['recipientType']='Kader'; source['items'].append(other)
        results=routine_documents([source])
        self.assertEqual([x['header']['recipientSubtype'] for x in results],['Guru','Kader'])
        self.assertTrue(all(x['items'][0]['quantity']==1 and x['items'][0]['unit']=='hari' for x in results))
        self.assertTrue(all(x['header']['paymentSnapshotVersion']==2 for x in results))

    def test_daily_only_allowlisted_pm_no_financial_identity_or_bast(self):
        source={'lpdhNumber':'OLD','rawMaterials':[{'amount':999}],'incentive':{'paid':999},'balance':{'closing':999},'signers':[{'signed':'Ya'}],
          'pm':{'rows':[{'code':'KS-09','targetPm':99,'distributed':0,'received':0,'bnba':'Tidak','bastNo':'OLD','bastLink':'https://example.test','paymentId':'old'}],
            'production':{'produced':5,'organoleptic':3,'retainedSample':2,'bukti':'old'}}}
        result=routine_daily(source)
        self.assertEqual(set(result),{'pm'})
        self.assertEqual(result['pm']['rows'],[{'code':'KS-09','distributed':0,'received':0,'bnba':'Tidak'}])
        self.assertEqual(result['pm']['production'],{'produced':5,'organoleptic':3,'retainedSample':2})

    def test_previous_api_is_read_only_site_scoped_and_no_future(self):
        api=fixture.api
        with self.assertRaises(fixture.HTTPException) as error:
            api.previous_routine('MAJA',date(2026,10,6),date(2026,10,6),'Bearer MAJA')
        self.assertEqual(error.exception.status_code,422)
        with self.assertRaises(fixture.HTTPException) as error:
            api.previous_routine('MAJA',date(2026,10,6),None,'Bearer CEMPLANG')
        self.assertEqual(error.exception.status_code,403)
        class Cur(fixture.FakeConnection):
            def execute(self,sql,params=()):
                self.calls.append((sql,params))
                if sql.startswith('select max(d.service_date)'): self.result={'source_date':date(2026,10,5)}
                elif sql.startswith('select data from lpdh_daily_state'): self.result={'data':{'pm':{'rows':[{'code':'KS-09','received':0,'bastNo':'OLD'}]}}}
                elif sql.startswith('select data,status,revision'): self.result={'data':{},'status':'DRAFT','revision':2}
                elif sql.startswith('select data from lpdh_site_state'): self.result={'data':{}}
        cur=Cur()
        @contextmanager
        def conn(): yield cur
        with patch.object(api,'connection',conn),patch.object(api,'load_documents',return_value=[doc()]),patch('backend.document_numbering.daily_number',return_value='002/LPDH/TEST/X/2026'):
            result=api.previous_routine('MAJA',date(2026,10,6),None,'Bearer MAJA')
        self.assertEqual(result['sourceDate'],date(2026,10,5))
        self.assertEqual(result['targetDailyRevision'],2)
        self.assertFalse(cur.committed)
        self.assertTrue(all(sql.startswith('select ') for sql,_ in cur.calls))
        self.assertNotIn('bastNo',result['dailyDefaults']['pm']['rows'][0])

    def test_copy_save_refuses_generated_or_changed_daily_under_lock(self):
        from backend import lpdh_api
        from starlette.requests import Request
        request=Request({'type':'http'}); request.state.sppg_role='MAJA'
        cur=fixture.FakeConnection()
        @contextmanager
        def conn(): yield cur
        payload=lpdh_api.DailyStateIn(site='MAJA',service_date=date(2026,10,6),data={},require_editable=True,expected_revision=2)
        for stored in ({'data':{},'status':'GENERATED','revision':2}, {'data':{},'status':'DRAFT','revision':3}, {'data':{'_historicalGeneratedSnapshot':True},'status':'DRAFT','revision':2}):
            with patch.object(lpdh_api,'connection',conn),patch.object(lpdh_api,'_load_daily',return_value=stored):
                with self.assertRaises(fixture.HTTPException) as error: lpdh_api.save_daily(payload,request)
            self.assertEqual(error.exception.status_code,409)
        self.assertFalse(cur.committed)
        self.assertTrue(all(sql.startswith('select ') for sql,_ in cur.calls))


class ExcelTests(unittest.TestCase):
    def test_invoice_snapshot_numbers_identifiers_and_literal_text(self):
        source=doc(); source['items'][0]['itemName']='=HYPERLINK("evil")'
        source['header']['issuerName']='=NOT_A_FORMULA'
        before=deepcopy(source)
        wb=load_workbook(io.BytesIO(render_document_excel(source)))
        ws=wb['Invoice']; cells=[c for row in ws for c in row]
        self.assertTrue(any(c.value=='0012345' or c.value=='Nomor rekening: 0012345' for c in cells))
        self.assertTrue(any(c.value==date(2026,10,5) or str(c.value).startswith('2026-10-05') for c in cells))
        self.assertEqual([c.data_type for c in cells if isinstance(c.value,str) and c.value.startswith('=')],['s','s'])
        self.assertEqual(sum(c.value==1250.5 for c in cells),3)
        self.assertEqual(ws.page_setup.fitToWidth,1)
        self.assertEqual(source,before)

    def test_aggregate_receipt_cover_and_complete_appendix_no_fabricated_signatures(self):
        source=doc('UPAH_RELAWAN'); source['items'][0].update(unit='hari',quantity=1)
        source['items']*=46; source['total']=57523.0
        wb=load_workbook(io.BytesIO(render_document_excel(source)))
        self.assertEqual(wb.sheetnames,['Kuitansi','Daftar Penerima'])
        details=[row for row in wb['Daftar Penerima'].iter_rows(values_only=True) if row[1]=='Item uji']
        self.assertEqual(len(details),46)
        self.assertEqual(sum(row[5] for row in details),source['total'])
        self.assertTrue(all(row[6]=='________________' for row in details))
        self.assertEqual(wb['Daftar Penerima'].print_title_rows,'$11:$11')

    def test_mismatched_snapshot_total_rejected(self):
        source=doc(); source['total']=999
        with self.assertRaises(ValueError): render_document_excel(source)

    def test_sender_artwork_embedded_privately_with_partial_overlap(self):
        from PIL import Image
        image=Image.new('RGBA',(100,100),(0,0,255,100)); content=io.BytesIO(); image.save(content,format='PNG')
        wb=load_workbook(io.BytesIO(render_document_excel(doc(),{'stampAssetId':content.getvalue(),'signatureAssetId':content.getvalue()})))
        pictures=wb['Invoice']._images
        self.assertEqual(len(pictures),2)
        signature,stamp=pictures
        self.assertEqual(signature.anchor._from.row,stamp.anchor._from.row)
        self.assertEqual(signature.anchor._from.col,stamp.anchor._from.col)
        self.assertGreater(stamp.anchor.ext.cx,signature.anchor._from.colOff)


class ArchiveTests(unittest.TestCase):
    def load_drive(self):
        module=types.ModuleType('backend.google_services')
        module.drive_auth_mode=lambda:'TEST'
        module.ensure_drive_folder=Mock(side_effect=lambda parent,name:parent+'/'+name)
        module.upload_bytes_to_drive=Mock(return_value='https://drive.google.com/file/d/synthetic/view')
        spec=importlib.util.spec_from_file_location('tested_accountant_drive',ROOT/'backend/accountant_drive.py')
        actual=importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules,{'backend.google_services':module}): spec.loader.exec_module(actual)
        return actual,module

    def test_new_pair_same_site_year_month_day_and_fixed_folder_on_retry(self):
        actual,module=self.load_drive()
        first=actual.upload_accountant_artifact(kind='invoice',site='MAJA',service_date='2026-10-05',filename='doc.pdf',data=b'%PDF',mime_type='application/pdf',artifact_key='doc-1-pdf')
        self.assertEqual(first['drivePath'],'MAJA/2026/10/2026-10-05')
        second=actual.upload_accountant_artifact(kind='invoice',site='MAJA',service_date='2026-10-05',target_folder_id=first['folderId'],filename='doc.xlsx',data=b'PK',mime_type='test/xlsx',artifact_key='doc-1-xlsx')
        self.assertEqual(first['folderId'],second['folderId'])
        self.assertEqual(module.ensure_drive_folder.call_count,4)
        self.assertEqual(module.upload_bytes_to_drive.call_args.kwargs['artifact_key'],'doc-1-xlsx')

    def test_date_folder_failure_never_silently_falls_back_to_root(self):
        actual,module=self.load_drive(); module.ensure_drive_folder.side_effect=RuntimeError('folder denied')
        with self.assertRaises(actual.AccountantDriveUploadError):
            actual.upload_accountant_artifact(kind='invoice',site='MAJA',service_date='2026-10-05',filename='doc.pdf',data=b'%PDF',mime_type='application/pdf')
        module.upload_bytes_to_drive.assert_not_called()

    def test_google_upload_key_recovers_existing_after_process_interruption(self):
        source=ast.parse((ROOT/'backend/google_services.py').read_text(encoding='utf-8'))
        functions=[x for x in source.body if isinstance(x,ast.FunctionDef) and x.name in {'upload_bytes_to_drive','drive_file_parent'}]
        service=Mock(); service.files.return_value.list.return_value.execute.return_value={'files':[{'id':'already-uploaded','webViewLink':'https://drive.google.com/file/d/already-uploaded/view'}]}
        globals_={'io':io,'re':re,'drive_service':lambda:service,'MediaIoBaseUpload':Mock(),'GoogleServicesNotConfigured':RuntimeError}
        exec(compile(ast.Module(body=functions,type_ignores=[]),'google_services.py','exec'),globals_)
        uri=globals_['upload_bytes_to_drive']('date-folder','doc.pdf',b'PDF','application/pdf',artifact_key='doc-1-pdf')
        self.assertIn('already-uploaded',uri)
        service.files.return_value.create.assert_not_called()
        self.assertIn("value='doc-1-pdf'",service.files.return_value.list.call_args.kwargs['q'])


if __name__=='__main__': unittest.main()
