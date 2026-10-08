import unittest
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
import test_generated_document_api as fixture
from backend.operational_item_recap import operational_recap, recap_excel


def doc(identity,day='2026-10-05',status='FINAL',kind='OPERASIONAL',unit='paket',qty=2):
    return {'id':identity,'serviceDate':day,'status':status,'documentType':kind,'documentNumber':f'{identity}/OP/X/2026',
        'items':[{'itemName':'APD','unit':unit,'category':'APD','quantity':qty,'unitPrice':100,'lineTotal':qty*100}]}


class RecapTests(unittest.TestCase):
    def test_final_period_units_and_duplicate_ids(self):
        documents=[doc(1),doc(1),doc(2,'2026-10-09',qty=3),doc(3,unit='pcs'),doc(4,status='CANCELLED'),
            doc(5,status='DRAFT'),doc(6,kind='UPAH_RELAWAN'),doc(7,'2026-10-10')]
        before=deepcopy(documents)
        result=operational_recap(documents,'2026-10-05','2026-10-09')
        self.assertEqual(result['invoiceCount'],3)
        self.assertEqual(result['totalAmount'],700)
        self.assertEqual(result['itemCount'],2)
        self.assertEqual(next(x for x in result['items'] if x['unit']=='paket')['quantity'],5)
        self.assertEqual(documents,before)

    def test_excel_snapshot_numbers_dates_and_safe_text(self):
        documents=[doc(1)]
        documents[0]['items'][0]['itemName']='=HYPERLINK("bad")'
        recap=operational_recap(documents,'2026-10-05','2026-10-09')
        wb=load_workbook(BytesIO(recap_excel(recap,'MAJA','2026-10-05','2026-10-09')))
        self.assertEqual(wb.sheetnames,['Rekap Item','Rincian Invoice'])
        self.assertEqual(wb['Rekap Item']['C6'].value,2)
        self.assertEqual(wb['Rekap Item']['E6'].value,200)
        self.assertEqual(wb['Rekap Item']['A6'].data_type,'s')
        self.assertEqual(wb['Rincian Invoice']['A6'].value.strftime('%Y-%m-%d'),'2026-10-05')
        self.assertEqual(wb['Rincian Invoice']['H6'].value,200)
        self.assertEqual(wb['Rekap Item'].freeze_panes,'A6')
        self.assertEqual(wb['Rekap Item']['E3'].value,200)

    def test_empty_period(self):
        self.assertEqual(operational_recap([],'2026-10-05','2026-10-09')['totalAmount'],0)


class RecapApiTests(fixture.DocumentApiTests):
    def test_item_recap_auth_and_export(self):
        base='/v1/accountant-documents/recap/operational-items'
        query='?site=MAJA&start_date=2026-10-05&end_date=2026-10-09'
        for path in (base,base+'.xlsx'):
            self.assertEqual(self.client.get(path+query).status_code,401)
            self.assertEqual(self.client.get(path+query.replace('site=MAJA','site=CEMPLANG'),headers=self.headers).status_code,403)
            self.assertEqual(self.client.get(path+query.replace('end_date=2026-10-09','end_date=2026-10-04'),headers=self.headers).status_code,422)
        with patch.object(fixture.api,'invoice_recap',return_value={'documents':[doc(1)]}):
            result=self.client.get(base+query,headers=self.headers)
            self.assertEqual(result.json()['totalAmount'],200)
            exported=self.client.get(base+'.xlsx'+query,headers=self.headers)
            self.assertEqual(exported.status_code,200)
            self.assertIn('spreadsheetml',exported.headers['content-type'])
            self.assertEqual(load_workbook(BytesIO(exported.content))['Rekap Item']['E6'].value,200)

