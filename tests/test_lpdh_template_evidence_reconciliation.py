"""Synthetic regression tests. No private workbook, live database or bank payment."""
import base64
import json
import unittest
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
import test_generated_document_api as fixture
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from backend import lpdh_api
from backend.legacy_payment_reconciliation import legacy_payment_plan, replace_legacy_payments
from backend.lpdh_logic import compute_preview, populate_workbook
from backend.generated_document_logic import merge_final_documents
from test_lpdh_fill_only_template import synthetic_template


def document():
    return {'id': 12, 'site': 'MAJA', 'serviceDate': '2026-10-05',
            'documentType': 'UPAH_RELAWAN', 'documentNumber': '012/TEST/X/2026',
            'status': 'DRAFT', 'header': {'paymentSnapshotVersion': 2}, 'total': 100,
            'items': [{'itemName': 'Synthetic A', 'quantity': 1, 'unit': 'hari',
                       'unitPrice': 100, 'lineTotal': 100, 'metadata': {}}]}


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.doc = document()
        self.daily = {'volunteerPayments': [
            {'name': 'Synthetic A', 'workDays': 1, 'dailyRate': 100},
            {'name': 'Synthetic B', 'workDays': 1, 'dailyRate': 20},
            {'name': 'Other final', 'workDays': 1, 'dailyRate': 30, 'sourceDocumentId': 99}],
            'balance': {'bank': 500}, 'incentiveRecipients': [{'name': 'Teacher', 'amount': 10}]}

    def test_confirmation_preserves_other_final_and_archives_exact_legacy(self):
        old = deepcopy(self.daily)
        plan = legacy_payment_plan(self.daily, self.doc)
        self.assertEqual((plan['legacyCount'], plan['legacyTotal'], plan['newTotal'], plan['removedCount']), (2, 120, 100, 1))
        replaced = replace_legacy_payments(self.daily, self.doc, plan['snapshotHash'], 'MAJA')
        self.assertEqual(self.daily, old, 'never mutate the input')
        self.assertEqual(replaced['volunteerPayments'], [old['volunteerPayments'][2]])
        self.assertEqual(replaced['_replacedLegacyPayments'][0]['rows'], old['volunteerPayments'][:2])
        self.assertEqual(replaced['balance'], old['balance'])
        self.assertEqual(replaced['incentiveRecipients'], old['incentiveRecipients'])

    def test_stale_legacy_or_document_hash_rejected(self):
        token = legacy_payment_plan(self.daily, self.doc)['snapshotHash']
        for daily, doc in (({**self.daily, 'volunteerPayments': self.daily['volunteerPayments'][:1]}, self.doc),
                           (self.daily, {**self.doc, 'total': 101})):
            with self.assertRaises(ValueError):
                replace_legacy_payments(daily, doc, token, 'MAJA')

    def test_no_replacement_for_unrelated_or_individual_receipt(self):
        self.doc['items'][0]['itemName'] = 'Unrelated'
        self.assertIsNone(legacy_payment_plan(self.daily, self.doc))
        self.doc = document(); self.doc['header'] = {}
        self.assertIsNone(legacy_payment_plan(self.daily, self.doc))

    def test_teacher_replacement_does_not_replace_cadres(self):
        self.doc.update(documentType='INSENTIF_GURU_KADER')
        self.doc['header']['recipientSubtype'] = 'Guru'
        self.daily['incentiveRecipients'] = [{'name': 'Synthetic A', 'type': 'Guru', 'amount': 10},
                                            {'name': 'Cadre', 'type': 'Kader', 'amount': 20}]
        plan = legacy_payment_plan(self.daily, self.doc)
        replaced = replace_legacy_payments(self.daily, self.doc, plan['snapshotHash'], 'MAJA')
        self.assertEqual(replaced['incentiveRecipients'], [self.daily['incentiveRecipients'][1]])

    def test_sync_requires_confirmation_and_blocks_generated(self):
        daily = {'volunteerPayments': self.daily['volunteerPayments'][:2]}
        draft = deepcopy(self.doc)
        final = {**draft, 'status': 'FINAL'}
        class Cursor:
            status = 'DRAFT'
            saved = None
            def execute(self, sql, params=()):
                if sql.startswith('insert into lpdh_daily_state'):
                    self.saved = json.loads(params[2])
            def fetchone(self): return {'data': daily, 'status': self.status}
        cur = Cursor()
        with patch.object(fixture.api, 'load_documents', return_value=[final]):
            with self.assertRaises(HTTPException) as caught:
                fixture.api._sync_daily(cur, 'MAJA', '2026-10-05', 'MAJA')
            self.assertEqual(caught.exception.status_code, 409)
            token = legacy_payment_plan(daily, draft)['snapshotHash']
            cur.status = 'GENERATED'
            with self.assertRaises(HTTPException):
                fixture.api._sync_daily(cur, 'MAJA', '2026-10-05', 'MAJA', (draft, token))
            self.assertIsNone(cur.saved)
            cur.status = 'DRAFT'
            result = fixture.api._sync_daily(cur, 'MAJA', '2026-10-05', 'MAJA', (draft, token))
            self.assertEqual(len(result['data']['volunteerPayments']), 1)
            self.assertEqual(result['data']['volunteerPayments'][0]['sourceDocumentId'], 12)
            self.assertEqual(len(result['data']['_replacedLegacyPayments']), 1)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.app = FastAPI()
        @self.app.middleware('http')
        async def role(request, call_next):
            request.state.sppg_role = request.headers.get('x-test-role', 'MAJA')
            return await call_next(request)
        self.app.include_router(lpdh_api.router)
        self.client = TestClient(self.app)
        self.cur = fixture.FakeConnection()
        original_execute = self.cur.execute
        def execute(sql, params=()):
            original_execute(sql, params)
            if sql.startswith('insert into lpdh_site_state'):
                self.cur.result = {'revision': 4}
        self.cur.execute = execute
        @contextmanager
        def connection(): yield self.cur
        self.patcher = patch.object(lpdh_api, 'connection', connection)
        self.patcher.start(); self.addCleanup(self.patcher.stop)

    def payload(self, names):
        wb = Workbook(); wb.remove(wb.active)
        for name in names: wb.create_sheet(name)
        output = BytesIO(); wb.save(output)
        content = synthetic_template() if 'A_PM' in names else output.getvalue()
        return {'site': 'MAJA', 'filename': 'synthetic.xlsx', 'content_base64': base64.b64encode(content).decode()}

    def test_json_body_upload_valid_and_does_not_overwrite_master(self):
        names = ['Identitas','A_PM','B_BahanBaku','C_Operasional','C1_Relawan','D_Insentif','E_Saldo','F_TopUp','G_CekPPK','H_RekapPPK','I_RegisterBukti','J_Pengesahan','Ref']
        response = self.client.put('/v1/lpdh/official-template', json=self.payload(names))
        self.assertEqual(response.status_code, 200, response.text)
        sql, params = self.cur.calls[-1]
        self.assertIn('lpdh_site_state.data || excluded.data', sql)
        self.assertEqual(set(json.loads(params[1])), {'_officialTemplateBase64', '_officialTemplateFilename',
            '_preparedTemplateBase64','_preparedTemplateVersion','_preparedTemplateSourceHash'})
        self.assertTrue(response.json()['fillOnly'])
        self.assertIn('_officialTemplateHistory',sql)

    def test_invalid_or_master_workbook_no_write_site_guard(self):
        for payload in (self.payload(['Master_Sekolah']), {'site':'MAJA','filename':'bad.xlsx','content_base64':'not base64'}):
            response = self.client.put('/v1/lpdh/official-template', json=payload)
            self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.cur.calls, [])
        self.assertEqual(self.client.put('/v1/lpdh/official-template', json=self.payload(['Master_Sekolah']), headers={'x-test-role':'CEMPLANG'}).status_code, 403)

    def test_filled_template_clears_unused_finances_preserves_formula_and_style(self):
        masters, daily = {}, {}
        preview = compute_preview(masters, daily, '2026-10-05', False)
        wb = load_workbook(BytesIO(populate_workbook(masters, daily, preview, '2026-10-05')))
        for sheet, coordinate in [('B_BahanBaku','C45'),('C1_Relawan','C65'),('E_Saldo','C23'),('A_PM','L6'),('A_PM','M6')]:
            wb[sheet][coordinate] = 'OLD PAYMENT MUST NOT CARRY'
            wb[sheet][coordinate].hyperlink = 'https://example.test/old-private-proof'
        style = wb['B_BahanBaku']['C45'].style_id
        formula = wb['B_BahanBaku']['I6'].value
        template = BytesIO(); wb.save(template)
        result = load_workbook(BytesIO(populate_workbook(masters, daily, preview, '2026-10-06', template.getvalue())))
        for sheet, coordinate in [('B_BahanBaku','C45'),('C1_Relawan','C65'),('E_Saldo','C23'),('A_PM','L6'),('A_PM','M6')]:
            self.assertIsNone(result[sheet][coordinate].value)
        self.assertEqual(result['B_BahanBaku']['C45'].style_id, style)
        self.assertEqual(result['B_BahanBaku']['I6'].value, formula)

    def test_incentive_amount_matches_official_eligibility_gate(self):
        daily = {'pm': {'rows': [], 'production': {'organoleptic': 3, 'retainedSample': 2}},
                 'incentive': {'eligibility': {'verified':'Ya','pmInputSipgn':'Ya'}}}
        self.assertEqual(compute_preview({}, daily, '2026-10-05', True)['incentiveCalculated'], 10000)
        daily['incentive']['eligibility']['contamination'] = 'Ya'
        result = compute_preview({}, daily, '2026-10-05', True)
        self.assertFalse(result['incentiveEligible'])
        self.assertEqual(result['incentiveCalculated'], 0)

    def test_common_evidence_and_reference_export_every_recipient(self):
        doc = document(); doc['status'] = 'FINAL'
        doc['items'].append({**deepcopy(doc['items'][0]), 'itemName':'Synthetic B'})
        doc['total'] = 200
        daily = merge_final_documents({}, [doc])
        for row in daily['volunteerPayments']:
            row.update(evidenceLink='https://example.test/signed.pdf',paymentReference='SYNTHETIC-REF')
        daily = merge_final_documents(daily, [doc])
        result = load_workbook(BytesIO(populate_workbook({}, daily, compute_preview({}, daily, '2026-10-05', False), '2026-10-05')))
        for r in (6,7):
            self.assertEqual(result['C1_Relawan'][f'J{r}'].value, doc['documentNumber'])
            self.assertEqual(result['C1_Relawan'][f'K{r}'].value, 'https://example.test/signed.pdf')
            self.assertIn('SYNTHETIC-REF', result['C1_Relawan'][f'J{r}'].comment.text)
        for r in (2,3):
            self.assertEqual(result['Lampiran_Dokumen'][f'I{r}'].value, 'SYNTHETIC-REF')
        daily.update(volunteerPaymentReference='OLD-REF',volunteerBatchEvidenceLink='https://example.test/old')
        for row in daily['volunteerPayments']: row.update(evidenceLink='',paymentReference='')
        cleared = compute_preview({}, merge_final_documents(daily,[doc]), '2026-10-05', False)
        self.assertTrue(all(not row['evidenceLink'] and not row['paymentReference'] for row in cleared['volunteers']))


class FinalizationApiTests(fixture.DocumentApiTests):
    # Reuse the isolated endpoint fixture, including its existing lifecycle checks.
    def test_preflight_site_scope_and_confirmed_body(self):
        self.conn.row.update(document_type='UPAH_RELAWAN', header_payload={**self.conn.row['header_payload'], 'paymentSnapshotVersion':2})
        self.conn.items = [{'item_name':'Synthetic A','category_code':'Upah','quantity':1,'unit':'hari','unit_price':100,'line_total':100,'item_payload':{}}]
        self.conn.row['total_amount'] = 100
        daily = {'volunteerPayments':[{'name':'Synthetic A','workDays':1,'dailyRate':100}]}
        original_execute = self.conn.execute
        def execute(sql, params=()):
            original_execute(sql, params)
            if sql.startswith('select data,status from lpdh_daily_state'):
                self.conn.result = {'data':daily,'status':'DRAFT'}
        self.conn.execute = execute
        url = '/v1/accountant-documents/1/finalization-check'
        self.assertEqual(self.client.get(url,headers={'Authorization':'Bearer CEMPLANG'}).status_code,403)
        response = self.client.get(url,headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        token = response.json()['legacyReplacement']['snapshotHash']
        self.assertEqual(len(token),64)
        with patch.object(fixture.api, '_sync_daily', return_value={'data':{},'imported':1}) as sync, patch.object(fixture.api,'_archive_document',return_value={'driveUploadStatus':'UPLOADED'}):
            response = self.client.patch('/v1/accountant-documents/1/finalize',headers=self.headers,json={'replace_legacy_snapshot':token})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(sync.call_args.args[4][1],token)
            self.assertEqual(sync.call_args.args[4][0]['status'],'DRAFT')
            self.assertTrue(self.conn.committed)


if __name__ == '__main__': unittest.main()
