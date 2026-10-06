"""Application runtime export of the exact saved document, not the LPDH workbook.

FINAL amounts are literal numbers from the same snapshot used by the PDF.
No payment, acceptance signature, or accounting event is inferred by this export.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO
from math import ceil

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
GREEN = '14513D'
MONEY = '"Rp "#,##0.00;[Red]("Rp "#,##0.00);"Rp "0.00'


def put(ws, row, col, value, bold=False, size=10):
    cell = ws.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = 's'  # Operator text, including = + @, never a formula.
    cell.font = Font(name='Calibri', size=size, bold=bold, color=GREEN if bold else '203B32')
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    return cell


def band(ws, row, start, end, text, size=10):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    put(ws, row, start, str(text or ''), True, size)


def artwork_image(ws, content, row, col, width, height, x=0):
    if not content:
        return
    picture = Image(BytesIO(content))
    ratio = min(width / picture.width, height / picture.height)
    picture.anchor = OneCellAnchor(_from=AnchorMarker(col=col-1, row=row-1, colOff=int(x*9525)),
                                  ext=XDRPositiveSize2D(int(picture.width*ratio*9525), int(picture.height*ratio*9525)))
    ws.add_image(picture)


def base_sheet(wb, name, document, title, artwork):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for col, width in zip('ABCDEFG', (5, 31, 15, 10, 11, 18, 20)):
        ws.column_dimensions[col].width = width
    h = document['header']
    kop = (artwork or {}).get('letterheadAssetId')
    if not kop:
        from backend.generated_document_pdf import ASSETS, asset_profile
        profile = asset_profile(document)
        if profile:
            path = ASSETS / ('maja-koperasi-0.png' if profile == 'maja-koperasi' else 'maja-yayasan-1.png')
            if path.exists():
                kop = path.read_bytes()
    if kop:
        probe = Image(BytesIO(kop))
        wide = probe.width >= probe.height * 3
        artwork_image(ws, kop, 1, 1, 715 if wide else 65, 74)
        ws.row_dimensions[1].height = 60
        first = 2
    else:
        first = 1
    band(ws, first, 1, 7, h.get('issuerName'), 16)
    band(ws, first+1, 1, 7, h.get('issuerSubtitle'))
    band(ws, first+2, 1, 7, h.get('issuerAddress'))
    ws.row_dimensions[first+2].height = 32
    band(ws, first+4, 1, 7, title, 14)
    band(ws, first+5, 1, 5, 'Nomor: ' + document['documentNumber'])
    put(ws, first+5, 6, document['status'], True)
    band(ws, first+6, 1, 5, 'Tanggal pelayanan / pembayaran')
    ws.merge_cells(start_row=first+6, start_column=6, end_row=first+6, end_column=7)
    put(ws, first+6, 6, date.fromisoformat(document['serviceDate'][:10])).number_format = 'dd-mm-yyyy'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    ws.oddFooter.right.text = 'Halaman &P dari &N'
    ws.oddFooter.left.text = document['documentNumber'].replace('&', '&&')
    return ws, first+8


def table(ws, start, headings, rows, total, amount_col=7):
    edge = Side(style='thin', color='CCDAD3')
    for col, value in enumerate(headings, 1):
        cell = put(ws, start, col, value, True)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    ws.row_dimensions[start].height = 36
    for index, values in enumerate(rows, start+1):
        lines = max(ceil(len(str(value))/width) for value, width in zip(values[:3],(5,31,15)))
        ws.row_dimensions[index].height = max(30, min(409, lines*16))
        for col, value in enumerate(values, 1):
            cell = put(ws, index, col, value)
            cell.border = Border(bottom=edge)
            if index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F1F7F4')
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = MONEY if col in {6, amount_col} else '#,##0.####' if col == 4 else '0'
    end = start+len(rows)+1
    band(ws, end, 1, amount_col-1, 'TOTAL PEMBAYARAN')
    put(ws, end, amount_col, float(total), True).number_format = MONEY
    ws.row_dimensions[end].height = 30
    ws.freeze_panes = f'D{start+1}'
    ws.print_title_rows = f'{start}:{start}'
    return end


def render_document_excel(document, artwork=None, _workbook=None, _prefix=''):
    try:
        from .payment_package import is_payment_package, payment_sections
    except ImportError:
        from payment_package import is_payment_package, payment_sections
    items = document['items']
    total = Decimal(str(document['total'])).quantize(Decimal('.01'))
    if sum((Decimal(str(item['lineTotal'])) for item in items), Decimal(0)).quantize(Decimal('.01')) != total:
        raise ValueError('Total snapshot dokumen tidak sama dengan rincian')
    wb = _workbook if _workbook is not None else Workbook()
    if _workbook is None:
        wb.remove(wb.active)
    wb.properties.title = document['documentNumber']
    wb.properties.creator = document['header'].get('issuerName') or 'SPPG'
    kind = document['documentType']
    receipt = kind in {'UPAH_RELAWAN', 'INSENTIF_GURU_KADER'}
    title = ('KUITANSI UPAH RELAWAN' if kind == 'UPAH_RELAWAN' else
             'KUITANSI INSENTIF ' + str(document['header'].get('recipientSubtype') or 'GURU / KADER').upper() if receipt else
             'INVOICE BAHAN BAKU' if kind == 'BAHAN_BAKU' else 'INVOICE OPERASIONAL')
    combined = is_payment_package(document)
    if kind == 'INSENTIF_MITRA':
        title = 'INVOICE INSENTIF MITRA / YAYASAN'
    if combined:
        title = 'INVOICE UPAH DAN INSENTIF HARIAN'
    ws, row = base_sheet(wb, 'Invoice Utama' if combined else _prefix or ('Kuitansi' if receipt else 'Invoice'), document, title, artwork)
    h = document['header']
    for label, value in [('Penerima', h.get('recipientName')), ('Perusahaan', h.get('recipientCompany')),
                         ('Alamat', h.get('recipientAddress')), ('Pengirim', h.get('senderName') or h.get('issuerName')),
                         ('Perusahaan pengirim', h.get('senderCompany'))]:
        band(ws, row, 1, 7, f'{label}: {value or ""}')
        ws.row_dimensions[row].height = 30 if label == 'Alamat' else 20
        row += 1
    row += 1
    details = ([[1, title + ' — ' + document['serviceDate'], f'{len(items)} penerima', 1, 'hari', float(total), float(total)]] if receipt else
               [[i, x['itemName'], x['category'], float(x['quantity']), x['unit'], float(x['unitPrice']), float(x['lineTotal'])] for i, x in enumerate(items, 1)])
    if combined:
        details = [[i, name, f"{len(section['items'])} penerima", 1, 'hari', section['total'], section['total']] for i, (name, section) in enumerate(payment_sections(document), 1)]
    end = table(ws, row, ['No', 'Nama barang / pembayaran', 'Kategori / penerima', 'Jumlah', 'Satuan', 'Harga satuan', 'Total harga'], details, total)
    row = end+2
    for label, value in [('Metode', h.get('paymentMethod') or 'Transfer'), ('Bank', h.get('bankName')),
                         ('Nomor rekening', h.get('accountNumber')), ('Atas nama', h.get('accountName'))]:
        band(ws, row, 1, 4, f'{label}: {value or ""}')
        ws.row_dimensions[row].height = 22
        row += 1
    signrow = end+3
    band(ws, signrow, 5, 7, 'Hormat Kami')
    ws.row_dimensions[signrow+1].height = 60
    artwork_image(ws, (artwork or {}).get('signatureAssetId'), signrow+1, 6, 118, 66, 20)
    artwork_image(ws, (artwork or {}).get('stampAssetId'), signrow+1, 6, 75, 75)
    band(ws, signrow+2, 5, 7, h.get('senderSignatory'))
    band(ws, row+1, 1, 7, 'Tanda terima penerima diisi setelah pembayaran diterima.' if receipt else 'Penerima: ' + str(h.get('recipientSignatory') or '________________'))
    if not receipt:
        ws.row_dimensions[row+2].height = 60
        artwork_image(ws, (artwork or {}).get('recipientSignatureAssetId'), row+2, 2, 120, 75)
    ws.print_area = f'A1:G{row+3}'
    if combined:
        for name, section in payment_sections(document):
            render_document_excel(section, artwork, _workbook=wb, _prefix=name)
    elif receipt:
        appendix, start = base_sheet(wb, ('Daftar ' + _prefix)[:31] if _prefix else 'Daftar Penerima', document, 'LAMPIRAN DAFTAR PENERIMA', artwork)
        band(appendix, start, 1, 7, 'Pembayaran harian 1 hari. Tanda terima dan tanggal penerimaan diisi manual setelah diterima.')
        appendix.row_dimensions[start].height = 32
        values = [[i, item['itemName'], (item.get('metadata') or {}).get('role') or (item.get('metadata') or {}).get('unitName') or '',
                   1, 'hari', float(item['lineTotal']), '________________'] for i, item in enumerate(items, 1)]
        end = table(appendix, start+2, ['No', 'Nama penerima', 'Unit / Tugas', 'Hari', 'Satuan', 'Nominal', 'Tanda terima / tanggal'], values, total, 6)
        appendix.print_area = f'A1:G{end}'
    if _workbook is not None:
        return
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
