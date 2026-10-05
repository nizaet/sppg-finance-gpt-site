from __future__ import annotations

import base64
import json
from datetime import date, datetime, timedelta
from typing import Any

from fastapi import APIRouter, HTTPException, Query, Request
from pydantic import BaseModel, Field

from backend.db import connection, database_ready
from backend.lpdh_logic import (
    compute_preview,
    encode_bytes,
    make_master_template,
    parse_master_workbook,
    populate_workbook,
    workbook_reference_rows,
    merge_master_import,
    validate_master_portions,
    master_target_by_group,
    normalize_daily_draft,
    validate_daily_financial_sources,
)
from backend.document_numbering import claim_number, daily_number

router = APIRouter(prefix="/v1/lpdh", tags=["lpdh"])


class MasterStateIn(BaseModel):
    site: str
    data: dict[str, Any] = Field(default_factory=dict)


class DailyStateIn(BaseModel):
    site: str
    service_date: date
    data: dict[str, Any] = Field(default_factory=dict)
    status: str = "DRAFT"


class EffectiveDaysIn(BaseModel):
    site: str
    month: str
    dates: list[date] = Field(default_factory=list)
    notes: dict[str, str] = Field(default_factory=dict)


class MasterImportIn(BaseModel):
    site: str
    filename: str = "master.xlsx"
    content_base64: str


class CalculatorFinalIn(BaseModel):
    site: str
    service_date: date
    source_plan_id: str | None = None
    plan_name: str | None = None
    payload: dict[str, Any] = Field(default_factory=dict)


class GenerateIn(BaseModel):
    site: str
    service_date: date


def _require_db() -> None:
    if not database_ready():
        raise HTTPException(503, "DATABASE_URL tidak tersedia")


def _role(request: Request) -> str:
    return str(getattr(request.state, "sppg_role", "") or "").upper().strip()


def _site(request: Request, requested: str) -> str:
    site = str(requested or "").upper().strip()
    if site not in {"MAJA", "CEMPLANG"}:
        raise HTTPException(400, "site harus MAJA atau CEMPLANG")
    role = _role(request)
    if role not in {"OWNER", site}:
        raise HTTPException(403, f"akun {role or '-'} tidak dapat mengakses LPDH {site}")
    return site


def _month_bounds(month: str) -> tuple[date, date]:
    try:
        year_text, month_text = month.split("-", 1)
        year = int(year_text)
        mon = int(month_text)
        start = date(year, mon, 1)
        if mon == 12:
            end = date(year + 1, 1, 1)
        else:
            end = date(year, mon + 1, 1)
        return start, end
    except Exception as exc:
        raise HTTPException(400, "month harus format YYYY-MM") from exc


def _load_master(site: str) -> dict[str, Any]:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("select data,revision,updated_by,updated_at from lpdh_site_state where site=%s", (site,))
            row = cur.fetchone()
    if not row:
        return {"data": {}, "revision": 0, "updatedBy": None, "updatedAt": None}
    return {
        "data": row["data"] or {},
        "revision": row["revision"],
        "updatedBy": row["updated_by"],
        "updatedAt": row["updated_at"],
    }


def _load_daily(site: str, service_date: date) -> dict[str, Any]:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "select data,status,revision,updated_by,updated_at from lpdh_daily_state where site=%s and service_date=%s",
                (site, service_date),
            )
            row = cur.fetchone()
    if not row:
        return {"data": {}, "status": "DRAFT", "revision": 0, "updatedBy": None, "updatedAt": None}
    return {
        "data": row["data"] or {},
        "status": row["status"],
        "revision": row["revision"],
        "updatedBy": row["updated_by"],
        "updatedAt": row["updated_at"],
    }


def _load_final_plan(site: str, service_date: date) -> dict[str, Any] | None:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """select source_plan_id,plan_name,payload,revision,finalized_by,finalized_at
                   from lpdh_final_plans where site=%s and service_date=%s""",
                (site, service_date),
            )
            row = cur.fetchone()
    if not row:
        return None
    return {
        "sourcePlanId": row["source_plan_id"],
        "planName": row["plan_name"],
        "payload": row["payload"] or {},
        "revision": row["revision"],
        "finalizedBy": row["finalized_by"],
        "finalizedAt": row["finalized_at"],
    }


def _effective_context(site: str, service_date: date) -> dict[str, Any]:
    week_start = service_date - timedelta(days=service_date.weekday())
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "select is_effective from lpdh_effective_days where site=%s and service_date=%s",
                (site, service_date),
            )
            row = cur.fetchone()
            effective = bool(row and row["is_effective"])
            if effective:
                cur.execute(
                    """select count(*) as n
                       from lpdh_effective_days
                       where site=%s
                         and is_effective=true
                         and service_date >= %s
                         and service_date <= %s""",
                    (site, week_start, service_date),
                )
                count_row = cur.fetchone()
                hpe_number = int((count_row or {}).get("n") or 0)
            else:
                hpe_number = 0
    return {"effective": effective, "hpeNumber": hpe_number}


def _effective(site: str, service_date: date) -> bool:
    return bool(_effective_context(site, service_date)["effective"])


def _daily_with_hpe(site: str, service_date: date, daily: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    context = _effective_context(site, service_date)
    normalized = dict(daily or {})
    normalized["dayStatus"] = "HPE" if context["effective"] else normalized.get("dayStatus") or "Tidak HPE"
    normalized["hpeNumber"] = context["hpeNumber"]
    return normalized, context


@router.get("/masters")
def get_masters(request: Request, site: str = Query()) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    result = _load_master(target)
    public_data = dict(result.get("data") or {})
    public_data.pop("_officialTemplateBase64", None)
    public_data["groupTargetAggregate"] = master_target_by_group(public_data)
    result = {**result, "data": public_data}
    return {"site": target, **result}


@router.put("/masters")
def save_masters(payload: MasterStateIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    actor = _role(request)
    current = _load_master(site)["data"] or {}
    incoming = dict(payload.data or {})
    try:
        validate_master_portions(incoming)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    for protected_key in ("_officialTemplateBase64", "_officialTemplateFilename"):
        if protected_key not in incoming and protected_key in current:
            incoming[protected_key] = current[protected_key]
    # The original workbook totals are read-only import history.
    incoming["groupTargets"] = current.get("groupTargets") or {}
    incoming.pop("groupTargetAggregate", None)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """insert into lpdh_site_state(site,data,revision,updated_by,updated_at)
                   values (%s,%s::jsonb,1,%s,now())
                   on conflict (site) do update
                   set data=excluded.data,
                       revision=lpdh_site_state.revision+1,
                       updated_by=excluded.updated_by,
                       updated_at=now()
                   returning revision,updated_at""",
                (site, json.dumps(incoming, ensure_ascii=False), actor),
            )
            row = cur.fetchone()
        conn.commit()
    return {"site": site, "saved": True, "revision": row["revision"], "updatedAt": row["updated_at"]}


@router.get("/daily")
def get_daily(request: Request, site: str = Query(), service_date: date = Query(alias="date")) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    daily = _load_daily(target, service_date)
    masters = _load_master(target)["data"]
    normalized = normalize_daily_draft(masters, daily.get("data") or {}, daily["status"])
    if daily["status"] != "GENERATED":
        from backend.accountant_generated_document_api import load_documents
        from backend.generated_document_logic import merge_final_documents
        with connection() as conn, conn.cursor() as cur:
            normalized = merge_final_documents(normalized, load_documents(cur, target, service_date, True))
    normalized, context = _daily_with_hpe(target, service_date, normalized)
    with connection() as conn, conn.cursor() as cur:
        normalized["lpdhNumber"] = daily_number(cur, target, service_date, _load_master(target)["data"], normalized.get("lpdhNumber"))
    return {
        "site": target,
        "serviceDate": service_date,
        **daily,
        "data": normalized,
        **context,
        "finalPlan": _load_final_plan(target, service_date),
    }


@router.put("/daily")
def save_daily(payload: DailyStateIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    requested_status = str(payload.status or "DRAFT").upper()
    if requested_status not in {"DRAFT", "READY", "GENERATED"}:
        raise HTTPException(400, "status daily tidak valid")
    if requested_status == "GENERATED":
        raise HTTPException(422, "Gunakan Generate LPDH untuk menyimpan snapshot final")
    from backend.accountant_generated_document_api import daily_lock, load_documents
    from backend.generated_document_logic import merge_final_documents
    with connection() as conn, conn.cursor() as cur:
        daily_lock(cur, site, payload.service_date)
        try:
            existing = _load_daily(site, payload.service_date)
            validate_daily_financial_sources(payload.data, existing["data"])
            if existing["status"] == "GENERATED" or existing["data"].get("_historicalGeneratedSnapshot"):
                payload.data["_historicalGeneratedSnapshot"] = True
            else:
                payload.data.pop("_historicalGeneratedSnapshot", None)
            payload.data = normalize_daily_draft(_load_master(site)["data"], payload.data, "DRAFT")
            payload.data = merge_final_documents(payload.data, load_documents(cur, site, payload.service_date, True))
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc
        result = _save_daily_locked(payload, request, site, cur)
        conn.commit()
    return result


def _save_daily_locked(payload, request, site, cur):
    masters = _load_master(site)["data"] or {}
    if not str(payload.data.get("lpdhNumber") or "").strip():
        cur.execute("select data from lpdh_daily_state where site=%s and service_date=%s", (site, payload.service_date))
        old = (cur.fetchone() or {}).get("data") or {}
        payload.data["lpdhNumber"] = daily_number(cur, site, payload.service_date, masters, old.get("lpdhNumber"))
    claim_number(cur, site, "LPDH", payload.data["lpdhNumber"], "DAY:" + payload.service_date.isoformat())
    requested_status = str(payload.status or "DRAFT").upper()
    status = requested_status
    if requested_status != "GENERATED":
        normalized, context = _daily_with_hpe(site, payload.service_date, payload.data or {})
        final_plan = _load_final_plan(site, payload.service_date)
        if final_plan and context["effective"]:
            validation = compute_preview(
                masters,
                normalized,
                payload.service_date.isoformat(),
                context["effective"],
                final_plan,
            )
            status = "READY" if validation["ready"] else "DRAFT"
        else:
            status = "DRAFT"
    actor = _role(request)
    cur.execute(
                """insert into lpdh_daily_state(site,service_date,data,status,revision,updated_by,updated_at)
                   values (%s,%s,%s::jsonb,%s,1,%s,now())
                   on conflict (site,service_date) do update
                   set data=excluded.data,status=excluded.status,
                       revision=lpdh_daily_state.revision+1,
                       updated_by=excluded.updated_by,updated_at=now()
                   returning revision,updated_at""",
                (site, payload.service_date, json.dumps(payload.data, ensure_ascii=False), status, actor),
            )
    row = cur.fetchone()
    return {"site": site, "serviceDate": payload.service_date, "saved": True, "status": status, "revision": row["revision"], "updatedAt": row["updated_at"]}


@router.delete("/daily")
def delete_daily(request: Request, site: str = Query(), service_date: date = Query(alias="date")) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("delete from lpdh_daily_state where site=%s and service_date=%s", (target, service_date))
            deleted = cur.rowcount
        conn.commit()
    return {"site": target, "serviceDate": service_date, "deleted": bool(deleted)}


@router.get("/calendar")
def calendar_month(request: Request, site: str = Query(), month: str = Query()) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    start, end = _month_bounds(month)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """select service_date,is_effective,note
                   from lpdh_effective_days
                   where site=%s and service_date >= %s and service_date < %s
                   order by service_date""",
                (target, start, end),
            )
            effective_rows = cur.fetchall()
            cur.execute(
                """select service_date,status,revision,updated_at
                   from lpdh_daily_state
                   where site=%s and service_date >= %s and service_date < %s
                   order by service_date""",
                (target, start, end),
            )
            daily_rows = cur.fetchall()
            cur.execute(
                """select service_date,revision,finalized_at
                   from lpdh_final_plans
                   where site=%s and service_date >= %s and service_date < %s
                   order by service_date""",
                (target, start, end),
            )
            plan_rows = cur.fetchall()

    effective_map = {row["service_date"].isoformat(): row for row in effective_rows}
    daily_map = {row["service_date"].isoformat(): row for row in daily_rows}
    plan_map = {row["service_date"].isoformat(): row for row in plan_rows}
    all_dates = sorted(set(effective_map) | set(daily_map) | set(plan_map))
    items = []
    for day in all_dates:
        e = effective_map.get(day)
        d = daily_map.get(day)
        p = plan_map.get(day)
        items.append({
            "serviceDate": day,
            "effective": bool(e and e["is_effective"]),
            "note": (e or {}).get("note") if e else "",
            "status": (d or {}).get("status") if d else "EMPTY",
            "dailyRevision": (d or {}).get("revision") if d else 0,
            "updatedAt": (d or {}).get("updated_at") if d else None,
            "finalized": bool(p),
            "finalPlanRevision": (p or {}).get("revision") if p else 0,
        })
    return {"site": target, "month": month, "items": items}


@router.get("/effective-days")
def get_effective_days(request: Request, site: str = Query(), month: str = Query()) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    start, end = _month_bounds(month)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """select service_date,is_effective,note,updated_by,updated_at
                   from lpdh_effective_days
                   where site=%s and service_date >= %s and service_date < %s
                   order by service_date""",
                (target, start, end),
            )
            rows = cur.fetchall()
    return {
        "site": target,
        "month": month,
        "dates": [row["service_date"] for row in rows if row["is_effective"]],
        "items": rows,
    }


@router.put("/effective-days")
def save_effective_days(payload: EffectiveDaysIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    actor = _role(request)
    start, end = _month_bounds(payload.month)
    chosen = sorted({item for item in payload.dates if start <= item < end})
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "delete from lpdh_effective_days where site=%s and service_date >= %s and service_date < %s",
                (site, start, end),
            )
            for item in chosen:
                cur.execute(
                    """insert into lpdh_effective_days(site,service_date,is_effective,note,updated_by,updated_at)
                       values (%s,%s,true,%s,%s,now())""",
                    (site, item, payload.notes.get(item.isoformat(), ""), actor),
                )
        conn.commit()
    return {"site": site, "month": payload.month, "saved": True, "count": len(chosen), "dates": chosen}


@router.post("/import-master")
def import_master(payload: MasterImportIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    try:
        content = base64.b64decode(payload.content_base64, validate=True)
        parsed = parse_master_workbook(content)
    except Exception as exc:
        raise HTTPException(400, f"File master tidak dapat dibaca: {exc}") from exc

    source_name = str((parsed.get("identity") or {}).get("sppgName") or "").upper()
    if parsed.get("source") == "OFFICIAL_LPDH" and site not in source_name:
        raise HTTPException(422, "Nama dapur workbook tidak cocok dengan dapur tujuan. Pilih dapur yang benar.")
    try:
        validate_master_portions(parsed)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc

    actor = _role(request)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("select pg_advisory_xact_lock(hashtextextended(%s,0))", ("lpdh-master:" + site,))
            cur.execute("select data from lpdh_site_state where site=%s for update", (site,))
            existing = (cur.fetchone() or {}).get("data") or {}
            merged, added, warnings = merge_master_import(existing, parsed, payload.filename)
            cur.execute(
                """insert into lpdh_site_state(site,data,revision,updated_by,updated_at)
                   values (%s,%s::jsonb,1,%s,now())
                   on conflict (site) do update
                   set data=excluded.data,revision=lpdh_site_state.revision+1,
                       updated_by=excluded.updated_by,updated_at=now()
                   returning revision""",
                (site, json.dumps(merged, ensure_ascii=False), actor),
            )
            revision = cur.fetchone()["revision"]
        conn.commit()

    return {
        "site": site,
        "filename": payload.filename,
        "imported": added,
        "groupTargets": len(parsed.get("groupTargets") or {}),
        "warnings": warnings,
        "revision": revision,
    }


@router.get("/master-template")
def master_template(request: Request, site: str = Query()) -> dict[str, Any]:
    _site(request, site)
    content = make_master_template()
    return {
        "filename": "Template_Import_Master_LPDH_Sekolah_Posyandu_v2.xlsx",
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "contentBase64": encode_bytes(content),
    }


@router.get("/reference")
def reference(request: Request, site: str = Query()) -> dict[str, Any]:
    _site(request, site)
    return {"sheet": "Ref", "rows": workbook_reference_rows()}




@router.get("/official-template")
def official_template_status(request: Request, site: str = Query()) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    masters = _load_master(target)["data"] or {}
    return {
        "site": target,
        "installed": bool(masters.get("_officialTemplateBase64")),
        "filename": masters.get("_officialTemplateFilename") or None,
    }


@router.put("/official-template")
def save_official_template(payload: OfficialTemplateIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    try:
        raw = base64.b64decode(payload.content_base64, validate=True)
        from openpyxl import load_workbook
        import io
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        required = {"Identitas","A_PM","B_BahanBaku","C_Operasional","C1_Relawan","D_Insentif","E_Saldo","F_TopUp","G_CekPPK","H_RekapPPK","I_RegisterBukti","J_Pengesahan","Ref"}
        missing = sorted(required.difference(workbook.sheetnames))
        if missing:
            raise ValueError("sheet wajib tidak ada: " + ", ".join(missing))
    except Exception as exc:
        raise HTTPException(400, f"Template LPDH tidak valid: {exc}") from exc
    current = _load_master(site)["data"] or {}
    current["_officialTemplateBase64"] = payload.content_base64
    current["_officialTemplateFilename"] = payload.filename
    actor = _role(request)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """insert into lpdh_site_state(site,data,revision,updated_by,updated_at)
                   values (%s,%s::jsonb,1,%s,now())
                   on conflict (site) do update
                   set data=excluded.data,revision=lpdh_site_state.revision+1,
                       updated_by=excluded.updated_by,updated_at=now()
                   returning revision""",
                (site, json.dumps(current, ensure_ascii=False), actor),
            )
            revision = cur.fetchone()["revision"]
        conn.commit()
    return {"site": site, "installed": True, "filename": payload.filename, "revision": revision}


@router.get("/calculator-final")
def get_calculator_final(request: Request, site: str = Query(), service_date: date = Query(alias="date")) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    plan = _load_final_plan(target, service_date)
    return {"site": target, "serviceDate": service_date, "finalized": bool(plan), "plan": plan}


@router.post("/calculator-final")
def save_calculator_final(payload: CalculatorFinalIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    actor = _role(request)
    if not payload.payload:
        raise HTTPException(400, "payload final kalkulator kosong")
    if not str(payload.payload.get("date") or payload.service_date.isoformat()).startswith(payload.service_date.isoformat()):
        raise HTTPException(400, "tanggal payload kalkulator tidak sesuai tanggal final")
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """insert into lpdh_final_plans(site,service_date,source_plan_id,plan_name,payload,revision,finalized_by,finalized_at)
                   values (%s,%s,%s,%s,%s::jsonb,1,%s,now())
                   on conflict (site,service_date) do update
                   set source_plan_id=excluded.source_plan_id,
                       plan_name=excluded.plan_name,
                       payload=excluded.payload,
                       revision=lpdh_final_plans.revision+1,
                       finalized_by=excluded.finalized_by,
                       finalized_at=now()
                   returning revision,finalized_at""",
                (
                    site, payload.service_date, payload.source_plan_id, payload.plan_name,
                    json.dumps(payload.payload, ensure_ascii=False), actor,
                ),
            )
            row = cur.fetchone()
        conn.commit()
    return {
        "site": site,
        "serviceDate": payload.service_date,
        "finalized": True,
        "revision": row["revision"],
        "finalizedAt": row["finalized_at"],
    }


@router.get("/preview")
def preview(request: Request, site: str = Query(), service_date: date = Query(alias="date")) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    masters = _load_master(target)["data"] or {}
    daily_state = _load_daily(target, service_date)
    raw_daily = normalize_daily_draft(masters, daily_state["data"] or {}, daily_state["status"])
    if daily_state["status"] != "GENERATED":
        from backend.accountant_generated_document_api import load_documents
        from backend.generated_document_logic import merge_final_documents
        with connection() as conn, conn.cursor() as cur:
            raw_daily = merge_final_documents(raw_daily, load_documents(cur, target, service_date, True))
    daily, context = _daily_with_hpe(target, service_date, raw_daily)
    final_plan = _load_final_plan(target, service_date)
    return compute_preview(masters, daily, service_date.isoformat(), context["effective"], final_plan)


@router.post("/generate")
def generate(payload: GenerateIn, request: Request) -> dict[str, Any]:
    _require_db()
    site = _site(request, payload.site)
    from backend.accountant_generated_document_api import daily_lock
    with connection() as conn, conn.cursor() as cur:
        daily_lock(cur, site, payload.service_date)
        result = _generate_locked(payload, request, site, cur)
        conn.commit()
    return result


def _generate_locked(payload, request, site, cur):
    actor = _role(request)
    masters = _load_master(site)["data"] or {}
    daily_state = _load_daily(site, payload.service_date)
    from backend.accountant_generated_document_api import load_documents
    from backend.generated_document_logic import merge_final_documents
    raw_daily = normalize_daily_draft(masters, daily_state["data"] or {}, daily_state["status"])
    if daily_state["status"] != "GENERATED":
        raw_daily = merge_final_documents(raw_daily, load_documents(cur, site, payload.service_date, True))
    daily, context = _daily_with_hpe(site, payload.service_date, raw_daily)
    daily["lpdhNumber"] = daily_number(cur, site, payload.service_date, masters, daily.get("lpdhNumber"))
    final_plan = _load_final_plan(site, payload.service_date)
    if not final_plan:
        raise HTTPException(409, {"message": "Data Kalkulator belum berstatus FINAL untuk tanggal ini"})
    if not context["effective"]:
        raise HTTPException(409, {"message": "Tanggal ini bukan Hari Pelayanan Efektif"})
    preview_data = compute_preview(masters, daily, payload.service_date.isoformat(), context["effective"], final_plan)
    if not preview_data["ready"]:
        issues = [row for row in preview_data["checks"] if not row["ok"]]
        raise HTTPException(
            409,
            {
                "message": "LPDH belum dapat digenerate karena validasi belum bersih",
                "errorCount": len(issues),
                "issues": issues,
            },
        )

    claim_number(cur, site, "LPDH", daily["lpdhNumber"], "DAY:" + payload.service_date.isoformat())

    template_bytes = None
    stored_template = masters.get("_officialTemplateBase64")
    if stored_template:
        try:
            template_bytes = base64.b64decode(stored_template, validate=True)
        except Exception:
            template_bytes = None
    content = populate_workbook(masters, daily, preview_data, payload.service_date.isoformat(), template_bytes=template_bytes)
    daily["_historicalGeneratedSnapshot"] = True
    filename = f"LPDH_{site}_{payload.service_date.isoformat()}.xlsx"
    cur.execute(
        """insert into lpdh_generation_log(
             site,service_date,filename,validation_status,validation_error_count,generated_by,payload
           ) values (%s,%s,%s,'OK',0,%s,%s::jsonb)""",
        (site, payload.service_date, filename, actor, json.dumps({"preview": {"errorCount": 0}}, ensure_ascii=False)),
    )
    cur.execute(
        """insert into lpdh_daily_state(site,service_date,data,status,revision,updated_by,updated_at)
           values (%s,%s,%s::jsonb,'GENERATED',1,%s,now())
           on conflict (site,service_date) do update
           set data=excluded.data,status='GENERATED',revision=lpdh_daily_state.revision+1,
               updated_by=excluded.updated_by,updated_at=now()""",
        (site, payload.service_date, json.dumps(daily, ensure_ascii=False), actor),
    )
    return {
        "filename": filename,
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "contentBase64": encode_bytes(content),
        "validation": {"ready": True, "errorCount": 0},
    }


@router.get("/history")
def history(request: Request, site: str = Query(), limit: int = Query(default=50, ge=1, le=200)) -> dict[str, Any]:
    _require_db()
    target = _site(request, site)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """select id,site,service_date,filename,validation_status,validation_error_count,generated_by,generated_at
                   from lpdh_generation_log where site=%s order by generated_at desc limit %s""",
                (target, limit),
            )
            rows = cur.fetchall()
    return {"site": target, "items": rows}
