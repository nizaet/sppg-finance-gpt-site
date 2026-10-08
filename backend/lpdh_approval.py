"""Print the official J sheet, preserving its source formulas and layout.

Only a disposable print copy is saved through openpyxl. Downloaded workbooks
continue to use lpdh_template's lossless, input-only package patching.
"""
import copy
import hashlib
import json
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from pypdf import PdfReader

ASSETS = {
    'approvalSppgStamp': ('D29', 90, 80),
    'approvalFoundationStamp': ('F29', 90, 80),
    'approvalFinanceSignature': ('B29', 170, 72),
    'approvalSppgSignature': ('D29', 170, 72),
    'approvalFoundationSignature': ('F29', 170, 72),
}


def incentive_defaults(daily, calculated, service_date, site, receipt=None):
    out = copy.deepcopy(daily)
    if out.get('_historicalGeneratedSnapshot'):
        return out
    inc = out.setdefault('incentive', {})
    automatic = set(inc.get('_automaticAmounts') or [])
    for field in ('statementAmount', 'paidAmount'):
        if inc.get(field) in (None, '') or field in automatic:
            inc[field] = calculated
            automatic.add(field)
    inc['_automaticAmounts'] = sorted(automatic)
    for field, value in {
        'paymentDate': service_date, 'receiptSigned': 'Ya',
        'proofNo': f'BYR-INS-{site}-{service_date.replace("-", "")}',
        'receiptNo': receipt,
    }.items():
        if inc.get(field) in (None, '') and value is not None:
            inc[field] = value
    return out


def approval_hash(masters, daily, service_date, preview):
    # The assigned Drive link and approval audit are outputs, not print inputs.
    data = copy.deepcopy(daily)
    data.pop('_approval', None)
    data.pop('_approvalHistory', None)
    inc = data.setdefault('incentive', {})
    inc.pop('evidenceLink', None)
    inc.pop('approvalEvidenceLink', None)
    material = {'date': service_date, 'daily': data,
                'approvalRevision': int((daily.get('_approval') or {}).get('revision') or 1) + (1 if (daily.get('_approval') or {}).get('status') == 'CANCELLED' else 0),
                'template': masters.get('_officialTemplateBase64'),
                'signers': masters.get('signers'), 'assets': masters.get('assets'),
                'identity': masters.get('identity'), 'parameters': masters.get('parameters'),
                'summary': {k: preview.get(k) for k in ('rawTotal', 'operationalTotal', 'incentiveCalculated', 'production', 'balance', 'topup')}}
    return hashlib.sha256(json.dumps(material, sort_keys=True, default=str).encode()).hexdigest()


def attach_approval(daily, archive, digest, actor, timestamp):
    out = copy.deepcopy(daily)
    link = str(archive.get('driveUri') or '')
    if not link.startswith('https://'):
        raise ValueError('Drive tidak mengembalikan link PDF pengesahan yang valid.')
    previous = out.get('_approval')
    if previous and (previous.get('hash') != digest or previous.get('status') != 'FINAL'):
        out.setdefault('_approvalHistory', []).append(previous)
    out['_approval'] = {'status': 'FINAL', 'hash': digest, 'pdfLink': link,
                        'revision': int((previous or {}).get('revision') or 1) + (1 if previous and previous.get('status') == 'CANCELLED' else 0),
                        'folderId': archive.get('folderId'), 'finalizedBy': actor,
                        'finalizedAt': timestamp, 'evidenceType': 'PENGESAHAN'}
    inc = out.setdefault('incentive', {})
    inc['approvalEvidenceLink'] = link
    # Preserve a separately entered real payment-evidence link.
    old_auto = (previous or {}).get('pdfLink')
    if not str(inc.get('evidenceLink') or '').strip().startswith('https://') or inc.get('evidenceLink') == old_auto:
        inc['evidenceLink'] = link
    return out


def cancel_approval(daily, actor, timestamp, reason):
    out = copy.deepcopy(daily)
    current = out.get('_approval') or {}
    if current.get('status') != 'FINAL':
        raise ValueError('Pengesahan tidak berstatus FINAL.')
    cancelled = {**current, 'status': 'CANCELLED', 'cancelledBy': actor,
                 'cancelledAt': timestamp, 'cancelReason': reason}
    out.setdefault('_approvalHistory', []).append(cancelled)
    out['_approval'] = cancelled
    inc = out.setdefault('incentive', {})
    for field in ('approvalEvidenceLink', 'evidenceLink'):
        if inc.get(field) == current.get('pdfLink'):
            inc[field] = ''
    return out


def print_copy(content, assets, validation=None):
    """Hide other sheets only in the print copy; references remain available."""
    if __package__:
        from .generated_document_settings import validate_artwork
    else:
        from generated_document_settings import validate_artwork
    wb = load_workbook(BytesIO(content), keep_links=False)
    # Never evaluate network/DDE formulas supplied in an uploaded workbook.
    import re
    for sheet in wb:
        for row in sheet:
            for cell in row:
                expression = re.sub(r'"(?:[^"]|"")*"', '""', str(cell.value))
                if cell.data_type == 'f' and re.search(r'WEBSERVICE\s*\(|RTD\s*\(|DDE\s*\(|\[[0-9]+\]|^[=+].*\|.*!', expression, re.I):
                    raise ValueError('Template cetak tidak boleh memakai rumus sumber eksternal.')
    if 'J_Pengesahan' not in wb.sheetnames:
        raise ValueError('Template tidak memiliki sheet J_Pengesahan.')
    ws = wb['J_Pengesahan']
    if validation is not None:
        # Disposable PDF copy: use the same authoritative checklist as Review.
        # Legacy workbook checks do not understand grouped invoices or dummy BAST.
        conclusion = 'LENGKAP: DAPAT DIPROSES' if validation.get('ready') is True else 'PERLU PERBAIKAN'
        if 'G_CekPPK' in wb.sheetnames:
            wb['G_CekPPK']['C33'] = conclusion
        # Official layouts may store a literal result rather than link C33.
        for row in ws:
            for cell in row:
                if str(cell.value or '').strip() in ('PERLU PERBAIKAN', 'LENGKAP: DAPAT DIPROSES'):
                    cell.value = conclusion
    for sheet in wb:
        sheet.sheet_state = 'visible' if sheet == ws else 'hidden'
    wb.active = wb.index(ws)
    ws.sheet_view.tabSelected = True
    # Preserve an explicit official print area. Legacy official templates omit it.
    if not ws.print_area:
        ws.print_area = 'A1:F34'
    for key, (anchor, width, height) in ASSETS.items():
        if not assets.get(key):
            continue
        raw, _ = validate_artwork(str(assets[key]).split(',', 1)[-1])
        # Ignore surrounding white/transparent padding when centering uploaded art.
        from PIL import Image as PillowImage, ImageChops
        art = PillowImage.open(BytesIO(raw)).convert('RGBA')
        paper = PillowImage.new('RGBA', art.size, 'white')
        paper.alpha_composite(art)
        mask = ImageChops.difference(paper.convert('RGB'), PillowImage.new('RGB', art.size, 'white')).convert('L')
        bounds = mask.point(lambda value: 255 if value > 12 else 0).getbbox()
        if not bounds:
            raise ValueError('Gambar TTD/stempel kosong. Unggah gambar yang terlihat.')
        art = art.crop(bounds)
        cropped = BytesIO(); art.save(cropped, format='PNG'); cropped.seek(0)
        image = Image(cropped)
        column, row = coordinate_from_string(anchor)
        col_index = column_index_from_string(column) - 1
        col_width = int((ws.column_dimensions[column].width or 13) * 7 + 5)
        area_height = sum(points_to_pixels(ws.row_dimensions[r].height or ws.sheet_format.defaultRowHeight or 15) for r in range(row,33))
        factor = min(width / image.width, min(height, area_height - 4) / image.height, (col_width - 12) / image.width)
        image.width *= factor
        image.height *= factor
        stamp = key.endswith('Stamp')
        center = col_width * (0.42 if stamp else 0.57)
        marker = AnchorMarker(col=col_index, row=row-1,
            colOff=pixels_to_EMU(max(0, min(col_width-image.width, center-image.width/2))),
            rowOff=pixels_to_EMU(max(0, (area_height-image.height)/2)))
        image.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height)))
        # Replace the placeholder only on the disposable print copy.
        placeholder = f'{column}31'
        if stamp and str(ws[placeholder].value or '').strip().lower().startswith('(cap '):
            ws[placeholder] = None
        ws.add_image(image)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def render_approval(content, assets, validation=None):
    binary = shutil.which('libreoffice') or shutil.which('soffice')
    if not binary:
        raise ValueError('Mesin cetak template belum tersedia. Dokumen tidak difinalkan.')
    with tempfile.TemporaryDirectory(prefix='lpdh-approval-') as folder:
        root = Path(folder)
        source = root / 'J_Pengesahan.xlsx'
        source.write_bytes(print_copy(content, assets, validation))
        # Independent profiles avoid concurrent print jobs sharing state.
        profile = (root / 'profile').as_uri()
        subprocess.run([binary, f'-env:UserInstallation={profile}', '--headless',
                        '--convert-to', 'pdf:calc_pdf_Export', '--outdir', str(root), str(source)],
                       check=True, timeout=60, capture_output=True)
        target = root / 'J_Pengesahan.pdf'
        if not target.exists():
            raise ValueError('PDF J_Pengesahan belum berhasil dibuat.')
        data = target.read_bytes()
        reader = PdfReader(BytesIO(data))
        text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        if not reader.pages or 'PENGESAHAN' not in text.upper() or '#REF!' in text:
            raise ValueError('Hasil cetak template tidak valid. Periksa template resmi.')
        return data

