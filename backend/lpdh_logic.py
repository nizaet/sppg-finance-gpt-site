from __future__ import annotations

import base64
import io
import re
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

ROOT = Path(__file__).resolve().parents[1]
LPDH_TEMPLATE = ROOT / "backend" / "templates" / "LPDH_SPPG_Format_Excel_PPK.xlsx"

GROUPS = [
    ("KS-01", "PAUD/TK/RA", "Kecil", "Sekolah"),
    ("KS-02", "SD/MI Kelas 1–3", "Kecil", "Sekolah"),
    ("KS-03", "SD/MI Kelas 4–6", "Besar", "Sekolah"),
    ("KS-04", "SMP/MTs", "Besar", "Sekolah"),
    ("KS-05", "SMA/MA/SMK/SLB", "Besar", "Sekolah"),
    ("KS-06", "Santri", "Besar", "Sekolah"),
    ("KS-07", "Ibu Hamil", "Besar", "3B"),
    ("KS-08", "Ibu Menyusui", "Besar", "3B"),
    ("KS-09", "Anak Balita (6–59 bulan)", "Kecil", "3B"),
    ("PTK", "Pendidik dan Tenaga Kependidikan", "Besar", "Sekolah"),
]

DEFAULT_PARAMETERS = {
    "incentiveTariff": 2000,
    "rawSmall": 8000,
    "rawLarge": 10000,
    "operationalPerPm": 3000,
    "maxVa": 500_000_000,
    "maxHpePerWeek": 5,
    "uploadHour": 6,
    "bufferPct": None,
    "dateTolerance": 1,
    "applyIndexRaw": True,
    "applyIndexOp": True,
    "cityIndex": 1.0,
    "cityIndexSource": "",
}

OPERATIONAL_DEFAULTS = [
    "Relawan",
    "Insentif Penanggung Jawab Satuan Pendidikan",
    "Insentif Kader Posyandu",
    "BPJS Ketenagakerjaan",
    "Listrik",
    "Air PDAM",
    "Air minum / galon",
    "Gas",
    "Sewa kendaraan",
    "BBM",
    "Pulsa dan internet",
    "ATK",
    "Alat kebersihan",
    "APD",
    "Biaya rapat koordinasi",
    "Lain-lain",
]


def as_number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("Rp", "").replace(" ", "")
    if not text:
        return default
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def as_int(value: Any, default: int = 0) -> int:
    return int(round(as_number(value, default)))


def yes(value: Any) -> bool:
    return str(value or "").strip().lower() in {"ya", "yes", "y", "true", "1", "signed", "terlampir", "ada"}


def https_url(value: Any) -> bool:
    return str(value or "").strip().lower().startswith("https://")


def parse_iso_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def parse_time(value: Any) -> time | None:
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    return None


def valid_id(identity_type: Any, number: Any) -> bool:
    digits = re.sub(r"\D", "", str(number or ""))
    kind = str(identity_type or "").strip().upper()
    if kind == "NIK":
        return len(digits) == 16
    if kind == "NIP":
        return len(digits) == 18
    return bool(digits)


def parameters(masters: dict[str, Any]) -> dict[str, Any]:
    out = deepcopy(DEFAULT_PARAMETERS)
    out.update(masters.get("parameters") or {})
    return out


def master_target_by_group(masters: dict[str, Any]) -> dict[str, float]:
    totals: dict[str, float] = defaultdict(float)
    for row in masters.get("beneficiaries") or []:
        if str(row.get("status") or "Aktif").strip().lower() in {"nonaktif", "inactive"}:
            continue
        code = str(row.get("groupCode") or row.get("kodeKelompok") or row.get("code") or "").upper().strip()
        if code in {x[0] for x in GROUPS}:
            totals[code] += as_number(row.get("targetPm") if "targetPm" in row else row.get("target"))
    return totals


def merged_pm_rows(masters: dict[str, Any], daily: dict[str, Any], effective: bool) -> list[dict[str, Any]]:
    targets = master_target_by_group(masters)
    supplied = daily.get("pm", {}).get("rows") or []
    supplied_by_code = {
        str(row.get("code") or row.get("groupCode") or "").upper().strip(): row
        for row in supplied
    }
    result: list[dict[str, Any]] = []
    for code, label, portion, pic in GROUPS:
        source = supplied_by_code.get(code, {})
        target = as_number(source.get("targetPm"), targets.get(code, 0))
        distributed = as_number(source.get("distributed"))
        received = as_number(source.get("received"))
        not_received = as_number(source.get("notReceived"))
        bast_no = str(source.get("bastNo") or "").strip()
        bast_link = str(source.get("bastLink") or "").strip()
        bnba = source.get("bnba")
        bast_ok = received <= 0 or (bool(bast_no) and https_url(bast_link))
        calculated = received if effective and yes(bnba) and bast_ok else 0
        result.append({
            "code": code,
            "label": label,
            "portion": portion,
            "pic": pic,
            "targetPm": target,
            "distributed": distributed,
            "received": received,
            "notReceived": not_received,
            "reason": str(source.get("reason") or "").strip(),
            "difference": distributed - received - not_received,
            "bnba": "Ya" if yes(bnba) else "Tidak",
            "bastNo": bast_no,
            "bastLink": bast_link,
            "bastStatus": "Terlampir" if bast_ok and received > 0 else ("Tidak wajib" if received <= 0 else "Belum lengkap"),
            "calculatedPm": calculated,
            "achievement": (calculated / target) if target > 0 else 0,
            "note": str(source.get("note") or ""),
        })
    return result


def final_plan_raw_rows(final_plan: dict[str, Any] | None, service_date: str) -> list[dict[str, Any]]:
    payload = (final_plan or {}).get("payload") or {}
    shopping = (payload.get("shoppingListJSON") or {}).get("shoppingList") or []
    rows = []
    for item in shopping[:40]:
        qty = as_number(item.get("jumlah") if "jumlah" in item else item.get("qty"))
        price = as_number(item.get("harga_satuan") if "harga_satuan" in item else item.get("price"))
        rows.append({
            "date": service_date,
            "name": str(item.get("item") or item.get("name") or item.get("source_ingredient") or "").strip(),
            "category": str(item.get("category") or item.get("supplier_category") or ""),
            "qty": qty,
            "unit": str(item.get("satuan") or item.get("unit") or ""),
            "price": price,
            "supplier": str(item.get("supplier") or item.get("vendor") or ""),
            "invoiceNo": "",
            "evidenceLink": "",
            "note": str(item.get("note") or ""),
            "source": "FINAL_KALKULATOR",
        })
    return [row for row in rows if row["name"]]


def raw_rows(daily: dict[str, Any], final_plan: dict[str, Any] | None, service_date: str) -> list[dict[str, Any]]:
    rows = deepcopy(daily.get("rawMaterials") or [])
    if not rows:
        rows = final_plan_raw_rows(final_plan, service_date)
    for row in rows:
        row["qty"] = as_number(row.get("qty"))
        row["price"] = as_number(row.get("price"))
        row["amount"] = row["qty"] * row["price"]
    return rows[:40]


def volunteer_rows(masters: dict[str, Any], daily: dict[str, Any]) -> list[dict[str, Any]]:
    master = {
        str(row.get("code") or row.get("kode") or row.get("name") or "").strip(): row
        for row in masters.get("volunteers") or []
        if str(row.get("status") or "Aktif").lower() not in {"nonaktif", "inactive"}
    }
    supplied = daily.get("volunteerPayments") or []
    if supplied:
        rows = deepcopy(supplied)
    else:
        rows = [{
            "volunteerCode": key,
            "name": value.get("name") or value.get("nama") or key,
            "role": value.get("role") or value.get("tugas") or "",
            "workDays": 0,
            "dailyRate": value.get("dailyRate") or value.get("rate") or 0,
            "paymentMethod": value.get("paymentMethod") or "",
            "receiptNo": "",
            "evidenceLink": "",
            "date": "",
        } for key, value in master.items()]
    for row in rows:
        key = str(row.get("volunteerCode") or row.get("code") or row.get("name") or "").strip()
        ref = master.get(key, {})
        row["name"] = row.get("name") or ref.get("name") or ref.get("nama") or key
        row["role"] = row.get("role") or ref.get("role") or ref.get("tugas") or ""
        row["dailyRate"] = as_number(row.get("dailyRate"), as_number(ref.get("dailyRate") or ref.get("rate")))
        row["workDays"] = as_number(row.get("workDays"))
        row["amount"] = row["dailyRate"] * row["workDays"]
        row["paymentMethod"] = row.get("paymentMethod") or ref.get("paymentMethod") or ""
    return rows[:60]


def incentive_recipient_rows(daily: dict[str, Any]) -> list[dict[str, Any]]:
    rows = deepcopy(daily.get("incentiveRecipients") or [])
    for row in rows:
        row["amount"] = as_number(row.get("amount"))
    return rows


def operational_rows(masters: dict[str, Any], daily: dict[str, Any]) -> list[dict[str, Any]]:
    rows = deepcopy(daily.get("operations") or [])
    master_by_code = {str(x.get("code") or ""): x for x in masters.get("operations") or []}
    for row in rows:
        ref = master_by_code.get(str(row.get("itemCode") or ""), {})
        row["description"] = row.get("description") or ref.get("name") or ref.get("description") or ""
        row["unit"] = row.get("unit") or ref.get("unit") or ""
        row["qty"] = as_number(row.get("qty"))
        row["price"] = as_number(row.get("price"), as_number(ref.get("defaultPrice")))
        row["amount"] = row["qty"] * row["price"]
    return rows


def assign_document_numbers(rows: list[dict[str, Any]], base_key: str, explicit_key: str = "proofNo") -> list[dict[str, Any]]:
    cloned = deepcopy(rows)
    groups: dict[str, list[int]] = defaultdict(list)
    for index, row in enumerate(cloned):
        base = str(row.get(explicit_key) or row.get(base_key) or "").strip()
        row["_baseProofNo"] = base
        if base:
            groups[base].append(index)
    for base, indexes in groups.items():
        if len(indexes) == 1:
            cloned[indexes[0]]["proofNoDerived"] = base
        else:
            width = max(2, len(str(len(indexes))))
            for seq, index in enumerate(indexes, start=1):
                cloned[index]["proofNoDerived"] = f"{base}-{seq:0{width}d}"
    return cloned


def build_register(
    raw: list[dict[str, Any]],
    operations: list[dict[str, Any]],
    volunteers: list[dict[str, Any]],
    incentive_recipients: list[dict[str, Any]],
    daily: dict[str, Any],
) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []

    raw_numbered = assign_document_numbers(raw, "invoiceNo")
    for idx, row in enumerate(raw_numbered, 1):
        if as_number(row.get("amount")) <= 0 and not row.get("_baseProofNo"):
            continue
        entries.append({
            "source": "B_BahanBaku",
            "code": f"BB-{idx:03d}",
            "proofNo": row.get("proofNoDerived") or row.get("_baseProofNo") or "",
            "date": row.get("date") or "",
            "amount": as_number(row.get("amount")),
            "link": row.get("evidenceLink") or "",
        })

    op_numbered = assign_document_numbers(operations, "invoiceNo")
    for idx, row in enumerate(op_numbered, 1):
        if as_number(row.get("amount")) <= 0 and not row.get("_baseProofNo"):
            continue
        entries.append({
            "source": "C_Operasional",
            "code": f"OP-{idx:03d}",
            "proofNo": row.get("proofNoDerived") or row.get("_baseProofNo") or "",
            "date": row.get("date") or "",
            "amount": as_number(row.get("amount")),
            "link": row.get("evidenceLink") or "",
        })

    volunteer_base = str(daily.get("volunteerReceiptBaseNo") or "").strip()
    volunteer_for_number = []
    for row in volunteers:
        copy = deepcopy(row)
        copy["receiptBase"] = copy.get("receiptNo") or volunteer_base
        volunteer_for_number.append(copy)
    volunteer_numbered = assign_document_numbers(volunteer_for_number, "receiptBase")
    for idx, row in enumerate(volunteer_numbered, 1):
        if as_number(row.get("amount")) <= 0:
            continue
        entries.append({
            "source": "C1_Relawan",
            "code": f"RL-{idx:03d}",
            "proofNo": row.get("proofNoDerived") or row.get("_baseProofNo") or "",
            "date": row.get("date") or "",
            "amount": as_number(row.get("amount")),
            "link": row.get("evidenceLink") or "",
        })

    incentive_base = str(daily.get("incentiveReceiptBaseNo") or "").strip()
    incentive_numbered = []
    for row in incentive_recipients:
        copy = deepcopy(row)
        copy["receiptBase"] = copy.get("receiptNo") or incentive_base
        incentive_numbered.append(copy)
    incentive_numbered = assign_document_numbers(incentive_numbered, "receiptBase")
    for idx, row in enumerate(incentive_numbered, 1):
        if as_number(row.get("amount")) <= 0:
            continue
        entries.append({
            "source": "C_Operasional",
            "code": f"IN-OP-{idx:03d}",
            "proofNo": row.get("proofNoDerived") or row.get("_baseProofNo") or "",
            "date": row.get("date") or "",
            "amount": as_number(row.get("amount")),
            "link": row.get("evidenceLink") or "",
        })

    incentive = daily.get("incentive") or {}
    if as_number(incentive.get("paidAmount")) > 0 or incentive.get("proofNo"):
        entries.append({
            "source": "D_Insentif",
            "code": "IN-001",
            "proofNo": str(incentive.get("proofNo") or "").strip(),
            "date": incentive.get("paymentDate") or "",
            "amount": as_number(incentive.get("paidAmount")),
            "link": incentive.get("evidenceLink") or "",
        })

    for idx, row in enumerate(daily.get("topups") or [], 1):
        amount = as_number(row.get("rawAmount")) + as_number(row.get("operationalAmount")) + as_number(row.get("incentiveAmount"))
        if amount <= 0 and not row.get("receiptNo"):
            continue
        entries.append({
            "source": "E_Saldo",
            "code": f"TU-{idx:03d}",
            "proofNo": str(row.get("receiptNo") or "").strip(),
            "date": row.get("date") or "",
            "amount": amount,
            "link": row.get("evidenceLink") or "",
        })

    counts = Counter(str(x["proofNo"]).strip().upper() for x in entries if str(x["proofNo"]).strip())
    for row in entries:
        proof = str(row["proofNo"]).strip()
        row["proofStatus"] = "UNIK" if proof and counts[proof.upper()] == 1 else ("DUPLIKAT" if proof else "BELUM ADA")
    return entries


def compute_preview(
    masters: dict[str, Any],
    daily: dict[str, Any],
    service_date: str,
    effective: bool,
    final_plan: dict[str, Any] | None = None,
) -> dict[str, Any]:
    params = parameters(masters)
    pm = merged_pm_rows(masters, daily, effective)
    raw = raw_rows(daily, final_plan, service_date)
    volunteers = volunteer_rows(masters, daily)
    incentive_recipients = incentive_recipient_rows(daily)
    operations = operational_rows(masters, daily)

    production = daily.get("pm", {}).get("production") or {}
    distributed = sum(row["distributed"] for row in pm)
    received = sum(row["received"] for row in pm)
    not_received = sum(row["notReceived"] for row in pm)
    calculated_pm = sum(row["calculatedPm"] for row in pm)
    produced_default = as_number((final_plan or {}).get("payload", {}).get("porsiKecil")) + as_number((final_plan or {}).get("payload", {}).get("porsiBesar"))
    produced = as_number(production.get("produced"), produced_default)
    organoleptic = as_number(production.get("organoleptic"))
    retained_sample = as_number(production.get("retainedSample"))
    production_not_distributed = as_number(production.get("notDistributed"))
    buffer_qty = as_number(production.get("buffer"))
    production_diff = produced - (distributed + organoleptic + retained_sample + production_not_distributed + buffer_qty)

    incentive_pm = calculated_pm + (organoleptic if effective else 0) + (retained_sample if effective else 0)
    index = as_number(params.get("cityIndex"), 1.0) or 1.0
    raw_small = as_number(params.get("rawSmall"), 8000) * (index if params.get("applyIndexRaw", True) else 1)
    raw_large = as_number(params.get("rawLarge"), 10000) * (index if params.get("applyIndexRaw", True) else 1)
    small_pm = sum(row["calculatedPm"] for row in pm if row["portion"] == "Kecil")
    large_pm = sum(row["calculatedPm"] for row in pm if row["portion"] == "Besar")
    weighted_raw_pagu = ((small_pm * raw_small + large_pm * raw_large) / calculated_pm) if calculated_pm else 0
    raw_total = sum(as_number(row.get("amount")) for row in raw)
    raw_per_portion = raw_total / produced if produced else 0

    volunteer_total = sum(as_number(row.get("amount")) for row in volunteers)
    incentive_recipient_total = sum(as_number(row.get("amount")) for row in incentive_recipients)
    normal_operational_total = sum(as_number(row.get("amount")) for row in operations)
    operational_total = volunteer_total + incentive_recipient_total + normal_operational_total
    op_pagu = as_number(params.get("operationalPerPm"), 3000) * (index if params.get("applyIndexOp", True) else 1)
    op_per_pm = operational_total / incentive_pm if incentive_pm else 0

    incentive = daily.get("incentive") or {}
    incentive_tariff = as_number(params.get("incentiveTariff"), 2000)
    incentive_calculated = incentive_pm * incentive_tariff
    register = build_register(raw, operations, volunteers, incentive_recipients, daily)

    balance = daily.get("balance") or {}
    opening = {
        "raw": as_number(balance.get("openingRaw")),
        "operational": as_number(balance.get("openingOperational")),
        "incentive": as_number(balance.get("openingIncentive")),
    }
    topups = daily.get("topups") or []
    topup_total = {
        "raw": sum(as_number(x.get("rawAmount")) for x in topups),
        "operational": sum(as_number(x.get("operationalAmount")) for x in topups),
        "incentive": sum(as_number(x.get("incentiveAmount")) for x in topups),
    }
    expenditure = {
        "raw": raw_total,
        "operational": operational_total,
        "incentive": as_number(incentive.get("paidAmount")),
    }
    closing = {key: opening[key] + topup_total[key] - expenditure[key] for key in opening}
    closing_total = sum(closing.values())
    bank_balance = as_number(balance.get("bankBalance"))
    bank_difference = bank_balance - closing_total

    proposal = daily.get("topupProposal") or {}
    proposal_total = as_number(proposal.get("raw")) + as_number(proposal.get("operational")) + as_number(proposal.get("incentive"))
    room_to_max = max(0, as_number(params.get("maxVa"), 500_000_000) - closing_total)

    service_dt = parse_iso_date(service_date)
    tolerance = as_int(params.get("dateTolerance"), 1)
    valid_from = service_dt - timedelta(days=tolerance) if service_dt else None
    valid_to = service_dt

    def register_date_ok(row: dict[str, Any]) -> bool:
        tx = parse_iso_date(row.get("date"))
        if as_number(row.get("amount")) <= 0:
            return True
        return bool(tx and valid_from and valid_to and valid_from <= tx <= valid_to)

    evidence_missing = [r for r in register if as_number(r.get("amount")) > 0 and not https_url(r.get("link"))]
    duplicates = [r for r in register if r.get("proofStatus") == "DUPLIKAT"]
    invalid_dates = [r for r in register if not register_date_ok(r)]

    identity = masters.get("identity") or {}
    required_identity = [
        "lpdhNumber", "sppgId", "sppgName", "village", "district", "city",
        "province", "foundation", "vaNumber", "bankName",
    ]
    identity_ok = all(str(identity.get(key) or "").strip() for key in required_identity)

    all_distributions_balance = all(abs(row["difference"]) < 0.0001 for row in pm)
    reasons_ok = all(row["notReceived"] <= 0 or bool(row["reason"]) for row in pm)
    receiving_docs_ok = all(row["received"] <= 0 or (row["bnba"] == "Ya" and bool(row["bastNo"])) for row in pm)
    bast_links_ok = all(row["received"] <= 0 or https_url(row["bastLink"]) for row in pm)
    target_ok = all(row["calculatedPm"] <= row["targetPm"] + 0.0001 for row in pm if row["targetPm"] > 0)

    volunteer_complete = all(
        as_number(row.get("amount")) <= 0 or (
            bool(str(row.get("receiptNo") or daily.get("volunteerReceiptBaseNo") or "").strip())
            and https_url(row.get("evidenceLink"))
            and parse_iso_date(row.get("date")) is not None
        )
        for row in volunteers
    )

    eligibility = incentive.get("eligibility") or {}
    incentive_eligible = (
        effective
        and not yes(eligibility.get("contamination"))
        and not yes(eligibility.get("fatalIncident"))
        and not yes(eligibility.get("suspended"))
        and yes(eligibility.get("verified"))
        and yes(eligibility.get("pmInputSipgn"))
    )

    paid = as_number(incentive.get("paidAmount"))
    statement = as_number(incentive.get("statementAmount"))
    incentive_payment_matches = paid == statement if (paid or statement) else False
    incentive_evidence_ok = paid <= 0 or (
        bool(str(incentive.get("proofNo") or "").strip())
        and bool(str(incentive.get("receiptNo") or "").strip())
        and yes(incentive.get("receiptSigned"))
        and https_url(incentive.get("evidenceLink"))
        and parse_iso_date(incentive.get("paymentDate")) is not None
    )

    topup_evidence_ok = all(
        (as_number(x.get("rawAmount")) + as_number(x.get("operationalAmount")) + as_number(x.get("incentiveAmount"))) <= 0
        or (bool(str(x.get("receiptNo") or "").strip()) and https_url(x.get("evidenceLink")) and parse_iso_date(x.get("date")) is not None)
        for x in topups
    )

    signers = masters.get("signers") or []
    signer_ok = len(signers) >= 3 and all(
        str(x.get("name") or "").strip()
        and valid_id(x.get("identityType"), x.get("identityNumber"))
        and yes(x.get("signed"))
        for x in signers[:3]
    )

    upload = daily.get("upload") or {}
    upload_date = parse_iso_date(upload.get("date"))
    upload_time = parse_time(upload.get("time"))
    deadline_date = service_dt + timedelta(days=1) if service_dt else None
    deadline_hour = as_int(params.get("uploadHour"), 6)
    upload_ok = bool(
        upload_date and upload_time and deadline_date
        and (
            upload_date < deadline_date
            or (upload_date == deadline_date and upload_time <= time(deadline_hour, 0))
        )
    )

    buffer_limit = params.get("bufferPct")
    buffer_ok = True
    if buffer_limit not in (None, "") and produced > 0:
        buffer_ok = (buffer_qty / produced) <= as_number(buffer_limit) / 100

    checks = [
        ("01", "Identitas wajib lengkap", identity_ok, "Lengkapi Identitas SPPG, VA, bank, dan yayasan."),
        ("02", "Hari pelayanan efektif / batas HPE", effective, "Tanggal belum ditandai sebagai Hari Pelayanan Efektif."),
        ("03", "Produksi = distribusi + sampel + sisa + buffer", abs(production_diff) < 0.0001, f"Selisih produksi {production_diff:,.2f}."),
        ("04", "POP dan Fleet seimbang", all_distributions_balance, "Ada kelompok dengan jumlah distribusi tidak sama dengan diterima + tidak diterima."),
        ("05", "Alasan tidak diterima lengkap", reasons_ok, "Isi alasan untuk PM yang tidak menerima."),
        ("06", "Buffer sesuai ambang", buffer_ok, "Buffer melebihi batas parameter."),
        ("07", "BNBA dan nomor BAST penerima lengkap", receiving_docs_ok, "Ada penerima dengan dokumen BNBA/BAST belum lengkap."),
        ("08", "Link BAST valid", bast_links_ok, "Link BAST wajib HTTPS untuk kelompok yang menerima."),
        ("09", "PM dihitung tidak melebihi target SPS", target_ok, "Ada PM dihitung melebihi target master."),
        ("10", "Indeks kemahalan tersedia", index > 0, "Isi indeks kemahalan pada Master Parameter."),
        ("11", "Biaya bahan baku per porsi dalam pagu", produced > 0 and raw_per_portion <= weighted_raw_pagu + 0.0001, f"Biaya/porsi Rp{raw_per_portion:,.0f}; pagu Rp{weighted_raw_pagu:,.0f}."),
        ("12", "Biaya operasional per PM dalam pagu", incentive_pm > 0 and op_per_pm <= op_pagu + 0.0001, f"Operasional/PM Rp{op_per_pm:,.0f}; pagu Rp{op_pagu:,.0f}."),
        ("13", "Bukti transaksi bernilai terlampir", len(evidence_missing) == 0, f"{len(evidence_missing)} transaksi belum punya link HTTPS."),
        ("14", "Nomor bukti transaksi unik", len(duplicates) == 0, f"{len(duplicates)} baris memakai nomor bukti duplikat."),
        ("15", "Tanggal transaksi dalam periode", len(invalid_dates) == 0, f"{len(invalid_dates)} transaksi di luar periode {valid_from} s.d. {valid_to}."),
        ("16", "Bukti pembayaran relawan lengkap", volunteer_complete, "Ada upah relawan tanpa kuitansi/tanggal/link bukti."),
        ("17", "Syarat insentif terpenuhi", incentive_eligible, "Syarat insentif belum lengkap atau ada kondisi penggugur."),
        ("18", "Perhitungan insentif = tarif × PM", abs(as_number(incentive.get("calculatedAmount"), incentive_calculated) - incentive_calculated) < 0.0001, f"Seharusnya Rp{incentive_calculated:,.0f}."),
        ("19", "Pembayaran insentif sesuai pernyataan PPK", incentive_payment_matches, "Nilai dibayar harus sama dengan nilai pernyataan PPK."),
        ("20", "Bukti pembayaran insentif lengkap", incentive_evidence_ok, "Nomor bukti/kuitansi/tanda tangan/link/tanggal insentif belum lengkap."),
        ("21", "Bukti penerimaan Top Up lengkap", topup_evidence_ok, "Ada top up tanpa kuitansi/tanggal/link bukti."),
        ("22", "Saldo komponen tidak negatif", all(value >= -0.0001 for value in closing.values()), "Ada saldo komponen negatif."),
        ("23", "Saldo komponen sama dengan saldo VA", abs(bank_difference) < 0.01, f"Selisih saldo VA Rp{bank_difference:,.0f}."),
        ("24", "Tiga pengesah lengkap dan identitas valid", signer_ok, "Lengkapi 3 pengesah, NIK/NIP, dan status tanda tangan."),
        ("25", "Rencana unggah H+1 sebelum batas waktu", upload_ok, "Isi tanggal dan jam upload; batas H+1 pukul yang ditentukan."),
        ("26", "Usulan top up tidak melebihi plafon VA", proposal_total <= room_to_max + 0.0001, f"Usulan Rp{proposal_total:,.0f}; ruang VA Rp{room_to_max:,.0f}."),
    ]
    check_rows = [{"no": no, "check": label, "ok": bool(ok), "status": "OK" if ok else "PERIKSA", "detail": detail if not ok else ""} for no, label, ok, detail in checks]
    error_count = sum(1 for row in check_rows if not row["ok"])

    return {
        "serviceDate": service_date,
        "effective": effective,
        "pmRows": pm,
        "production": {
            "produced": produced,
            "distributed": distributed,
            "received": received,
            "notReceived": not_received,
            "organoleptic": organoleptic,
            "retainedSample": retained_sample,
            "notDistributed": production_not_distributed,
            "buffer": buffer_qty,
            "difference": production_diff,
            "calculatedPm": calculated_pm,
            "incentivePm": incentive_pm,
        },
        "rawMaterials": raw,
        "rawTotal": raw_total,
        "rawPerPortion": raw_per_portion,
        "weightedRawPagu": weighted_raw_pagu,
        "rawStatus": "DALAM PAGU" if produced > 0 and raw_per_portion <= weighted_raw_pagu + 0.0001 else "MELEBIHI PAGU",
        "operations": operations,
        "volunteers": volunteers,
        "incentiveRecipients": incentive_recipients,
        "volunteerTotal": volunteer_total,
        "incentiveRecipientTotal": incentive_recipient_total,
        "operationalTotal": operational_total,
        "operationalPerPm": op_per_pm,
        "operationalPagu": op_pagu,
        "operationalStatus": "DALAM PAGU" if incentive_pm > 0 and op_per_pm <= op_pagu + 0.0001 else "MELEBIHI PAGU",
        "incentiveCalculated": incentive_calculated,
        "register": register,
        "balance": {
            "opening": opening,
            "topups": topup_total,
            "expenditure": expenditure,
            "closing": closing,
            "closingTotal": closing_total,
            "bankBalance": bank_balance,
            "bankDifference": bank_difference,
        },
        "topup": {"proposalTotal": proposal_total, "roomToMax": room_to_max, "withinMax": proposal_total <= room_to_max + 0.0001},
        "validPeriod": {"from": valid_from.isoformat() if valid_from else "", "to": valid_to.isoformat() if valid_to else ""},
        "checks": check_rows,
        "errorCount": error_count,
        "ready": error_count == 0,
        "parameters": params,
        "finalPlan": final_plan or {},
    }


def _excel_date(value: Any) -> date | None:
    return parse_iso_date(value)


def _set_if(ws, cell: str, value: Any) -> None:
    if value is not None:
        ws[cell] = value


def populate_workbook(
    masters: dict[str, Any],
    daily: dict[str, Any],
    preview: dict[str, Any],
    service_date: str,
    template_bytes: bytes | None = None,
) -> bytes:
    wb = load_lpdh_workbook(template_bytes)
    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    identity = masters.get("identity") or {}
    ws = wb["Identitas"]
    identity_cells = {
        "B5": identity.get("lpdhNumber"),
        "B6": identity.get("sppgId"),
        "B7": identity.get("sppgName"),
        "B8": identity.get("village"),
        "B9": identity.get("district"),
        "B10": identity.get("city"),
        "B11": identity.get("province"),
        "B12": identity.get("foundation"),
        "B13": identity.get("vaNumber"),
        "B14": identity.get("bankName"),
        "B15": _excel_date(service_date),
        "B18": daily.get("dayStatus") or "Hari Pelayanan Efektif",
        "B19": as_int(daily.get("hpeNumber"), 1),
        "B25": _excel_date((daily.get("upload") or {}).get("date")),
        "B26": parse_time((daily.get("upload") or {}).get("time")),
    }
    for cell, value in identity_cells.items():
        _set_if(ws, cell, value)
    signers = masters.get("signers") or []
    for offset in range(3):
        if offset >= len(signers):
            continue
        signer = signers[offset]
        row = 36 + offset
        _set_if(ws, f"B{row}", signer.get("name"))
        _set_if(ws, f"C{row}", signer.get("identityType"))
        _set_if(ws, f"D{row}", signer.get("identityNumber"))
        _set_if(ws, f"F{row}", "Ya" if yes(signer.get("signed")) else "Tidak")

    ref = wb["Ref"]
    params = preview["parameters"]
    ref_values = {
        "B5": as_number(params.get("incentiveTariff"), 2000),
        "B6": as_number(params.get("rawSmall"), 8000),
        "B7": as_number(params.get("rawLarge"), 10000),
        "B8": as_number(params.get("operationalPerPm"), 3000),
        "B9": as_number(params.get("maxVa"), 500_000_000),
        "B10": as_int(params.get("maxHpePerWeek"), 5),
        "B11": as_int(params.get("uploadHour"), 6),
        "B12": None if params.get("bufferPct") in (None, "") else as_number(params.get("bufferPct")),
        "B13": as_int(params.get("dateTolerance"), 1),
        "B14": "Ya" if params.get("applyIndexRaw", True) else "Tidak",
        "B15": "Ya" if params.get("applyIndexOp", True) else "Tidak",
    }
    for cell, value in ref_values.items():
        _set_if(ref, cell, value)

    pm_ws = wb["A_PM"]
    for idx, row_data in enumerate(preview["pmRows"], start=6):
        _set_if(pm_ws, f"E{idx}", row_data["targetPm"])
        _set_if(pm_ws, f"F{idx}", row_data["distributed"])
        _set_if(pm_ws, f"G{idx}", row_data["received"])
        _set_if(pm_ws, f"H{idx}", row_data["notReceived"])
        _set_if(pm_ws, f"I{idx}", row_data["reason"])
        _set_if(pm_ws, f"K{idx}", row_data["bnba"])
        _set_if(pm_ws, f"L{idx}", row_data["bastNo"])
        _set_if(pm_ws, f"M{idx}", row_data["bastLink"])
    prod = preview["production"]
    for cell, value in {
        "C19": prod["produced"], "C21": prod["organoleptic"], "C22": prod["retainedSample"],
        "C23": prod["notDistributed"], "C24": prod["buffer"],
    }.items():
        _set_if(pm_ws, cell, value)

    raw_ws = wb["B_BahanBaku"]
    raw_numbered = assign_document_numbers(preview["rawMaterials"], "invoiceNo")
    for i, item in enumerate(raw_numbered[:40], start=6):
        for col, value in {
            "C": _excel_date(item.get("date")), "D": item.get("name"), "E": item.get("category"),
            "F": as_number(item.get("qty")), "G": item.get("unit"), "H": as_number(item.get("price")),
            "J": item.get("supplier"), "K": item.get("proofNoDerived") or item.get("_baseProofNo") or "",
            "L": item.get("evidenceLink"), "O": item.get("note"),
        }.items():
            _set_if(raw_ws, f"{col}{i}", value)

    op_ws = wb["C_Operasional"]
    op_rows = preview["operations"]
    op_numbered = assign_document_numbers(op_rows, "invoiceNo")
    volunteer_total = preview["volunteerTotal"]
    incentive_total = preview["incentiveRecipientTotal"]
    special = [
        {"description": "Relawan", "amount": volunteer_total, "date": service_date, "proofNoDerived": daily.get("volunteerReceiptBaseNo") or "", "evidenceLink": daily.get("volunteerBatchEvidenceLink") or "", "unit": "orang", "qty": len([x for x in preview["volunteers"] if as_number(x.get("amount")) > 0]), "price": 0},
        {"description": "Insentif Penanggung Jawab Satuan Pendidikan", "amount": sum(as_number(x.get("amount")) for x in preview["incentiveRecipients"] if str(x.get("type") or "").lower() in {"guru","sekolah","penanggung jawab satuan pendidikan"}), "date": service_date, "proofNoDerived": daily.get("incentiveReceiptBaseNo") or "", "evidenceLink": daily.get("incentiveBatchEvidenceLink") or "", "unit": "orang", "qty": 0, "price": 0},
        {"description": "Insentif Kader Posyandu", "amount": sum(as_number(x.get("amount")) for x in preview["incentiveRecipients"] if str(x.get("type") or "").lower() in {"kader","posyandu","kader posyandu"}), "date": service_date, "proofNoDerived": daily.get("incentiveReceiptBaseNo") or "", "evidenceLink": daily.get("incentiveBatchEvidenceLink") or "", "unit": "orang", "qty": 0, "price": 0},
    ]
    by_name = {str(x.get("description") or "").strip().lower(): x for x in op_numbered}
    for pos, default_name in enumerate(OPERATIONAL_DEFAULTS, start=6):
        item = by_name.get(default_name.lower())
        if default_name == "Relawan":
            item = special[0]
        elif default_name == "Insentif Penanggung Jawab Satuan Pendidikan":
            item = special[1]
        elif default_name == "Insentif Kader Posyandu":
            item = special[2]
        if not item:
            continue
        _set_if(op_ws, f"C{pos}", _excel_date(item.get("date")))
        _set_if(op_ws, f"D{pos}", item.get("description") or default_name)
        _set_if(op_ws, f"E{pos}", as_number(item.get("qty")))
        _set_if(op_ws, f"F{pos}", item.get("unit"))
        if as_number(item.get("price")):
            _set_if(op_ws, f"G{pos}", as_number(item.get("price")))
        elif as_number(item.get("amount")) and as_number(item.get("qty")):
            _set_if(op_ws, f"G{pos}", as_number(item.get("amount")) / as_number(item.get("qty")))
        _set_if(op_ws, f"I{pos}", item.get("proofNoDerived") or item.get("_baseProofNo") or "")
        _set_if(op_ws, f"J{pos}", item.get("evidenceLink"))
        _set_if(op_ws, f"M{pos}", item.get("note"))

    rel_ws = wb["C1_Relawan"]
    volunteer_numbered = []
    volunteer_base = str(daily.get("volunteerReceiptBaseNo") or "")
    for x in preview["volunteers"]:
        y = deepcopy(x)
        y["receiptBase"] = y.get("receiptNo") or volunteer_base
        volunteer_numbered.append(y)
    volunteer_numbered = assign_document_numbers(volunteer_numbered, "receiptBase")
    for i, item in enumerate(volunteer_numbered[:60], start=6):
        _set_if(rel_ws, f"C{i}", item.get("name"))
        _set_if(rel_ws, f"D{i}", item.get("role"))
        _set_if(rel_ws, f"E{i}", _excel_date(item.get("date")))
        _set_if(rel_ws, f"F{i}", as_number(item.get("workDays")))
        _set_if(rel_ws, f"G{i}", as_number(item.get("dailyRate")))
        _set_if(rel_ws, f"I{i}", item.get("paymentMethod"))
        _set_if(rel_ws, f"J{i}", item.get("proofNoDerived") or item.get("_baseProofNo") or "")
        _set_if(rel_ws, f"K{i}", item.get("evidenceLink"))

    ins_ws = wb["D_Insentif"]
    incentive = daily.get("incentive") or {}
    elig = incentive.get("eligibility") or {}
    # Rows 6-10 are the editable eligibility questions in the official template.
    for cell, value in {
        "C6": "Ya" if yes(elig.get("contamination")) else "Tidak",
        "C7": "Ya" if yes(elig.get("fatalIncident")) else "Tidak",
        "C8": "Ya" if yes(elig.get("suspended")) else "Tidak",
        "C9": "Ya" if yes(elig.get("verified")) else "Tidak",
        "C10": "Ya" if yes(elig.get("pmInputSipgn")) else "Tidak",
        "C19": incentive.get("ppkStatementNo"),
        "C20": as_number(incentive.get("statementAmount")),
        "C21": as_number(incentive.get("paidAmount")),
        "C22": _excel_date(incentive.get("paymentDate")),
        "C23": incentive.get("proofNo"),
        "C24": incentive.get("receiptNo"),
        "C25": "Ya" if yes(incentive.get("receiptSigned")) else "Tidak",
        "C26": incentive.get("evidenceLink"),
        "C27": incentive.get("vaReference"),
    }.items():
        _set_if(ins_ws, cell, value)

    saldo_ws = wb["E_Saldo"]
    bal = daily.get("balance") or {}
    for cell, value in {
        "B5": as_number(bal.get("openingRaw")), "B6": as_number(bal.get("openingOperational")),
        "B7": as_number(bal.get("openingIncentive")), "E10": as_number(bal.get("bankBalance")),
    }.items():
        _set_if(saldo_ws, cell, value)
    for idx, item in enumerate((daily.get("topups") or [])[:5], start=19):
        for col, value in {
            "C": _excel_date(item.get("date")), "D": item.get("reference"),
            "E": as_number(item.get("rawAmount")), "F": as_number(item.get("operationalAmount")),
            "G": as_number(item.get("incentiveAmount")), "I": item.get("receiptNo"),
            "J": item.get("evidenceLink"),
        }.items():
            _set_if(saldo_ws, f"{col}{idx}", value)

    topup_ws = wb["F_TopUp"]
    proposal = daily.get("topupProposal") or {}
    for cell, value in {"C5": as_number(proposal.get("raw")), "C6": as_number(proposal.get("operational")), "C7": as_number(proposal.get("incentive"))}.items():
        _set_if(topup_ws, cell, value)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_master_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Petunjuk"
    ws.append(["TEMPLATE IMPORT MASTER LPDH"])
    ws.append(["Isi sheet Master_Penerima, Master_Relawan, Master_Operasional, lalu upload ke aplikasi. Jangan ubah nama kolom."])

    pm = wb.create_sheet("Master_Penerima")
    pm.append(["Kode Unit","Jenis Unit","Nama Sekolah / Posyandu","Kode Kelompok","Kelompok Sasaran","Kategori Porsi","Jenis PIC","Target PM","Nama PIC","No. HP PIC","Alamat","Status Aktif","Catatan"])
    pm.append(["SKL-001","Sekolah","SD Contoh 01","KS-02","SD/MI Kelas 1–3","Kecil","Sekolah",220,"Nama PIC","081234567890","Alamat","Aktif",""])

    rv = wb.create_sheet("Master_Relawan")
    rv.append(["Kode Relawan","Nama Relawan","Tugas","Status Aktif","Besaran Harian (Rp)","Metode Bayar Default","Nama Bank","No. Rekening","Nama Pemilik Rekening","No. HP","Catatan"])
    rv.append(["RL-001","Contoh Relawan","Juru masak","Aktif",90000,"Transfer","BRI","1234567890","Contoh Relawan","081234567890",""])

    op = wb.create_sheet("Master_Operasional")
    op.append(["Kode Item","Nama Item Operasional","Kategori","Satuan Default","Harga Default (Rp)","Sifat Biaya","Vendor Default","Status Aktif","Catatan"])
    for idx, name in enumerate(OPERATIONAL_DEFAULTS[3:], 1):
        op.append([f"OP-{idx:03d}", name, "Operasional", "unit", 0, "Rutin", "", "Aktif", ""])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def parse_master_workbook(content: bytes) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    result: dict[str, list[dict[str, Any]]] = {"beneficiaries": [], "volunteers": [], "operations": []}

    def records(sheet_name: str) -> list[dict[str, Any]]:
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            out.append({headers[idx]: value for idx, value in enumerate(row) if idx < len(headers) and headers[idx]})
        return out

    for row in records("Master_Penerima"):
        result["beneficiaries"].append({
            "code": str(row.get("Kode Unit") or "").strip(),
            "unitType": str(row.get("Jenis Unit") or "").strip(),
            "unitName": str(row.get("Nama Sekolah / Posyandu") or "").strip(),
            "groupCode": str(row.get("Kode Kelompok") or "").strip().upper(),
            "groupName": str(row.get("Kelompok Sasaran") or "").strip(),
            "portionCategory": str(row.get("Kategori Porsi") or "").strip(),
            "picType": str(row.get("Jenis PIC") or "").strip(),
            "targetPm": as_number(row.get("Target PM")),
            "picName": str(row.get("Nama PIC") or "").strip(),
            "phone": str(row.get("No. HP PIC") or "").strip(),
            "address": str(row.get("Alamat") or "").strip(),
            "status": str(row.get("Status Aktif") or "Aktif").strip(),
            "note": str(row.get("Catatan") or "").strip(),
        })

    volunteer_sheet = "Master_Relawan" if "Master_Relawan" in wb.sheetnames else "Import_Relawan"
    for row in records(volunteer_sheet):
        result["volunteers"].append({
            "code": str(row.get("Kode Relawan") or "").strip(),
            "name": str(row.get("Nama Relawan") or "").strip(),
            "role": str(row.get("Tugas") or "").strip(),
            "status": str(row.get("Status Aktif") or "Aktif").strip(),
            "dailyRate": as_number(row.get("Besaran Harian (Rp)")),
            "paymentMethod": str(row.get("Metode Bayar Default") or "").strip(),
            "bankName": str(row.get("Nama Bank") or "").strip(),
            "accountNumber": str(row.get("No. Rekening") or "").strip(),
            "accountName": str(row.get("Nama Pemilik Rekening") or "").strip(),
            "phone": str(row.get("No. HP") or "").strip(),
            "note": str(row.get("Catatan") or "").strip(),
        })

    for row in records("Master_Operasional"):
        result["operations"].append({
            "code": str(row.get("Kode Item") or "").strip(),
            "name": str(row.get("Nama Item Operasional") or "").strip(),
            "category": str(row.get("Kategori") or "").strip(),
            "unit": str(row.get("Satuan Default") or "").strip(),
            "defaultPrice": as_number(row.get("Harga Default (Rp)")),
            "costNature": str(row.get("Sifat Biaya") or "").strip(),
            "vendor": str(row.get("Vendor Default") or "").strip(),
            "status": str(row.get("Status Aktif") or "Aktif").strip(),
            "note": str(row.get("Catatan") or "").strip(),
        })

    return result


def workbook_reference_rows() -> list[list[Any]]:
    if not LPDH_TEMPLATE.is_file():
        return []
    wb = load_workbook(LPDH_TEMPLATE, data_only=False, read_only=True)
    ws = wb["Ref"]
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=min(ws.max_column, 4), values_only=True):
        rows.append(list(row))
    return rows


def encode_bytes(content: bytes) -> str:
    return base64.b64encode(content).decode("ascii")
