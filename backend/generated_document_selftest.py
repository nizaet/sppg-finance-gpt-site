"""Business invariants for document-first daily accounting; runs without a DB."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import argparse

from openpyxl import load_workbook
from pypdf import PdfReader

from generated_document_logic import line_amount, merge_final_documents
from generated_document_pdf import render_document_pdf
from generated_document_reference import MAJA_PROFILES
from lpdh_selftest import fixture
from lpdh_logic import compute_preview, populate_workbook


def document(identifier, kind, items, status="FINAL"):
    prefixes = {"BAHAN_BAKU": "INV-BHN", "OPERASIONAL": "INV-OPS", "UPAH_RELAWAN": "KWT-REL", "INSENTIF_GURU_KADER": "KWT-GURU"}
    header = deepcopy(MAJA_PROFILES["YAYASAN" if kind.startswith("UPAH") or kind.startswith("INSENTIF") else "KOPERASI"])
    header["evidenceLink"] = "https://example.test/signed-evidence.pdf"
    rows = [{"itemName": name, "category": category, "quantity": qty, "unit": unit,
             "unitPrice": price, "lineTotal": line_amount(qty, price), "metadata": metadata}
            for name, category, qty, unit, price, metadata in items]
    return {"id": identifier, "site": "MAJA", "documentType": kind,
            "documentNumber": f"{prefixes[kind]}-MAJA-20261005-{identifier:03d}", "serviceDate": "2026-10-05",
            "status": status, "header": header, "total": sum(x["lineTotal"] for x in rows), "items": rows}


def fixtures():
    return [
        document(1, "BAHAN_BAKU", [("Beras", "Beras", 2, "kg", 1000, {}), ("Telur", "Protein", 1, "kg", 1500, {})]),
        document(2, "OPERASIONAL", [("Gas 50 kg", "Gas", .5, "tabung", 1000, {}), ("Masker", "APD (masker, sarung tangan, penutup kepala)", 1, "pack", 50, {})]),
        document(3, "OPERASIONAL", [("Gas 12 kg", "Gas", 1, "tabung", 200, {}), ("Mama Lemon", "Alat kebersihan", 2, "pouch", 100, {})]),
        document(4, "UPAH_RELAWAN", [("Relawan Uji A", "Upah", 1, "hari", 90000, {"volunteerCode": "V1", "role": "Masak"}), ("Relawan Uji B", "Upah", 1, "hari", 120000, {"volunteerCode": "V2", "role": "Distribusi"})]),
        document(5, "INSENTIF_GURU_KADER", [("Guru Uji", "Insentif", 1, "hari", 10000, {"recipientType": "Guru", "unitName": "Sekolah Uji"}), ("Kader Uji", "Insentif", 1, "hari", 8000, {"recipientType": "Kader", "unitName": "Posyandu Uji"})]),
    ]


def run(output_dir=None):
    service_date, masters, initial, final_plan = fixture()
    initial.update(rawMaterials=[], operations=[], volunteerPayments=[], incentiveRecipients=[])
    docs = fixtures()
    archived = deepcopy(docs)
    for doc in archived:
        doc['header']['evidenceLink'] = ''
        doc['driveUri'] = f"https://drive.google.com/file/d/test-{doc['id']}/view"
    before_archive = deepcopy(archived)
    for doc in before_archive:
        doc.pop('driveUri')
    blank_daily = merge_final_documents(initial, before_archive)
    auto_daily = merge_final_documents(blank_daily, archived)
    for key in ('rawMaterials', 'operations', 'volunteerPayments', 'incentiveRecipients'):
        assert all(row['evidenceLink'] == f"https://drive.google.com/file/d/test-{row['sourceDocumentId']}/view" for row in auto_daily[key]), 'Every document line inherits its own archived PDF'
    auto_preview = compute_preview(masters, auto_daily, service_date, True, final_plan)
    auto_workbook = load_workbook(BytesIO(populate_workbook(masters, auto_daily, auto_preview, service_date)), data_only=False)
    assert auto_workbook['B_BahanBaku']['L6'].value.endswith('test-1/view')
    assert auto_workbook['C1_Relawan']['K6'].value.endswith('test-4/view')
    assert auto_workbook['C1_Relawan']['K7'].value.endswith('test-4/view')
    assert auto_workbook['C_Operasional']['J7'].value.endswith('test-5/view')
    assert auto_workbook['C_Operasional']['J8'].value.endswith('test-5/view')
    auto_daily['volunteerPayments'][0]['evidenceLink'] = 'https://example.test/manual-signed.pdf'
    assert merge_final_documents(auto_daily, archived)['volunteerPayments'][0]['evidenceLink'].endswith('manual-signed.pdf'), 'Preserve manually completed evidence'
    draft_only = deepcopy(docs[0]); draft_only["status"] = "DRAFT"
    assert merge_final_documents(initial, [draft_only]) == initial, "Draft must not affect daily costs"
    initial["rawMaterials"] = [{"source": "FINAL_KALKULATOR", "name": "Estimasi", "qty": 99, "price": 999}]
    imported = merge_final_documents(initial, docs)
    assert imported["balance"] == initial["balance"], "Unrelated daily data must be preserved"
    assert merge_final_documents(imported, docs) == imported, "Repeated import must be idempotent"
    manual_docs = deepcopy(docs)
    manual_docs[0]['documentNumber'] = 'INV/MAJA/001'
    manual_docs[3]['items'][0]['metadata']['receiptNo'] = 'KWT/REL/101'
    manual_docs[3]['items'][1]['metadata']['receiptNo'] = 'KWT/REL/102'
    manual_daily = merge_final_documents(initial, manual_docs)
    manual_preview = compute_preview(masters, manual_daily, service_date, True, final_plan)
    manual_workbook = load_workbook(BytesIO(populate_workbook(masters, manual_daily, manual_preview, service_date)), data_only=False)
    assert manual_workbook['B_BahanBaku']['K6'].value == 'INV/MAJA/001'
    assert manual_workbook['C1_Relawan']['J6'].value == 'KWT/REL/101'
    assert 'KWT/REL/101' in PdfReader(BytesIO(render_document_pdf(manual_docs[3]))).pages[0].extract_text()
    cancelled = deepcopy(docs)
    cancelled[1]["status"] = "CANCELLED"
    reduced = merge_final_documents(imported, cancelled)
    assert all(x["sourceDocumentId"] != docs[1]["id"] for x in reduced["operations"])
    assert docs[1]["id"] not in reduced["_generatedDocumentIds"]
    for doc in cancelled:
        doc["status"] = "CANCELLED"
    cleared = merge_final_documents(imported, cancelled)
    assert all(not cleared[key] for key in ("rawMaterials", "operations", "volunteerPayments", "incentiveRecipients")), "Cancelling last final must remove its costs"
    assert cleared["balance"] == imported["balance"]
    assert not cleared["_generatedDocumentIds"]
    assert merge_final_documents(cleared, cancelled) == cleared
    assert len(merge_final_documents(cleared, docs)["operations"]) == len(imported["operations"]), "Replacement can be imported without double count"
    assert "DIBATALKAN" in PdfReader(BytesIO(render_document_pdf(cancelled[0]))).pages[0].extract_text()
    assert all(x["workDays"] == 1 for x in imported["volunteerPayments"])
    tampered = deepcopy(imported); tampered["operations"][0]["price"] = 9999999
    tampered["operations"][0]["evidenceLink"] = "https://example.test/new-signed.pdf"
    restored = merge_final_documents(tampered, docs)
    assert restored["operations"][0]["price"] == 1000, "Final amounts must remain canonical"
    assert restored["operations"][0]["evidenceLink"].endswith("new-signed.pdf"), "Signed evidence can be completed after printing"
    duplicate = document(6, "UPAH_RELAWAN", [("Relawan Uji A", "Upah", 1, "hari", 90000, {"role": "Tugas berbeda"})])
    try:
        merge_final_documents(imported, docs + [duplicate])
        raise AssertionError("Duplicate daily wage must be refused")
    except ValueError:
        pass
    oversized = document(8, "BAHAN_BAKU", [(f"Item {i}", "Bahan", 1, "kg", 1, {}) for i in range(41)])
    try:
        merge_final_documents({}, [oversized])
        raise AssertionError("Workbook capacity overflow must be refused")
    except ValueError:
        pass
    preview = compute_preview(masters, imported, service_date, True, final_plan)
    assert preview["rawTotal"] == 3500
    assert preview["operationalTotal"] == 228950
    assert len([x for x in preview["register"] if x["source"] in {"B_BahanBaku", "C_Operasional", "C1_Relawan"}]) == 7
    assert all(x["proofStatus"] == "UNIK" for x in preview["register"])
    assert next(x for x in preview["checks"] if x["no"] == "14")["ok"]
    workbook = load_workbook(BytesIO(populate_workbook(masters, imported, preview, service_date)), data_only=False)
    assert workbook["B_BahanBaku"]["K6"].value == docs[0]["documentNumber"]
    assert workbook["B_BahanBaku"]["K7"].value == docs[0]["documentNumber"]
    assert workbook["B_BahanBaku"]["I6"].value == "=ROUND(F6*H6,2)"
    assert workbook["C_Operasional"]["G13"].value == 700, "Both gas invoices must sum into the Gas category"
    assert workbook["C_Operasional"]["E13"].value == 1
    assert workbook["C_Operasional"]["H13"].value == "=ROUND(E13*G13,2)"
    assert docs[1]["documentNumber"] in workbook["C_Operasional"]["I13"].value
    assert docs[2]["documentNumber"] in workbook["C_Operasional"]["I13"].value
    assert workbook["C1_Relawan"]["J6"].value == docs[3]["documentNumber"] + "-001"
    assert workbook["C_Operasional"]["I7"].value == docs[4]["documentNumber"] + "-001"
    assert workbook["C_Operasional"]["I8"].value == docs[4]["documentNumber"] + "-002"
    assert workbook["C_Operasional"]["H6"].value.startswith("="), "Relawan formulas stay active"
    # Re-run with an installed official-like template; the populated register must
    # still use the actual invoice/receipt numbers, with validation formulas intact.
    from lpdh_logic import fallback_lpdh_workbook
    buf = BytesIO(); fallback_lpdh_workbook().save(buf)
    official = load_workbook(BytesIO(populate_workbook(masters, imported, preview, service_date, template_bytes=buf.getvalue())))
    assert official["I_RegisterBukti"]["C5"].value == docs[0]["documentNumber"]
    assert official["I_RegisterBukti"]["G5"].value.startswith("=")
    for doc in (docs[0], docs[1], docs[3], docs[4]):
        pdf = render_document_pdf(doc)
        reader = PdfReader(BytesIO(pdf))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert doc["documentNumber"] in text and "FINAL" in text
        for item in doc["items"]:
            assert item["itemName"] in text
        if doc["documentType"] in {"UPAH_RELAWAN", "INSENTIF_GURU_KADER"}:
            assert doc["documentNumber"] + "-001" in text
            assert doc["documentNumber"] + "-002" in text
            assert "1 hari" in text
        if output_dir:
            Path(output_dir).mkdir(parents=True, exist_ok=True)
            (Path(output_dir) / (doc["documentType"] + ".pdf")).write_bytes(pdf)
    large = document(9, "BAHAN_BAKU", [(f"Nama bahan panjang untuk uji pemisahan halaman {i}", "Bahan", 1, "kg", 1000, {}) for i in range(40)])
    assert len(PdfReader(BytesIO(render_document_pdf(large))).pages) >= 2
    cadre_batch = document(10, "INSENTIF_GURU_KADER", [(f"Kader {i}", "Insentif", 1, "hari", 1000, {"recipientType": "Kader"}) for i in range(12)])
    assert len(PdfReader(BytesIO(render_document_pdf(cadre_batch))).pages) >= 6
    print("DOCUMENT SELFTEST OK: draft exclusion, idempotent import, daily wages, duplicate guard, category totals, register/workbook parity, PDFs and pagination")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--output-dir")
    run(parser.parse_args().output_dir)
