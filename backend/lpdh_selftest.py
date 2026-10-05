from __future__ import annotations

from io import BytesIO
from openpyxl import load_workbook

import importlib.util
from pathlib import Path

_LOGIC_PATH = Path(__file__).with_name("lpdh_logic.py")
_spec = importlib.util.spec_from_file_location("lpdh_logic_standalone", _LOGIC_PATH)
_logic = importlib.util.module_from_spec(_spec)
assert _spec and _spec.loader
_spec.loader.exec_module(_logic)

compute_preview = _logic.compute_preview
make_master_template = _logic.make_master_template
parse_master_workbook = _logic.parse_master_workbook
populate_workbook = _logic.populate_workbook
fallback_lpdh_workbook = _logic.fallback_lpdh_workbook


def fixture():
    service_date = "2026-10-05"
    masters = {
        "identity": {
            "lpdhNumber": "LPDH-TEST-001",
            "sppgId": "SPPG-TEST",
            "sppgName": "SPPG TEST",
            "village": "Cemplang",
            "district": "Test",
            "city": "Kabupaten Test",
            "province": "Jawa Barat",
            "foundation": "Yayasan Test",
            "vaNumber": "8888888888",
            "bankName": "Bank Test",
        },
        "signers": [
            {"name": "Pengawas Keuangan", "identityType": "NIK", "identityNumber": "3201010101010101", "signed": "Ya"},
            {"name": "Kepala SPPG", "identityType": "NIP", "identityNumber": "199001012020011001", "signed": "Ya"},
            {"name": "Perwakilan Yayasan", "identityType": "NIK", "identityNumber": "3201010101010102", "signed": "Ya"},
        ],
        "beneficiaries": [
            {"code": "SKL-01", "groupCode": "KS-01", "targetPm": 100, "status": "Aktif"},
        ],
        "volunteers": [
            {"code": "RL-001", "name": "Relawan Test", "role": "Juru Masak", "dailyRate": 1000, "status": "Aktif", "paymentMethod": "Transfer"},
        ],
        "operations": [
            {"code": "OP-001", "name": "Gas", "unit": "tabung", "defaultPrice": 1000, "status": "Aktif"},
        ],
        "parameters": {
            "incentiveTariff": 2000,
            "rawSmall": 8000,
            "rawLarge": 10000,
            "operationalPerPm": 3000,
            "maxVa": 500000000,
            "maxHpePerWeek": 5,
            "uploadHour": 6,
            "dateTolerance": 1,
            "applyIndexRaw": True,
            "applyIndexOp": True,
            "cityIndex": 1,
        },
    }
    daily = {
        "dayStatus": "HPE",
        "hpeNumber": 1,
        "pm": {
            "rows": [{
                "code": "KS-01", "distributed": 100, "received": 100, "notReceived": 0,
                "bnba": "Ya", "bastNo": "BAST-001", "bastLink": "https://example.test/bast"
            }],
            "production": {"produced": 102, "organoleptic": 1, "retainedSample": 1, "notDistributed": 0, "buffer": 0},
        },
        "rawMaterials": [{
            "date": service_date, "name": "Beras", "category": "Bahan", "qty": 1, "unit": "kg",
            "price": 1000, "supplier": "Vendor Test", "invoiceNo": "INV-RAW-001",
            "evidenceLink": "https://example.test/raw", "note": ""
        }],
        "operations": [{
            "date": service_date, "itemCode": "OP-001", "description": "Gas", "qty": 1, "unit": "tabung",
            "price": 1000, "invoiceNo": "INV-OP-001", "evidenceLink": "https://example.test/op"
        }],
        "volunteerReceiptBaseNo": "KWT-RL-001",
        "volunteerPayments": [{
            "volunteerCode": "RL-001", "name": "Relawan Test", "role": "Juru Masak",
            "date": service_date, "workDays": 1, "dailyRate": 1000, "paymentMethod": "Transfer",
            "receiptNo": "KWT-RL-001", "evidenceLink": "https://example.test/relawan"
        }],
        "incentiveRecipients": [],
        "incentive": {
            "eligibility": {"contamination": "Tidak", "fatalIncident": "Tidak", "suspended": "Tidak", "verified": "Ya", "pmInputSipgn": "Ya"},
            "ppkStatementNo": "PPK-001", "statementAmount": 204000, "paidAmount": 204000,
            "paymentDate": service_date, "proofNo": "IN-001", "receiptNo": "KWT-IN-001",
            "receiptSigned": "Ya", "evidenceLink": "https://example.test/insentif", "vaReference": "VA-001"
        },
        "balance": {"openingRaw": 10000, "openingOperational": 10000, "openingIncentive": 500000, "bankBalance": 313000},
        "topups": [],
        "upload": {"date": "2026-10-06", "time": "05:00"},
    }
    final_plan = {
        "sourcePlanId": "plan-test",
        "planName": "Final Test",
        "payload": {
            "date": service_date, "porsiKecil": 100, "porsiBesar": 0,
            "shoppingListJSON": {"shoppingList": [{"item": "Beras", "jumlah": 1, "satuan": "kg", "harga_satuan": 1000}]},
        },
        "revision": 1,
    }
    return service_date, masters, daily, final_plan


def run():
    service_date, masters, daily, final_plan = fixture()

    invalid = dict(daily)
    invalid["rawMaterials"] = [dict(daily["rawMaterials"][0], invoiceNo="IN-001")]
    bad = compute_preview(masters, invalid, service_date, True, final_plan)
    assert bad["ready"] is False, "duplicate/invalid fixture should be blocked"
    assert any(not row["ok"] for row in bad["checks"]), "bad fixture must have validation failures"

    good = compute_preview(masters, daily, service_date, True, final_plan)
    assert good["ready"] is True, [x for x in good["checks"] if not x["ok"]]
    assert good["errorCount"] == 0
    assert good["rawStatus"] == "DALAM PAGU"
    assert good["operationalStatus"] == "DALAM PAGU"
    assert good["hpeEligible"] is True
    assert good["topup"]["requiredTotal"] == good["rawTotal"] + good["operationalTotal"] + good["incentiveCalculated"]
    assert good["checks"][5]["status"] == "BELUM DITETAPKAN"
    assert all(row["proofStatus"] == "UNIK" for row in good["register"])

    content = populate_workbook(masters, daily, good, service_date)
    wb = load_workbook(BytesIO(content), data_only=False)
    expected = ["Petunjuk","Identitas","A_PM","B_BahanBaku","C_Operasional","C1_Relawan","D_Insentif","E_Saldo","F_TopUp","G_CekPPK","H_RekapPPK","I_RegisterBukti","J_Pengesahan","Ref"]
    assert wb.sheetnames == expected
    assert isinstance(wb["Identitas"]["B16"].value, str) and wb["Identitas"]["B16"].value.startswith("=")
    assert wb["B_BahanBaku"]["I6"].value.startswith("=")
    assert wb["C_Operasional"]["H6"].value.startswith("=")
    assert wb["C_Operasional"]["E6"].value == "='C1_Relawan'!C66"
    assert wb["D_Insentif"]["B6"].value == "Tidak"
    assert wb["D_Insentif"]["C6"].value.startswith("=")
    assert wb["F_TopUp"]["B5"].value.startswith("=")
    assert wb["F_TopUp"]["C5"].value is None
    assert wb["G_CekPPK"]["C5"].value.startswith("=")
    assert wb["I_RegisterBukti"]["G5"].value.startswith("=")

    base_wb = fallback_lpdh_workbook()
    base_io = BytesIO()
    base_wb.save(base_io)
    official_like = populate_workbook(masters, daily, good, service_date, template_bytes=base_io.getvalue())
    official_wb = load_workbook(BytesIO(official_like), data_only=False)
    assert official_wb["G_CekPPK"]["C5"].value.startswith("="), "official-template formulas must be preserved"
    assert official_wb["I_RegisterBukti"]["G5"].value.startswith("="), "official register formulas must be preserved"

    template = make_master_template()
    parsed = parse_master_workbook(template)
    assert len(parsed["beneficiaries"]) == 0
    assert len(parsed["volunteers"]) == 0, "blank template must not import invented people"
    assert len(parsed["schools"]) == len(parsed["posyandu"]) == 0
    template_wb = load_workbook(BytesIO(template))
    assert template_wb["Master_Sekolah"]["F1"].value == "Porsi Kecil"
    assert template_wb["Master_Sekolah"]["G1"].value == "Porsi Besar"
    assert len(parsed["operations"]) > 0

    print("LPDH SELFTEST OK: official semantics, invalid-blocking, 26 validations, formula workbook, master import")


if __name__ == "__main__":
    run()
