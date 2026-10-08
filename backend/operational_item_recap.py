"""Read-only item recap and application Excel export from saved FINAL invoices."""
from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO


def operational_recap(documents, start_date, end_date):
    groups, details, seen = {}, [], set()
    for doc in documents:
        if doc.get('status') != 'FINAL' or doc.get('documentType') != 'OPERASIONAL':
            continue
        day = str(doc.get('serviceDate') or '')[:10]
        if not str(start_date) <= day <= str(end_date):
            continue
        identity = doc.get('id')
        if identity is not None and identity in seen:
            continue
        seen.add(identity)
        for item in doc.get('items') or []:
            name = str(item.get('itemName') or '').strip() or 'Tanpa nama item'
            unit = str(item.get('unit') or '').strip() or 'Tidak diisi'
            category = str(item.get('category') or '').strip()
            key = (name.casefold(), unit.casefold(), category.casefold())
            qty = Decimal(str(item.get('quantity') or 0))
            amount = Decimal(str(item.get('lineTotal') or 0))
            group = groups.setdefault(key, dict(itemName=name, unit=unit, category=category,
                quantity=Decimal(0), amount=Decimal(0), invoices=set(), dates=set()))
            group['quantity'] += qty
            group['amount'] += amount
            group['invoices'].add(identity if identity is not None else doc.get('documentNumber'))
            group['dates'].add(day)
            details.append(dict(serviceDate=day, documentNumber=doc.get('documentNumber') or '',
                documentId=identity, itemName=name, category=category, unit=unit,
                quantity=float(qty), unitPrice=float(item.get('unitPrice') or 0), amount=float(amount)))
    rows = [dict(itemName=g['itemName'], category=g['category'], unit=g['unit'],
        quantity=float(g['quantity']), amount=float(g['amount']), invoiceCount=len(g['invoices']),
        dayCount=len(g['dates'])) for _, g in sorted(groups.items())]
    return dict(items=rows, details=details, totalAmount=float(sum((g['amount'] for g in groups.values()), Decimal(0))),
        invoiceCount=len({(d['documentId'],d['documentNumber']) for d in details}), itemCount=len(rows))


def recap_excel(recap, site, start_date, end_date):
    # Reuse the application's existing XLSX writer and safe text handling.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from backend.generated_document_excel import put, MONEY
    wb = Workbook()
    wb.remove(wb.active)
    definitions = [
        ('Rekap Item', ['Item','Kategori','Jumlah','Satuan','Total biaya (Rp)','Jumlah invoice','Hari pelayanan'],
         [[x['itemName'],x['category'],x['quantity'],x['unit'],x['amount'],x['invoiceCount'],x['dayCount']] for x in recap['items']], [3], [5]),
        ('Rincian Invoice', ['Tanggal pelayanan','Nomor invoice','Item','Kategori','Jumlah','Satuan','Harga satuan (Rp)','Nilai (Rp)','ID dokumen'],
         [[date.fromisoformat(x['serviceDate']),x['documentNumber'],x['itemName'],x['category'],x['quantity'],x['unit'],x['unitPrice'],x['amount'],x['documentId']] for x in recap['details']], [5], [7,8])]
    for name, headers, rows, qty_cols, money_cols in definitions:
        ws=wb.create_sheet(name)
        put(ws,1,1,f'Rekap Item Operasional {site}',True,15)
        put(ws,2,1,f'Periode {start_date} sampai {end_date} — hanya invoice FINAL')
        put(ws,3,1,'Jumlah invoice'); put(ws,3,2,recap['invoiceCount'])
        put(ws,3,4,'Total biaya (Rp)'); put(ws,3,5,recap['totalAmount']).number_format=MONEY
        for col, heading in enumerate(headers,1):
            cell=put(ws,5,col,heading,True)
            cell.fill=PatternFill('solid',fgColor='14513D');cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
        for rowno, values in enumerate(rows,6):
            for col,value in enumerate(values,1):
                cell=put(ws,rowno,col,value)
                if col in money_cols: cell.number_format=MONEY
                elif col in qty_cols: cell.number_format='#,##0.####'
                elif isinstance(value,date): cell.number_format='dd-mm-yyyy'
        for col in range(1,len(headers)+1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col)].width=34 if headers[col-1] in ('Item','Nomor invoice') else 22
        ws.row_dimensions[5].height=32
        ws.freeze_panes='A6';ws.auto_filter.ref=f'A5:{get_column_letter(len(headers))}{max(5,5+len(rows))}'
        ws.sheet_view.showGridLines=False
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_title_rows='1:5'
    output=BytesIO();wb.save(output);return output.getvalue()

