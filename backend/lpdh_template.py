"""Prepare a private server template once; exports replace input values only.

The uploaded bytes are retained separately. Never re-create formulas, styles,
register structure or worksheets while filling a prepared template.
"""
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import math
import posixpath
import re
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.defined_name import DefinedName

VERSION = 1
MARKER = '_LPDH_FILL_ONLY_VERSION'
YELLOW = 'FFFFF2CC'
APPENDIX_ROWS = 500
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def input_cells():
    cells = {}
    def add(sheet, rows, columns):
        cells.setdefault(sheet, set()).update(f'{c}{r}' for r in rows for c in columns)
    add('Identitas', range(5,16), 'B')
    add('Identitas', [17,18,19,25,26], 'B')
    add('Identitas', range(36,39), 'BCDF')
    add('A_PM', range(6,16), 'EFGHIKLM')
    add('A_PM', range(6,16), 'R')
    add('A_PM', [19,21,22,23,24], 'C')
    add('B_BahanBaku', range(6,46), 'CDEFGHJKLO')
    add('C_Operasional', range(7,22), 'CEFGIJM')
    add('C_Operasional', [6], 'M')
    add('C1_Relawan', range(6,66), 'CDEFGIJK')
    add('D_Insentif', range(6,11), 'B')
    add('D_Insentif', range(19,28), 'C')
    add('E_Saldo', [5,6,7], 'B')
    add('E_Saldo', [10], 'E')
    add('E_Saldo', [12], 'E')
    add('E_Saldo', [5,6,7], 'F')
    add('E_Saldo', range(19,24), 'CDEFGIJ')
    add('Ref', range(5,16), 'B')
    add('Ref', range(65,214), 'ABCD')
    add('I_RegisterBukti', range(5,131), 'ABCDEF')
    add('Lampiran_Dokumen', range(2,APPENDIX_ROWS+2), 'ABCDEFGHI')
    return cells


HEADERS = {
    ('A_PM','E5'): 'Target PM (SPS)', ('A_PM','K5'): 'Data BNBA Tersedia',
    ('B_BahanBaku','D5'): 'Nama Bahan', ('B_BahanBaku','K5'): 'No. Bukti/Nota',
    ('C_Operasional','G5'): 'Harga Satuan (Rp)',
    ('C1_Relawan','C5'): 'Nama Relawan', ('C1_Relawan','J5'): 'No. Bukti Pembayaran',
}


def prepare_template(raw):
    """Validate mapping BEFORE replacement. Changes are made to a copy only."""
    wb = load_workbook(BytesIO(raw), data_only=False)
    required = {'Identitas','A_PM','B_BahanBaku','C_Operasional','C1_Relawan',
                'D_Insentif','E_Saldo','F_TopUp','G_CekPPK','H_RekapPPK',
                'I_RegisterBukti','J_Pengesahan','Ref'}
    missing = sorted(required.difference(wb.sheetnames))
    if missing:
        raise ValueError('Gunakan workbook LPDH lengkap, bukan template impor master. Sheet wajib tidak ada: ' + ', '.join(missing))
    for (sheet, cell), label in HEADERS.items():
        if str(wb[sheet][cell].value or '').strip().casefold() != label.casefold():
            raise ValueError(f'Susunan template berubah pada {sheet}!{cell}. Template lama tetap dipakai; pemetaan perlu disesuaikan.')
    mapped = input_cells()
    for sheet in wb.sheetnames:
        if sheet == 'Petunjuk': continue  # The yellow legend is not an input.
        for row in wb[sheet]:
            for cell in row:
                if cell.fill.fgColor.type == 'rgb' and cell.fill.fgColor.rgb == YELLOW and cell.coordinate not in mapped.get(sheet,set()):
                    raise ValueError(f'Sel input baru belum dipetakan: {sheet}!{cell.coordinate}. Template lama tetap dipakai sampai pemetaan disesuaikan.')
    # The source register was line-based. One package now has one register entry,
    # while recipient/item details keep the shared document number unchanged.
    if 'Lampiran_Dokumen' not in wb.sheetnames:
        ws = wb.create_sheet('Lampiran_Dokumen')
        ws.append(['Sheet sumber','ID dokumen','Baris dokumen','Nomor invoice/kuitansi',
                   'Item/penerima','Tanggal','Nominal (Rp)','Link bukti','Referensi pembayaran'])
        for c in ws[1]: c.font = Font(bold=True)
        ws.freeze_panes = 'E2'
        ws.sheet_view.showGridLines = False
        for c, width in {'A':21,'B':18,'C':18,'D':32,'E':40,'F':15,'G':20,'H':60,'I':36}.items():
            ws.column_dimensions[c].width = width
    inputs = input_cells()
    for sheet, addresses in inputs.items():
        for address in addresses:
            cell = wb[sheet][address]
            if isinstance(cell, MergedCell):
                raise ValueError(f'Sel input berpindah/tergabung: {sheet}!{address}. Sesuaikan pemetaan terlebih dahulu.')
            if sheet != 'I_RegisterBukti' and cell.data_type == 'f':
                raise ValueError(f'Sel input {sheet}!{address} berisi rumus. Template baru tidak dipasang agar rumus tidak tertimpa.')
            # Reference lookup data is not prior-day transaction data.
            if sheet != 'Ref':
                cell.value = None
                cell.hyperlink = None
                cell.comment = None
            cell.fill = PatternFill('solid', fgColor=YELLOW)
            if sheet == 'Lampiran_Dokumen':
                if address.startswith('F'): cell.number_format = 'dd-mm-yyyy'
                if address.startswith('G'): cell.number_format = '#,##0.00'
    for r in range(5,131):
        wb['I_RegisterBukti'][f'G{r}'] = f'=IF(C{r}="","",IF(COUNTIF($C$5:$C$130,C{r})>1,"DUPLIKAT","Unik"))'
    wb.defined_names[MARKER] = DefinedName(MARKER, attr_text=str(VERSION))
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    out = BytesIO(); wb.save(out)
    return out.getvalue()


def source_hash(raw):
    return sha256(raw).hexdigest()


def _values(masters, daily, preview, service_date, wb):
    from backend.lpdh_logic import (as_number, as_int, yes, parse_time, _excel_date,
                                   assign_document_numbers, grouped_operations, OPERATIONAL_DEFAULTS)
    values = {s: {c: None for c in cc} for s, cc in input_cells().items()}
    # Keep the prepared template's reference lookup catalog. Override selected
    # kitchen parameters below without blanking other reference records.
    values['Ref'] = {}
    def put(sheet, cell, value): values[sheet][cell] = value
    identity = masters.get('identity') or {}
    for c,v in {'B5':daily.get('lpdhNumber') or identity.get('lpdhNumber'), 'B6':identity.get('sppgId'),
                'B7':identity.get('sppgName'),'B8':identity.get('village'),'B9':identity.get('district'),
                'B10':identity.get('city'),'B11':identity.get('province'),'B12':identity.get('foundation'),
                'B13':identity.get('vaNumber'),'B14':identity.get('bankName'),'B15':_excel_date(service_date),
                'B17':as_int(daily.get('weekOfMonth'),(date.fromisoformat(service_date).day-1)//7+1),
                'B18':daily.get('dayStatus') or 'HPE','B19':as_int(daily.get('hpeNumber'),1),
                'B25':_excel_date((daily.get('upload') or {}).get('date')),
                'B26':parse_time((daily.get('upload') or {}).get('time'))}.items(): put('Identitas',c,v)
    for r,item in enumerate((masters.get('signers') or [])[:3],36):
        for c,v in {'B':item.get('name'),'C':item.get('identityType'),'D':item.get('identityNumber'),
                    'F':'Ya' if yes(item.get('signed')) else 'Tidak'}.items(): put('Identitas',f'{c}{r}',v)
    params = preview['parameters']
    for r,key,default in [(5,'incentiveTariff',2000),(6,'rawSmall',8000),(7,'rawLarge',10000),
                          (8,'operationalPerPm',3000),(9,'maxVa',500000000),(10,'maxHpePerWeek',5),
                          (11,'uploadHour',6),(13,'dateTolerance',1)]:
        put('Ref',f'B{r}',as_number(params.get(key),default))
    put('Ref','B12',None if params.get('bufferPct') in (None,'') else as_number(params['bufferPct']))
    for r,key in [(14,'applyIndexRaw'),(15,'applyIndexOp')]: put('Ref',f'B{r}','Ya' if params.get(key,True) else 'Tidak')
    city = str(identity.get('city') or '').strip()
    if city:
        rr = next((r for r in range(65,214) if str(wb['Ref'][f'B{r}'].value or '').strip().casefold()==city.casefold()),None)
        if rr is None: rr = next((r for r in range(65,214) if not wb['Ref'][f'B{r}'].value),None)
        if rr is None: raise ValueError('Tabel indeks kota pada template penuh. Sesuaikan template terlebih dahulu.')
        for c,v in {'A':identity.get('province'),'B':city,'C':as_number(params.get('cityIndex'),1) or 1,
                    'D':params.get('cityIndexSource') or 'Master LPDH'}.items(): put('Ref',f'{c}{rr}',v)
    for r,item in enumerate(preview['pmRows'],6):
        for c,key in {'E':'targetPm','F':'distributed','G':'received','H':'notReceived','I':'reason','K':'bnba','L':'bastNo','M':'bastLink'}.items():
            put('A_PM',f'{c}{r}',item[key])
    for r,key in [(19,'produced'),(21,'organoleptic'),(22,'retainedSample'),(23,'notDistributed'),(24,'buffer')]:
        put('A_PM',f'C{r}',preview['production'][key])
    raw = assign_document_numbers(preview['rawMaterials'],'invoiceNo')
    volunteers = assign_document_numbers([{**x,'receiptBase':x.get('receiptNo') or daily.get('volunteerReceiptBaseNo') or ''} for x in preview['volunteers']],'receiptBase')
    if len(raw)>40 or len(volunteers)>60 or len(daily.get('topups') or [])>5:
        raise ValueError('Jumlah rincian melebihi kapasitas template (40 bahan, 60 relawan, 5 top up). Tidak ada baris yang dipotong; unggah template yang telah disesuaikan.')
    def note(item):
        return '; '.join(filter(None,[item.get('note'),f"Ref pembayaran: {item['paymentReference']}" if item.get('paymentReference') else '']))
    for r,item in enumerate(raw,6):
        for c,v in {'C':_excel_date(item.get('date')),'D':item.get('name'),'E':item.get('category'),
                    'F':as_number(item.get('qty')),'G':item.get('unit'),'H':as_number(item.get('price')),
                    'J':item.get('supplier'),'K':item.get('proofNoDerived') or item.get('_baseProofNo') or '',
                    'L':item.get('evidenceLink'),'O':note(item)}.items(): put('B_BahanBaku',f'{c}{r}',v)
    for r,item in enumerate(volunteers,6):
        for c,v in {'C':item.get('name'),'D':item.get('role'),'E':_excel_date(item.get('date')),
                    'F':as_number(item.get('workDays')),'G':as_number(item.get('dailyRate')),
                    'I':item.get('paymentMethod'),'J':item.get('proofNoDerived') or item.get('_baseProofNo') or '',
                    'K':item.get('evidenceLink')}.items(): put('C1_Relawan',f'{c}{r}',v)
    groups = {str(x.get('description') or '').strip().lower():x for x in grouped_operations(assign_document_numbers(preview['operations'],'invoiceNo'))}
    recipients = preview['incentiveRecipients']
    for r,types,suffix,key in [(7,{'guru','sekolah','penanggung jawab satuan pendidikan'},'GURU','schoolPicOperationalProofNo'),
                               (8,{'kader','posyandu','kader posyandu'},'KADER','cadreOperationalProofNo')]:
        rows = [x for x in recipients if str(x.get('type') or '').strip().lower() in types]
        if not rows: continue
        sourced = any(x.get('sourceDocumentId') for x in rows)
        base = daily.get('incentiveReceiptBaseNo') or ''
        join = lambda k: '; '.join(dict.fromkeys(x.get(k) or '' for x in rows if x.get(k)))
        groups[OPERATIONAL_DEFAULTS[r-6].lower()] = {
            'date':daily.get('incentivePaymentDate') or service_date,'qty':len(rows),
            'unit':'satuan' if r==7 else 'orang','price':sum(as_number(x.get('amount')) for x in rows)/len(rows),
            'invoiceNo':join('receiptNo') if sourced else daily.get(key) or (f'{base}-{suffix}' if base else ''),
            'evidenceLink':join('evidenceLink') if sourced else daily.get('incentiveBatchEvidenceLink') or '',
            'paymentReference':join('paymentReference'),
            'note':f'Lampiran: {len(rows)} penerima; Kuitansi gabungan' if any(x.get('aggregatePayment') for x in rows) else f'Kuitansi individual: {len(rows)} lembar'}
    for r,label in enumerate(OPERATIONAL_DEFAULTS,6):
        if r==6: continue  # This entire row remains linked to C1_Relawan.
        item = groups.get(label.lower())
        if not item: continue
        for c,v in {'C':_excel_date(item.get('date')),'E':as_number(item.get('qty')),'F':item.get('unit'),
                    'G':as_number(item.get('price')),'I':item.get('invoiceNo') or item.get('proofNoDerived') or item.get('_baseProofNo') or '',
                    'J':item.get('evidenceLink'),'M':note(item)}.items(): put('C_Operasional',f'{c}{r}',v)
    ins = daily.get('incentive') or {}; elig = ins.get('eligibility') or {}
    for r,key in [(6,'contamination'),(7,'fatalIncident'),(8,'suspended'),(9,'verified'),(10,'pmInputSipgn')]:
        put('D_Insentif',f'B{r}','Ya' if yes(elig.get(key)) else 'Tidak')
    for r,v in {19:ins.get('ppkStatementNo'),20:as_number(ins.get('statementAmount')),21:as_number(ins.get('paidAmount')),
                22:_excel_date(ins.get('paymentDate')),23:ins.get('proofNo'),24:ins.get('receiptNo'),
                25:'Ya' if yes(ins.get('receiptSigned')) else 'Tidak',26:ins.get('evidenceLink'),27:ins.get('vaReference')}.items(): put('D_Insentif',f'C{r}',v)
    bal = daily.get('balance') or {}
    for c,key in {'B5':'openingRaw','B6':'openingOperational','B7':'openingIncentive','E10':'bankBalance'}.items(): put('E_Saldo',c,as_number(bal.get(key)))
    for r,item in enumerate(daily.get('topups') or [],19):
        for c,v in {'C':_excel_date(item.get('date')),'D':item.get('reference'),'E':as_number(item.get('rawAmount')),
                    'F':as_number(item.get('operationalAmount')),'G':as_number(item.get('incentiveAmount')),
                    'I':item.get('receiptNo'),'J':item.get('evidenceLink')}.items(): put('E_Saldo',f'{c}{r}',v)
    register = preview.get('register') or []
    evidence = [(s,x) for s,k in [('B_BahanBaku','rawMaterials'),('C_Operasional','operations'),('C1_Relawan','volunteers'),('C_Operasional','incentiveRecipients')] for x in preview.get(k) or [] if x.get('sourceDocumentId')]
    if len(register)>126 or len(evidence)>APPENDIX_ROWS:
        raise ValueError('Register/lampiran melebihi kapasitas template; sesuaikan template terlebih dahulu. Data tidak dipotong.')
    for r,item in enumerate(register,5):
        for c,v in {'A':item.get('source'),'B':item.get('code'),'C':item.get('proofNo'),'D':_excel_date(item.get('date')),
                    'E':as_number(item.get('amount')),'F':item.get('link')}.items(): put('I_RegisterBukti',f'{c}{r}',v)
    for r,(sheet,item) in enumerate(evidence,2):
        for c,v in zip('ABCDEFGHI',[sheet,f"DOC-{item['sourceDocumentId']}",item.get('sourceLine'),
                                    item.get('invoiceNo') or item.get('receiptNo'),item.get('name') or item.get('description'),
                                    _excel_date(item.get('date')),as_number(item.get('amount')),item.get('evidenceLink') or '',item.get('paymentReference') or '']):
            put('Lampiran_Dokumen',f'{c}{r}',v)
    return values


def fill_template(template, masters, daily, preview, service_date):
    wb = load_workbook(BytesIO(template),data_only=False)
    marker = wb.defined_names.get(MARKER)
    if marker is None or marker.attr_text != str(VERSION):
        raise ValueError('Template server belum disiapkan. Unggah ulang template LPDH resmi.')
    values = _values(masters,daily,preview,service_date,wb)
    allowed = input_cells()
    # Patch the existing XLSX package. All non-worksheet parts, including styles,
    # drawings, names and calculation settings, remain byte-for-byte identical.
    with ZipFile(BytesIO(template)) as src:
        root = ET.fromstring(src.read('xl/workbook.xml'))
        rels = {x.attrib['Id']:x.attrib['Target'] for x in ET.fromstring(src.read('xl/_rels/workbook.xml.rels'))}
        paths = {s.attrib['name']:posixpath.normpath('xl/'+rels[s.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']].lstrip('/')) for s in root.find('s:sheets',NS)}
        paths = {s:p[3:] if p.startswith('xl/xl/') else p for s,p in paths.items()}
        updates = {}
        for sheet, inputs in values.items():
            xml = src.read(paths[sheet]).decode('utf-8')
            found = set()
            def replace(match):
                address = match.group('address')
                if address not in inputs: return match.group()
                value = inputs[address]
                found.add(address)
                cell = wb[sheet][address]
                if address not in allowed[sheet] or cell.data_type=='f' or cell.fill.fgColor.rgb!=YELLOW:
                    raise ValueError(f'Penulisan di luar sel input kuning ditolak: {sheet}!{address}')
                attrs = re.match(r'<c\b([^>]*?)(?:/?>)',match.group()).group(1).rstrip('/ ')
                attrs = re.sub(r'\s+t="[^"]*"','',attrs)
                if value is None: body = ''; typ = ''
                elif isinstance(value,str):
                    # Literal text including '=' is NEVER a formula.
                    if len(value)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value):
                        raise ValueError(f'Teks tidak kompatibel dengan Excel: {sheet}!{address}')
                    body = '<is><t xml:space="preserve">'+escape(value)+'</t></is>'; typ = ' t="inlineStr"'
                elif isinstance(value,bool): body = f'<v>{int(value)}</v>'; typ = ' t="b"'
                else:
                    if isinstance(value,(date,datetime,time)): value = to_excel(value,wb.epoch)
                    if not isinstance(value,(int,float)) or not math.isfinite(value): raise ValueError(f'Nilai input tidak valid: {sheet}!{address}')
                    body = f'<v>{value}</v>'; typ = ' t="n"'
                return f'<c{attrs}{typ}>{body}</c>'
            xml = re.sub(r'<c\b[^>]*?\br="(?P<address>[A-Z]+[0-9]+)"[^>]*?(?:/>|>.*?</c>)',replace,xml,flags=re.DOTALL)
            if found != set(inputs):
                raise ValueError(f'Sel input tidak ditemukan pada template {sheet}: ' + ', '.join(sorted(set(inputs)-found)[:5]))
            updates[paths[sheet]] = xml.encode('utf-8')
        out = BytesIO()
        with ZipFile(out,'w',ZIP_DEFLATED) as dst:
            for info in src.infolist(): dst.writestr(info,updates.get(info.filename,src.read(info.filename)))
    return out.getvalue()
